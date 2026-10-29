"""Informes por área con gráficos, tendencias y recomendaciones (observaciones al demo v0.5, 13-oct)."""
from urllib.parse import urlencode

from flask import Blueprint, abort, flash, g, redirect, render_template, request, send_file, url_for

from ..services import charts
from ..services import reports as rsvc
from .common import f_date, login_required

bp = Blueprint("rep", __name__, url_prefix="/informes")

# informe -> (título, permisos que lo habilitan)
INFORMES = {
    "gestion": ("Gestión comercial", ("informes.comercial",)),
    "comercial": ("Comercial", ("informes.comercial",)),
    "reempaque": ("Reempaque", ("informes.reempaque",)),
    "produccion": ("Producción", ("informes.produccion",)),
    "distribucion": ("Distribución", ("informes.distribucion",)),
    "almacen": ("Movimiento de stock del almacén", ("informes.almacen",)),
    "cumplimiento": ("Cumplimiento por persona", ("informes.cumplimiento",)),
    "merma": ("Merma y remanente de MP", ("informes.produccion",)),
}
# V12: tableros de Producción que también son informes de su área (se abren en su propia pantalla)
EXTERNOS = {"control": ("Control de producción", ("produccion.partes_ver",), "prod.control"),
            "oee": ("Tablero OEE", ("produccion.partes_ver",), "prod.oee_dashboard")}
# V12: informes de Producción; a Producción no le suman los de otras áreas (p. ej. movimiento de stock)
INFORMES_PRODUCCION = ("produccion", "merma", "reempaque", "control", "oee", "cumplimiento")


def _area_produccion():
    return getattr(g.user.role, "area", "") == "Producción"


def disponibles():
    ok = lambda perms: any(p in g.perms for p in perms) or "vista.total" in g.perms  # noqa: E731
    items = [(k, t) for k, (t, perms) in INFORMES.items() if ok(perms)]
    items += [(k, t) for k, (t, perms, _) in EXTERNOS.items() if ok(perms)]
    if _area_produccion():
        items = sorted([x for x in items if x[0] in INFORMES_PRODUCCION], key=lambda x: INFORMES_PRODUCCION.index(x[0]))
    return items


@bp.app_template_global()
def informe_url(k):
    return url_for(EXTERNOS[k][2]) if k in EXTERNOS else url_for("rep.show", nombre=k)


def _periodo():
    try:
        return f_date("desde"), f_date("hasta")
    except Exception:  # noqa: BLE001
        return None, None


def _resumen(fuentes, fuente, cols, r, indicadores, elementos):
    """Resumen legible de la configuración del informe personalizado (se muestra antes del resultado y va en el PDF)."""
    a = request.args
    nom = {c: t for c, t, _ in cols}
    sel = [nom[c] for c in a.getlist("col") if c in nom] or [t for _, t, _ in cols]
    ag = nom.get(a.get("agrupar", ""), "")
    return [("Fuente de datos", fuentes[fuente]), ("Periodo", f"{r['desde']:%d/%m/%Y} – {r['hasta']:%d/%m/%Y}" if r else "—"),
            ("Columnas", ", ".join(sel)), ("Agrupado por", ag or "sin agrupar"), ("Texto buscado", a.get("q") or "—"),
            ("Indicadores", ", ".join(x["nombre"] for x in indicadores) or "ninguno"),
            ("Secciones", ", ".join({"kpis": "totales", "grafico": "gráfico", "tabla": "tabla",
                                     "recomendaciones": "recomendaciones", "indicadores": "indicadores"}.get(e, e)
                                    for e in elementos))]


@bp.route("/personalizado", methods=["GET", "POST"])
@login_required
def custom():
    """Personalizar informe (obs. demo v0.7; rediseño 28-oct-2026): fuente, columnas, periodo, filtro, agrupación e
    indicadores del área; vista previa con el resumen de la configuración; PDF o Excel; plantillas personales y
    compartidas. Solo se ofrece y se calcula lo que el perfil puede ver (matriz de indicadores)."""
    import io
    from werkzeug.datastructures import MultiDict
    from ..services import custom_report as cr
    from ..services import indicators as ind
    from ..services import pdfreport
    from ..services import report_center as rc
    from .common import act
    fuentes = cr.sources_for(g.user)
    if not fuentes:
        abort(403)
    if request.method == "POST":
        f = request.form
        if f.get("accion") == "eliminar":
            return act(lambda: rc.delete_template(g.user, f.get("id")), "Plantilla eliminada.", url_for("rep.custom"))
        args = MultiDict(__import__("urllib.parse", fromlist=["x"]).parse_qsl(f.get("qs", ""), keep_blank_values=False))
        return act(lambda: rc.save_template(g.user, f.get("nombre"), rc.config_from_args(args), f.get("compartir") == "1"),
                   lambda t: f"Plantilla «{t.nombre}» guardada" + (f" y compartida con {t.area}." if t.alcance == "COMPARTIDA" else "."),
                   url_for("rep.custom"), redirect_fn=lambda t: url_for("rep.custom", plantilla=t.id))
    if request.args.get("plantilla") and not request.args.get("fuente"):
        from datetime import date, timedelta
        try:
            t = rc.get_template(g.user, request.args["plantilla"])
        except Exception as e:  # noqa: BLE001
            flash(str(e), "error")
            return redirect(url_for("rep.custom"))
        cfg = dict(t.config)
        q = [(k, v) for k, vs in cfg.items() if k != "dias" for v in (vs if isinstance(vs, list) else [vs])]
        if cfg.get("dias") is not None:
            q.append(("desde", (date.today() - timedelta(days=int(cfg["dias"]))).isoformat()))
        return redirect(url_for("rep.custom") + "?" + urlencode(q + [("ver", "1"), ("cfg", "1"), ("plantilla", t.id)]))
    fuente = request.args.get("fuente") or next(iter(fuentes))
    if fuente not in fuentes:
        abort(403)
    columnas = cr.columns(fuente, g.user)
    desde, hasta = _periodo()
    # sin configurar (primera vez o enlace): los indicadores del informe predeterminado del área; «cfg» marca que el
    # usuario ya eligió (aunque haya desmarcado todos)
    sel_ind = request.args.getlist("ind") if request.args.get("cfg") else sorted(ind.default_set(g.user))
    elementos = request.args.getlist("elem") or ["indicadores", "kpis", "grafico", "tabla", "recomendaciones"]
    r, indicadores = None, []
    if request.args.get("ver") or request.args.get("formato"):
        r = cr.run(g.user, fuente, request.args.getlist("col"), desde, hasta, request.args.get("q", ""),
                   request.args.get("agrupar", ""))
        if "indicadores" in elementos:
            indicadores = ind.values(g.user, sel_ind, r["desde"], r["hasta"])
        nombre = f"informe_{fuente}_{r['hasta']:%Y%m%d}"
        resumen = _resumen(fuentes, fuente, columnas, r, indicadores, elementos)
        notas_pdf = [f"{k}: {v}" for k, v in resumen] + [f"{x['nombre']}: {x['valor']} ({x['detalle']})" for x in indicadores] \
            + r["notas"] + ([f"{k}: {v}" for k, v in r["kpis"]] if "kpis" in elementos else []) + (
            ["Recomendación: " + x for x in r["recomendaciones"]] if "recomendaciones" in elementos else [])
        if request.args.get("formato") == "pdf":
            return send_file(io.BytesIO(pdfreport.table_pdf(r["titulo"], r["columnas"], r["filas"] if "tabla" in elementos else [],
                                                            notas_pdf)),
                             mimetype="application/pdf", as_attachment=True, download_name=nombre + ".pdf")
        if request.args.get("formato") == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font
            wb = Workbook()
            ws = wb.active
            ws.title = r["titulo"][:30]
            ws.append([r["titulo"]])
            ws["A1"].font = Font(bold=True, size=12)
            for n in notas_pdf:
                ws.append([n])
            if "tabla" in elementos:
                ws.append(r["columnas"])
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
                for fila in r["filas"]:
                    ws.append(list(fila))
            out = io.BytesIO()
            wb.save(out)
            return send_file(io.BytesIO(out.getvalue()), as_attachment=True, download_name=nombre + ".xlsx",
                             mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    grafico = None
    if r is not None and "grafico" in elementos:
        med = request.args.get("medida")
        grafico = cr.chart(r, request.args.get("tipo_grafico", "barras"), int(med) if (med or "").isdigit() else None)
    qs = urlencode([(k, v) for k, v in request.args.items(multi=True) if k not in ("ver", "formato", "plantilla")])
    return render_template("informe_personalizado.html", fuentes=fuentes, fuente=fuente, columnas=columnas,
                           r=r, sel=request.args.getlist("col"), items=disponibles(), por_area=cr.sources_by_area(g.user),
                           elementos=elementos, grafico=grafico, grupos_ind=ind.for_user(g.user), sel_ind=sel_ind,
                           indicadores=indicadores, resumen=_resumen(fuentes, fuente, columnas, r, indicadores, elementos) if r else [],
                           plantillas=rc.templates_for(g.user), qs=qs, plantilla_id=request.args.get("plantilla", type=int),
                           predeterminado=rc.default_for(g.user, dict(disponibles()), fuentes))


@bp.route("/")
@login_required
def home():
    """Portada de informes: el informe predeterminado del área y «Personalizar informe», más los demás informes
    autorizados y las plantillas guardadas."""
    from ..services import report_center as rc
    from ..services.custom_report import sources_for
    items = disponibles()
    fuentes = sources_for(g.user)
    if not items and not fuentes:
        abort(403)
    pred = rc.default_for(g.user, dict(items), fuentes)
    return render_template("informes_home.html", items=items, pred=pred, plantillas=rc.templates_for(g.user) if fuentes else [],
                           personalizar=bool(fuentes))


def _descarga(html, titulo):
    """Obs. demo v0.7: cada informe se descarga en PDF (y sus tablas en Excel) tal como se ve en pantalla."""
    import io
    from datetime import date
    from ..services import pdfreport
    formato = request.args.get("formato")
    nombre = "informe_" + "".join(ch if ch.isalnum() else "_" for ch in titulo.lower()).strip("_") + f"_{date.today():%Y%m%d}"
    if formato == "pdf":
        return send_file(io.BytesIO(pdfreport.html_to_pdf(html, f"Informe {titulo}")), mimetype="application/pdf",
                         as_attachment=True, download_name=nombre + ".pdf")
    if formato == "xlsx":
        return send_file(io.BytesIO(pdfreport.html_to_xlsx(html, f"Informe {titulo}")), as_attachment=True,
                         download_name=nombre + ".xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return html


@bp.route("/<nombre>")
@login_required
def show(nombre):
    if nombre not in INFORMES or nombre not in dict(disponibles()):
        abort(403)
    if nombre == "gestion" and request.args.get("formato") == "pptx":
        import io
        from datetime import date
        from ..services.pptxreport import gestion_pptx
        return send_file(io.BytesIO(gestion_pptx(_gestion_data())), as_attachment=True,
                         download_name=f"informe_gestion_comercial_{date.today():%Y%m%d}.pptx",
                         mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation")
    return _descarga(_show(nombre), INFORMES[nombre][0])


def _gestion_data():
    from ..services import commercial_mgmt as cm
    desde, hasta = _periodo()
    a = request.args
    f = {"canales": a.getlist("canal"), "modo_canal": a.get("modo_canal", "incluir"), "tipo": a.get("tipo", ""),
         "vendedores": a.getlist("vendedor"), "clientes": a.getlist("cliente")}
    if g.user.role_code == "EJECUTIVO_VENTAS":  # el ejecutivo ve lo suyo
        f["vendedores"] = [str(g.user.id)]
    return cm.gestion(desde, hasta, f, ver_costos="costos.ver" in g.perms, user=g.user)


def _show(nombre):
    desde, hasta = _periodo()
    if nombre == "gestion":
        from ..services import commercial_mgmt as cm
        d = _gestion_data()
        nv = lambda r: r["nivel"]  # noqa: E731
        graf = {
            "vend": charts.columns([n[:14] for n, _ in d["por"]["vendedor"][:10]],
                                   [("S/", [r["soles"] for _, r in d["por"]["vendedor"][:10]])], "S/", title="Venta por vendedor"),
            "canal": charts.bars([(n, r["soles"]) for n, r in d["por"]["canal"]], "S/", title="Venta por canal"),
            "otif_vend": charts.bars([(n, nv(r)["otif"] or 0) for n, r in d["por"]["vendedor"] if nv(r)["n"]], "%",
                                     title="OTIF por vendedor", color=2),
            "embudo": charts.bars(d["embudo"], "", title="Embudo comercial", color=6),
            "menos": charts.bars([(f"{p.code} {p.descripcion[:16]}", v) for p, v in d["menos_vendidos"]], "S/",
                                 title="Productos con menos venta", color=1),
        }
        return render_template("informe_gestion.html", d=d, graf=graf, items=disponibles(), actual=nombre, op=cm.options(),
                               charts=charts)
    if nombre == "comercial":
        fa = request.args
        vend = fa.get("vendedor", "")
        if g.user.role_code == "EJECUTIVO_VENTAS":  # el ejecutivo ve lo suyo
            vend = g.user.nombre
        d = rsvc.comercial(desde, hasta, ver_costos="costos.ver" in g.perms, canal=fa.get("canal", ""),
                           customer_id=fa.get("customer_id") or None, vendedor=vend, tipo=fa.get("tipo", ""))
        from sqlalchemy import select
        from ..db import db
        from ..models import Customer
        from ..services import kpis as kpisvc
        from ..services import prices as prsvc
        clientes = list(db().execute(select(Customer).order_by(Customer.nombre)).scalars())
        filtros = {"canales": sorted({prsvc.channel_of(c) for c in clientes} - {"", None}), "clientes": clientes,
                   "vendedores": sorted({v for v, _ in d["vendedores"]})}
        k = kpisvc.compute(desde, hasta)
        d["tiempos"] = {"preparacion_h": k.get("tiempo_preparacion_h"), "despacho_h": k.get("tiempo_despacho_h")}
        graf = {"semanal": charts.line([("Ventas S/", d["serie_semanal"])], "S/", title="Ventas semanales"),
                "vend_soles": charts.bars([(k, v["soles"]) for k, v in d["vendedores"]], "S/", title="Ventas por vendedor (S/)"),
                "vend_kg": charts.bars([(k, v["kg"]) for k, v in d["vendedores"]], "kg", title="Ventas por vendedor (kg)", color=2),
                "top": charts.bars([(f"{k} {n[:18]}", v) for k, n, v, _ in d["top_productos"]], "S/", title="Productos más vendidos"),
                "pres": charts.bars([(str(k), v) for k, v in d["presentaciones"]], "", title="Presentaciones preferidas", color=6),
                "margen": charts.bars([(f"{k} {n[:16]}", m) for k, n, m in d["margen_producto"][:15]], "%",
                                      title="Margen por producto (menores primero)", color=2)}
        from ..services.customer_behavior import behavior
        cb = behavior(d["desde"], d["hasta"])
        graf["cumplimiento"] = charts.bars([(x["customer"].nombre[:22], x["cumplimiento"] or 0) for x in cb["clientes"][:15]], "%",
                                           title="Cumplimiento por cliente (entregados completos / pedidos)", color=2)
        graf["top_kg"] = charts.bars([(f"{k} {n[:18]}", v) for k, n, v, _ in d["top_kg"]], "kg", title="Productos más vendidos (kg)", color=2)
        return render_template("informe_comercial.html", d=d, cb=cb, graf=graf, items=disponibles(), actual=nombre, filtros=filtros)
    if nombre == "almacen":
        from ..services.core import warehouse_scopes
        propios = warehouse_scopes().get(g.user.role_code)
        if propios is None and g.user.almacen:
            propios = [g.user.almacen]
        propios = [w for w in (propios or []) if not w.endswith("*")]
        wh = request.args.get("almacen") or ""
        if propios and wh and wh not in propios:
            abort(403)
        d = rsvc.almacen(desde, hasta, wh or None, propios)
        graf = charts.line([("Stock total (unidades)", d["serie"])], "und", title="Stock en cada carga de SAP")
        return render_template("informe_almacen.html", d=d, graf=graf, propios=propios, items=disponibles(), actual=nombre)
    if nombre == "produccion":
        d = rsvc.produccion(desde, hasta)
        cap = rsvc.produccion_capacidad(desde, hasta, request.args.get("agrupar", "DIA"), request.args.get("n") or 7)
        etiqueta = (lambda b: f"{b:%m/%Y}") if cap["agrupar"] == "MES" else (lambda b: f"{b:%d/%m}")
        graf_cap = charts.line([(f["grupo"], [(etiqueta(x["bucket"]), x["uso"] or 0) for x in f["serie"]]) for f in cap["filas"]
                                if f["capacidad"]], "%", title="Uso de capacidad por línea")
        graf_prog = charts.bars([(f["grupo"], f["cumplimiento"] or 0) for f in cap["filas"] if f["programado"]], "%",
                                title="Programado vs ejecutado (cumplimiento)", color=2)
        from ..services import prod_analytics as pa
        reco_cap, _ = pa.capacity_recommendations(d["desde"], d["hasta"])
        an = pa.analyze(d["desde"], d["hasta"])
        d["recomendaciones"] = list(d["recomendaciones"]) + reco_cap + [x for x in an["recomendaciones"]
                                                                        if x[1].startswith(("La línea", "Hay "))]
        d["merma_kg"], d["remanente_kg"], d["merma_pct"] = an["merma_kg"], an["remanente_kg"], an["merma_pct"]
        return render_template("informe_produccion.html", d=d, cap=cap, graf_cap=graf_cap, graf_prog=graf_prog,
                               charts=charts, items=disponibles(), actual=nombre)
    if nombre == "cumplimiento":
        from datetime import date, timedelta
        from ..services.shifts import compliance
        d = compliance(desde or date.today() - timedelta(days=29), hasta or date.today(), user=g.user)
        pp = d["produccion"] or []
        graf = {"kg_h": charts.bars([(p["nombre"][:22], p["kg_h"]) for p in pp if p["kg_h"]], "kg/h",
                                    title="Rendimiento por persona (kg/h)"),
                "ausencia": charts.bars([(p["nombre"][:22], p["ausencia_pct"]) for p in pp if p["ausencia_pct"]], "%",
                                        title="Tasa de ausencia (%)", color=1)}
        return render_template("informe_cumplimiento.html", d=d, graf=graf, items=disponibles(), actual=nombre)
    if nombre == "merma":
        from ..services import prod_analytics as pa
        d = pa.analyze(desde, hasta)
        reco_cap, cap = pa.capacity_recommendations(d["desde"], d["hasta"])
        d["recomendaciones"] = d["recomendaciones"] + reco_cap
        graf = {"lineas": charts.columns([x["linea"] for x in d["lineas"]],
                                         [("Merma kg", [x["merma"] for x in d["lineas"]]),
                                          ("Remanente MP kg", [x["remanente"] for x in d["lineas"]])], "kg",
                                         title="Merma y remanente por línea"),
                "serie": charts.line([("Merma kg", [(a, m) for a, m, _ in d["serie"]]),
                                      ("Remanente kg", [(a, r) for a, _, r in d["serie"]])], "kg",
                                     title="Merma y remanente por día"),
                "personas": charts.bars([(p["nombre"][:22], p["merma_pct"]) for p in d["personas"] if p["merma_pct"] is not None],
                                        "%", title="Merma por persona (%)", color=1),
                "rem_personas": charts.bars([(p["nombre"][:22], p["remanente"]) for p in d["personas"] if p["remanente"]],
                                            "kg", title="Remanente de MP por persona", color=2)}
        from ..services.core import param
        return render_template("informe_merma.html", d=d, graf=graf, items=disponibles(), actual=nombre,
                               umbral_merma=float(param("umbral_merma_pct", 5) or 5))
    fn = getattr(rsvc, nombre)
    d = fn(desde, hasta, solo_aprobadas=_area_produccion()) if nombre == "reempaque" else fn(desde, hasta)
    return render_template(f"informe_{nombre}.html", d=d, charts=charts, items=disponibles(), actual=nombre)


@bp.route("/comercial/cliente/<int:cid>")
@login_required
def customer_behavior(cid):
    """Detalle del comportamiento de un cliente: cada OV con su atención, retraso, rechazos, modificaciones y penalidad."""
    if "informes.comercial" not in g.perms and "vista.total" not in g.perms:
        abort(403)
    return _descarga(_customer_behavior(cid), "Comportamiento del cliente")


def _customer_behavior(cid):
    from ..services.customer_behavior import behavior
    desde, hasta = _periodo()
    d = behavior(desde, hasta, cid)
    if not d["clientes"]:
        from ..db import db
        from ..models import Customer
        c = db().get(Customer, cid) or abort(404)
        return render_template("informe_cliente.html", d=d, c=None, cliente=c, items=disponibles(), actual="comercial")
    return render_template("informe_cliente.html", d=d, c=d["clientes"][0], cliente=d["clientes"][0]["customer"],
                           items=disponibles(), actual="comercial")
