"""Presentación PowerPoint corporativa del informe de gestión comercial (V7): portada con logo, indicadores clave,
nivel de servicio con su fórmula, metas, desempeño por persona y recomendaciones."""
import io
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

AZUL = RGBColor(0x0B, 0x3D, 0x6E)
VERDE = RGBColor(0x1B, 0xAF, 0x7A)
ROJO = RGBColor(0xC6, 0x28, 0x28)
AMBAR = RGBColor(0xE0, 0x8A, 0x00)
GRIS = RGBColor(0x5F, 0x6B, 0x7A)
CLARO = RGBColor(0xF3, 0xF6, 0xFA)
LOGO = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "img", "logo_corporacion_productiva.png")


def _n(v, dec=0):
    if v is None:
        return "—"
    return f"{v:,.{dec}f}".replace(",", " ")


def _title(slide, texto, sub=""):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(10.5), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = texto
    p.font.size, p.font.bold, p.font.color.rgb = Pt(26), True, AZUL
    if sub:
        q = tb.text_frame.add_paragraph()
        q.text = sub
        q.font.size, q.font.color.rgb = Pt(12), GRIS
    if os.path.exists(LOGO):
        slide.shapes.add_picture(LOGO, Inches(11.3), Inches(0.3), height=Inches(0.6))
    bar = slide.shapes.add_shape(1, Inches(0.5), Inches(1.15), Inches(12.3), Emu(30000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = AZUL
    bar.line.fill.background()


def _tile(slide, x, y, valor, etiqueta, sub="", color=AZUL):
    sh = slide.shapes.add_shape(5, Inches(x), Inches(y), Inches(2.9), Inches(1.45))
    sh.fill.solid()
    sh.fill.fore_color.rgb = CLARO
    sh.line.color.rgb = color
    tf = sh.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = valor
    p.font.size, p.font.bold, p.font.color.rgb = Pt(24), True, color
    p.alignment = PP_ALIGN.CENTER
    for t, size in ((etiqueta, 12), (sub, 10)):
        if t:
            q = tf.add_paragraph()
            q.text = t
            q.font.size, q.font.color.rgb = Pt(size), GRIS
            q.alignment = PP_ALIGN.CENTER


def _table(slide, headers, rows, y=1.5, widths=None):
    rows = rows[:14]
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.5), Inches(y), Inches(12.3),
                                   Inches(0.35 * (len(rows) + 1)))
    t = shape.table
    for j, h in enumerate(headers):
        c = t.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = AZUL
        para = c.text_frame.paragraphs[0]
        para.font.size, para.font.bold, para.font.color.rgb = Pt(11), True, RGBColor(255, 255, 255)
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = str(v)
            c.text_frame.paragraphs[0].font.size = Pt(10)
    if widths:
        for j, w in enumerate(widths):
            t.columns[j].width = Inches(w)


def _sem(v, meta):
    if v is None:
        return GRIS
    return VERDE if v >= meta else (AMBAR if v >= meta - 10 else ROJO)


def gestion_pptx(d, titulo="Gestión comercial"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    periodo = f"Periodo {d['desde']:%d/%m/%Y} – {d['hasta']:%d/%m/%Y}"
    # portada
    s = prs.slides.add_slide(blank)
    if os.path.exists(LOGO):
        s.shapes.add_picture(LOGO, Inches(0.8), Inches(1.0), height=Inches(1.2))
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.5), Inches(2))
    p = tb.text_frame.paragraphs[0]
    p.text = f"Informe de {titulo.lower()}"
    p.font.size, p.font.bold, p.font.color.rgb = Pt(40), True, AZUL
    q = tb.text_frame.add_paragraph()
    q.text = periodo
    q.font.size, q.font.color.rgb = Pt(18), GRIS
    # indicadores clave
    s = prs.slides.add_slide(blank)
    _title(s, "Indicadores clave", periodo)
    i, a, pe, fa, nv = d["ingresadas"], d["atendidas"], d["pendientes"], d["facturacion"], d["nivel"]
    _tile(s, 0.5, 1.5, f"S/ {_n(i['soles'])}", "OV ingresadas", f"{i['n']} OV · {_n(i['kg'])} kg")
    _tile(s, 3.6, 1.5, f"S/ {_n(a['soles'])}", "Atendido", f"{a['n']} OV · {_n(a['kg'])} kg", VERDE)
    _tile(s, 6.7, 1.5, f"S/ {_n(pe['soles'])}", "Pendiente", f"{pe['n']} OV · {pe['vencidas']} vencidas",
          ROJO if pe["vencidas"] else AMBAR)
    _tile(s, 9.8, 1.5, f"S/ {_n(fa['soles'])}", "Facturación", fa["fuente"][:40])
    meta = 95
    for k, (key, lbl) in enumerate((("otif", "OTIF"), ("on_time", "On time"), ("in_full", "In full"), ("fill_unidades", "Fill rate"))):
        v = nv[key]
        _tile(s, 0.5 + 3.1 * k, 3.4, f"{_n(v, 1)} %" if v is not None else "—", lbl, d["formulas"].get(key, "")[:60],
              _sem(v, meta))
    tb = s.shapes.add_textbox(Inches(0.5), Inches(5.2), Inches(12.3), Inches(1.5))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = f"Base del nivel de servicio: {nv['n']} OV que salieron del almacén en el periodo."
    p.font.size, p.font.color.rgb = Pt(11), GRIS
    # metas
    if d["metas"]["filas"]:
        s = prs.slides.add_slide(blank)
        _title(s, f"Avance de metas {d['metas']['mes']}", f"Avance esperado a la fecha: {d['metas']['esperado']} %")
        _table(s, ["Alcance", "Nombre", "Meta S/", "Real S/", "Avance S/", "Meta kg", "Real kg", "Avance kg", "Estado"],
               [(m["alcance"].capitalize(), m["nombre"], _n(m["meta_soles"]), _n(m["soles"]), f"{_n(m['avance_soles'], 1)} %",
                 _n(m["meta_kg"]), _n(m["kg"]), f"{_n(m['avance_kg'], 1)} %" if m["meta_kg"] else "—",
                 "Atrasada" if m["atrasada"] else "En línea") for m in d["metas"]["filas"]])
    # desempeño por persona
    if d["personas"]:
        s = prs.slides.add_slide(blank)
        _title(s, "Desempeño por vendedor", "Venta, meta, nivel de servicio, modificaciones y embudo")
        _table(s, ["Vendedor", "OV", "S/", "kg", "Meta S/ %", "OTIF %", "OV modif. %", "Oport.", "Ganadas", "Penalidad S/"],
               [(p["nombre"], p["ovs"], _n(p["soles"]), _n(p["kg"]),
                 f"{_n((p['meta'] or {}).get('avance_soles'), 1)}" if p["meta"] else "—", _n(p["nivel"]["otif"], 1),
                 _n(p["pct_modificadas"], 1), p["oportunidades"], p["ganadas"], _n(p["penalidad_calc"])) for p in d["personas"]])
    # por canal
    s = prs.slides.add_slide(blank)
    _title(s, "Venta y servicio por canal", periodo)
    _table(s, ["Canal", "OV", "S/", "kg", "Ticket S/", "OTIF %", "Fill rate %"],
           [(n, r["ovs"], _n(r["soles"]), _n(r["kg"]), _n(r["ticket"]), _n(r["nivel"]["otif"], 1), _n(r["nivel"]["fill_unidades"], 1))
            for n, r in d["por"]["canal"]])
    # recomendaciones
    s = prs.slides.add_slide(blank)
    _title(s, "Recomendaciones", "Calculadas con reglas y umbrales configurables; requieren revisión del responsable")
    y = 1.5
    for nivel, texto in d["recomendaciones"][:7]:
        color = ROJO if nivel == "ALTA" else (AMBAR if nivel == "MEDIA" else VERDE)
        sh = s.shapes.add_shape(5, Inches(0.5), Inches(y), Inches(12.3), Inches(0.72))
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.fill.background()
        tf = sh.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = texto
        tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold = Pt(14), True
        tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        y += 0.8
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
