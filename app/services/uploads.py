"""Carga de archivos SAP/Excel con validación previa, corrección y autorización (IF §2, §6.4, §8.1, §19).

Flujo: subir -> validar (errores y advertencias por fila; diferencias) -> quien carga corrige/excluye filas ->
autoriza (confirmar). No se incorpora nada mientras haya errores sin corregir.

Formatos soportados (encabezados tal como aparecen en los archivos fuente):
- STOCK       : hoja STOCK_LOTE de 'PCP 2026 - J.xlsm' (exportación SAP de stock por lote).
- BOM         : 'LISTA DE MATERIALES 2026.xlsx' (Hoja1).
- OV          : hoja OVS de 'PCP 2026 - J.xlsm' (exportación SAP de órdenes de venta).
- LIQUIDACION : hoja 'Dat.Prod' de 'Recep. pt.xlsx' (órdenes de fabricación SAP).
- PRODUCTOS   : hoja DATA de 'PCP 2026 - J.xlsm' (maestro: código, descripción, PPB, MAQ, TIPO, CALIDAD).
- SALIDAS     : hoja 'Qy' de 'Rotacion de inventario - Almacenes.xlsx' (historial de ventas/salidas).
"""
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

import openpyxl
from sqlalchemy import select

from ..db import db
from ..models import (BOMLine, BOMVersion, Customer, Liquidation, Product, SalesHistory, SalesOrder, StockLot,
                      UploadBatch, UploadRow, User, Warehouse, now)
from . import alerts
from .core import BusinessError, audit, has_perm, param


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


SPECS = {
    "STOCK": {
        "nombre": "Stock por lote (SAP)", "perm": "stock.cargar", "confirm_perm": "stock.autorizar_carga",
        "hojas": ["STOCK_LOTE"],
        "cols": {"articulo": ["Número de artículo"], "descripcion": ["Descripción del artículo"], "lote": ["Lote"],
                 "almacen": ["Almacén"], "cantidad": ["Cantidad"], "fecha_venc": ["Fecha de vencimiento"],
                 "fecha_fab": ["Fecha de fabricación"], "fecha_admision": ["Fecha de admisión"]},
        "req": ["articulo", "lote", "almacen", "cantidad"],
    },
    "BOM": {
        "nombre": "Lista de materiales (Excel)", "perm": "produccion.bom_cargar", "confirm_perm": "produccion.bom_cargar",
        "hojas": ["Hoja1"],
        "cols": {"padre": ["Articulo Padre"], "padre_desc": ["Descripcion Padre"], "tipo_bom": ["Tipo BOM"],
                 "cantidad_bom": ["Cantidad BOM"], "almacen_pt": ["Almacen PT"], "fecha_creacion": ["Fecha Creacion"],
                 "fecha_actualizacion": ["Fecha Actualizacion"], "orden": ["Orden"], "componente": ["Articulo Componente"],
                 "componente_desc": ["Descripcion Componente"], "cantidad": ["Cantidad Componente"],
                 "almacen_comp": ["Almacen Componente"], "um": ["UM"], "metodo": ["Metodo Emision"]},
        "req": ["padre", "padre_desc", "cantidad_bom", "componente", "componente_desc", "cantidad", "almacen_comp"],
    },
    "OV": {
        "nombre": "Órdenes de venta (exportación SAP)", "perm": "ov.importar", "confirm_perm": "ov.importar",
        "hojas": ["OVS"],
        "cols": {"numero": ["Número de documento"], "estado": ["Estado"], "fecha_emision": ["Fecha de Emisión"],
                 "fecha_compromiso": ["Fecha de Vencimiento"], "vendedor": ["Vendedor"], "ruc": ["Ruc"],
                 "cliente": ["Nombre de cliente"], "grupo": ["Grupo de cliente"], "direccion": ["Dirección"],
                 "distrito": ["Distrito"], "provincia": ["Provincia"], "producto": ["Código de Producto"],
                 "descripcion": ["Descripción de Producto"], "solicitado": ["Solicitado"], "entregado": ["Entregado"],
                 "um": ["Unidad de medida"], "peso": ["Peso"], "tipo": ["T"], "comentarios": ["Comentarios"],
                 "precio": ["P. Unitario SOLES"]},
        "req": ["numero", "fecha_compromiso", "ruc", "cliente", "producto", "solicitado"],
    },
    "LIQUIDACION": {
        "nombre": "Liquidación de producción (SAP/Excel)", "perm": "produccion.liquidar", "confirm_perm": "produccion.liquidar",
        "hojas": ["Dat.Prod"],
        "cols": {"numero": ["Número de documento"], "tipo_op": ["Tipo de Operación"], "almacen": ["Almacén"],
                 "fecha": ["Fecha de creación"], "producto": ["Código de producto"], "descripcion": ["Descripción de producto"],
                 "planificada": ["Cantidad Planificada"], "completada": ["Cantidad Completada"], "um": ["Unidad de medida"],
                 "valorizado": ["Valorizado"], "estado": ["Estado"], "realizado_por": ["Realizado por"],
                 "comentarios": ["Comentarios"], "lote": ["Lote"], "fecha_fab": ["Fecha de admisión"],
                 "fecha_venc": ["Fecha de vencimiento"]},
        "req": ["numero", "producto", "completada", "lote"],
    },
    "PRODUCTOS": {
        "nombre": "Maestro de artículos (tabla DATA)", "perm": "config.editar", "confirm_perm": "config.editar",
        "hojas": ["DATA"],
        "cols": {"codigo": ["Código de Producto"], "descripcion": ["Descripción de Producto"], "ppb": ["PPB"],
                 "maq": ["MAQ"], "tipo": ["TIPO"], "calidad": ["CALIDAD"]},
        "req": ["codigo", "descripcion"],
    },
    "SALIDAS": {
        "nombre": "Historial de salidas/ventas", "perm": "stock.cargar", "confirm_perm": "stock.autorizar_carga",
        "hojas": ["Qy"],
        "cols": {"producto": ["Código de Producto"], "cantidad": ["Cantidad"], "fecha": ["Fecha de Emisión"],
                 "documento": ["Número de documento"]},
        "req": ["producto", "cantidad", "fecha"],
    },
}


# ------------------------------- lectura -----------------------------------
def _cell(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str):
        v = v.strip()
        if v.startswith("="):
            return None  # fórmula sin valor calculado
        return v
    return v


def read_rows(path, tipo, hoja=None):
    spec = SPECS[tipo]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    candidates = [hoja] if hoja else [h for h in spec["hojas"] if h in wb.sheetnames] + wb.sheetnames
    alias = {norm(a): key for key, names in spec["cols"].items() for a in names}
    for sheet in dict.fromkeys(candidates):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows_iter = ws.iter_rows(values_only=True)
        header_idx = None
        buffer = []
        for i, row in enumerate(rows_iter):
            buffer.append(row)
            mapping = {}
            for j, v in enumerate(row):
                k = alias.get(norm(v))
                if k and k not in mapping:
                    mapping[k] = j
            if all(r in mapping for r in spec["req"]):
                header_idx = i
                break
            if i > 15:
                break
        if header_idx is None:
            continue
        out = []
        n = header_idx + 1
        for row in rows_iter:
            n += 1
            if row is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                continue
            datos = {k: _cell(row[j]) if j < len(row) else None for k, j in mapping.items()}
            out.append((n, datos))
        wb.close()
        return sheet, out, sorted(set(spec["cols"]) - set(mapping))
    wb.close()
    raise BusinessError(f"No se encontró una hoja con los encabezados requeridos para '{spec['nombre']}': "
                        + ", ".join(spec["cols"][r][0] for r in spec["req"]))


# ------------------------------- conversiones -----------------------------
def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        raise ValueError(f"'{v}' no es numérico")


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # número de serie de fecha de Excel
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.strptime(s[:10] if fmt == "%Y-%m-%d" else s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"fecha '{v}' no reconocida")


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


PRESENT_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*(KG|K|G|GR|GRS)\b", re.I)


def parse_presentation(desc):
    """'ARVEJA VERDE PARTIDA * 500 G - CANTACLARO' -> ('500 G', 0.5)."""
    m = None
    for m in PRESENT_RE.finditer(desc or ""):
        pass
    if not m:
        return "", None
    q = float(m.group(1).replace(",", "."))
    unit = m.group(2).upper()
    kg = q if unit in ("KG", "K") else q / 1000.0
    return f"{m.group(1)} {unit}", kg


def product_base(desc):
    d = (desc or "").upper()
    base = d.split("*")[0] if "*" in d else d.split(" - ")[0]
    return re.sub(r"\s+", " ", base).strip()


# ------------------------------- validación -------------------------------
def create_batch(user, tipo, path, filename, hoja=None):
    spec = SPECS.get(tipo)
    if spec is None:
        raise BusinessError("Tipo de carga no válido.")
    if not has_perm(user, spec["perm"]):
        raise BusinessError(f"Su rol no puede cargar '{spec['nombre']}'.")
    sheet, rows, faltantes = read_rows(path, tipo, hoja)
    b = UploadBatch(tipo=tipo, archivo=filename, hoja=sheet, user_id=user.id, estado="LISTO")
    if faltantes:
        b.errores_generales = [f"Columnas opcionales no encontradas: {', '.join(faltantes)}"]
    for n, datos in rows:
        b.rows.append(UploadRow(fila=n, datos=datos))
    db().add(b)
    db().flush()
    validate_batch(b)
    audit(user, "CARGA_SUBIDA", "UploadBatch", b.id, {"tipo": tipo, "archivo": filename, "hoja": sheet, "filas": len(rows)})
    return b


def validate_batch(b):
    fn = {"STOCK": _v_stock, "BOM": _v_bom, "OV": _v_ov, "LIQUIDACION": _v_liq, "PRODUCTOS": _v_prod,
          "SALIDAS": _v_salidas}[b.tipo]
    for r in b.rows:
        r.errores, r.advertencias = [], []
    fn(b)
    activos = [r for r in b.rows if not r.excluida]
    n_err = sum(1 for r in activos if r.errores)
    n_adv = sum(1 for r in activos if r.advertencias)
    res = dict(b.resumen or {})
    res.update({"filas": len(b.rows), "activas": len(activos), "excluidas": len(b.rows) - len(activos),
                "con_error": n_err, "con_advertencia": n_adv, "editadas": sum(1 for r in b.rows if r.editada)})
    b.resumen = res
    if b.estado not in ("CONFIRMADO", "DESCARTADO"):
        b.estado = "VALIDADO_CON_ERRORES" if n_err else "LISTO"
    return b


def _dup_check(rows, keyfn, label):
    c = Counter()
    keys = {}
    for r in rows:
        try:
            k = keyfn(r.datos)
        except Exception:  # noqa: BLE001
            continue
        keys[r.id] = k
        c[k] += 1
    for r in rows:
        k = keys.get(r.id)
        if k is not None and c[k] > 1:
            r.errores = r.errores + [f"Duplicado ({label}: {' / '.join(map(str, k))})"]


def _v_stock(b):
    s = db()
    activos = [r for r in b.rows if not r.excluida]
    whs = {w.code for w in s.execute(select(Warehouse)).scalars()}
    prods = set(s.execute(select(Product.code)).scalars())
    for r in activos:
        d = r.datos
        e, w = [], []
        if not _txt(d.get("articulo")):
            e.append("Número de artículo vacío")
        if not _txt(d.get("almacen")):
            e.append("Almacén vacío")
        elif _txt(d.get("almacen")) not in whs:
            w.append(f"Almacén {_txt(d.get('almacen'))} no configurado: se creará con uso 'OTRO' (no reservable)")
        if not _txt(d.get("lote")):
            w.append("Lote vacío: se registrará como 'SIN LOTE'")
        elif not re.match(r"^[A-Za-z0-9/\-]+$", _txt(d.get("lote"))):
            w.append(f"Lote con caracteres inusuales: '{_txt(d.get('lote'))}' (verifique en SAP)")
        try:
            q = _num(d.get("cantidad"))
            if q is None:
                e.append("Cantidad vacía")
            elif q < 0:
                e.append("Cantidad negativa")
        except ValueError as ex:
            e.append(f"Cantidad: {ex}")
        for k in ("fecha_venc", "fecha_fab", "fecha_admision"):
            try:
                _date(d.get(k))
            except ValueError as ex:
                e.append(f"{k}: {ex}")
        if _txt(d.get("articulo")) and _txt(d.get("articulo")) not in prods:
            w.append("Artículo no existe en el catálogo: se creará con la descripción del archivo")
        r.errores, r.advertencias = e, w
    _dup_check(activos, lambda d: (_txt(d["articulo"]), _txt(d.get("lote")) or "SIN LOTE", _txt(d["almacen"])),
               "artículo/lote/almacén")
    # diferencias contra el stock del sistema (en los almacenes incluidos en el archivo)
    file_whs = {_txt(r.datos.get("almacen")) for r in activos}
    actual = defaultdict(float)
    reservado = defaultdict(float)
    from .stock import reserved_qty
    for lot in s.execute(select(StockLot).where(StockLot.warehouse_code.in_(file_whs))).scalars():
        k = (lot.product_code, lot.lote, lot.warehouse_code)
        actual[k] += lot.cantidad
        reservado[k] += reserved_qty(lot)
    nuevo = defaultdict(float)
    for r in activos:
        if r.errores:
            continue
        d = r.datos
        nuevo[(_txt(d["articulo"]), _txt(d.get("lote")) or "SIN LOTE", _txt(d["almacen"]))] += _num(d["cantidad"]) or 0
    nuevos = [k for k in nuevo if k not in actual]
    cambiados = [(k, actual[k], nuevo[k]) for k in nuevo if k in actual and abs(actual[k] - nuevo[k]) > 1e-9]
    ausentes = [k for k in actual if k not in nuevo and actual[k] > 0]
    afecta = [(k, reservado[k], nuevo.get(k, 0)) for k in reservado if reservado[k] > nuevo.get(k, 0) + 1e-9]
    b.resumen = {"almacenes_en_archivo": sorted(file_whs), "lotes_nuevos": len(nuevos), "lotes_modificados": len(cambiados),
                 "lotes_ausentes_quedaran_en_0": len(ausentes), "reservas_afectadas": len(afecta),
                 "detalle_modificados": [{"clave": list(k), "antes": a, "despues": n} for k, a, n in cambiados[:200]],
                 "detalle_ausentes": [list(k) for k in ausentes[:200]],
                 "detalle_reservas_afectadas": [{"clave": list(k), "reservado": rv, "nuevo": n} for k, rv, n in afecta[:200]]}


def _v_bom(b):
    s = db()
    prods = set(s.execute(select(Product.code)).scalars())
    activos = [r for r in b.rows if not r.excluida]
    for r in activos:
        d = r.datos
        e, w = [], []
        for k, lbl in (("padre", "Artículo padre"), ("padre_desc", "Descripción padre"), ("componente", "Artículo componente"),
                       ("componente_desc", "Descripción componente"), ("almacen_comp", "Almacén componente")):
            if not _txt(d.get(k)):
                e.append(f"{lbl} vacío")
        for k, lbl in (("cantidad_bom", "Cantidad BOM"), ("cantidad", "Cantidad componente")):
            try:
                v = _num(d.get(k))
                if v is None or v <= 0:
                    e.append(f"{lbl} debe ser > 0")
            except ValueError as ex:
                e.append(f"{lbl}: {ex}")
        if _txt(d.get("padre")) == _txt(d.get("componente")) and _txt(d.get("padre")):
            e.append("El componente es igual al artículo padre")
        if _txt(d.get("orden")):
            try:
                int(float(d.get("orden")))
            except (TypeError, ValueError):
                e.append("Orden no numérico")
        if _txt(d.get("padre")) and _txt(d.get("padre")) not in prods:
            w.append("Producto padre no está en el catálogo (se creará)")
        if _txt(d.get("componente")) and _txt(d.get("componente")) not in prods:
            w.append("Componente no está en el catálogo (se creará)")
        pres, _ = parse_presentation(_txt(d.get("padre_desc")))
        if not pres:
            w.append("No se identificó la presentación en la descripción del padre")
        r.errores, r.advertencias = e, w
    _dup_check(activos, lambda d: (_txt(d["padre"]), _txt(d["componente"])), "padre/componente")
    padres = {_txt(r.datos.get("padre")) for r in activos}
    prev = s.execute(select(BOMVersion).where(BOMVersion.activa.is_(True))).scalars().first()
    diff = {}
    if prev:
        old = {(l.padre, l.componente): l.cantidad for l in prev.lines}
        new = {}
        for r in activos:
            if not r.errores:
                new[(_txt(r.datos["padre"]), _txt(r.datos["componente"]))] = _num(r.datos["cantidad"])
        diff = {"version_anterior": prev.version,
                "componentes_agregados": len([k for k in new if k not in old]),
                "componentes_eliminados": len([k for k in old if k not in new]),
                "cantidades_modificadas": len([k for k in new if k in old and abs(old[k] - new[k]) > 1e-12])}
    b.resumen = {"productos_padre": len(padres), "comparacion": diff}


def _v_ov(b):
    s = db()
    prods = set(s.execute(select(Product.code)).scalars())
    activos = [r for r in b.rows if not r.excluida]
    for r in activos:
        d = r.datos
        e, w = [], []
        if not _txt(d.get("numero")):
            e.append("Número de documento vacío")
        if not _txt(d.get("ruc")):
            e.append("RUC vacío")
        if not _txt(d.get("producto")):
            e.append("Código de producto vacío")
        try:
            if _date(d.get("fecha_compromiso")) is None:
                e.append("Fecha de vencimiento (comprometida) vacía")
        except ValueError as ex:
            e.append(str(ex))
        try:
            q = _num(d.get("solicitado"))
            if q is None or q <= 0:
                e.append("Cantidad solicitada debe ser > 0")
        except ValueError as ex:
            e.append(f"Solicitado: {ex}")
        try:
            _num(d.get("precio"))
            _num(d.get("entregado"))
        except ValueError as ex:
            e.append(str(ex))
        if (_txt(d.get("estado")) or "").lower() == "cerrado":
            w.append("OV cerrada en SAP: no se importará")
        if _txt(d.get("producto")) and _txt(d.get("producto")) not in prods:
            w.append("Producto no está en el catálogo (se creará con la descripción)")
        existing = s.execute(select(SalesOrder).where(SalesOrder.numero == _txt(d.get("numero")))).scalars().first()
        if existing is not None:
            w.append(f"La OV ya existe (estado {existing.estado}): se aplicarán cambios como modificación")
        r.errores, r.advertencias = e, w
    rep = Counter((_txt(r.datos.get("numero")), _txt(r.datos.get("producto"))) for r in activos)
    for r in activos:
        if rep[(_txt(r.datos.get("numero")), _txt(r.datos.get("producto")))] > 1:
            r.advertencias = r.advertencias + ["Producto repetido en la misma OV: las cantidades se consolidarán en una línea"]
    b.resumen = {"ovs": len({_txt(r.datos.get('numero')) for r in activos})}


def _v_liq(b):
    s = db()
    prods = set(s.execute(select(Product.code)).scalars())
    existentes = set(s.execute(select(Liquidation.numero_documento, Liquidation.lote)).all())
    por_doc = Counter(_txt(r.datos.get("numero")) for r in b.rows if not r.excluida)
    activos = [r for r in b.rows if not r.excluida]
    for r in activos:
        d = r.datos
        e, w = [], []
        num = _txt(d.get("numero"))
        if not num:
            e.append("Número de documento vacío")
        elif (num, _txt(d.get("lote"))) in existentes:
            e.append("Liquidación ya registrada (duplicado documento/lote)")
        if num and por_doc[num] > 1:
            w.append("Documento con varios lotes: verificar si la cantidad completada se repite por lote (clave documento+lote pendiente de validar)")
        if _txt(d.get("producto")) not in prods:
            e.append("Producto no existe en el catálogo")
        if not _txt(d.get("lote")):
            e.append("Lote vacío")
        elif not re.match(r"^[A-Za-z0-9/\-]+$", _txt(d.get("lote"))):
            w.append(f"Lote con caracteres inusuales: '{_txt(d.get('lote'))}' (verifique)")
        try:
            q = _num(d.get("completada"))
            if q is None or q < 0:
                e.append("Cantidad completada inválida")
            pl = _num(d.get("planificada"))
            if pl is not None and q is not None and abs(pl - q) > 1e-9:
                w.append(f"Completada ({q:g}) distinta de planificada ({pl:g})")
        except ValueError as ex:
            e.append(str(ex))
        for k in ("fecha", "fecha_fab", "fecha_venc"):
            try:
                _date(d.get(k))
            except ValueError as ex:
                e.append(f"{k}: {ex}")
        r.errores, r.advertencias = e, w
    _dup_check(activos, lambda d: (_txt(d["numero"]), _txt(d["lote"])), "documento/lote")
    b.resumen = {}


def _v_prod(b):
    activos = [r for r in b.rows if not r.excluida]
    for r in activos:
        d = r.datos
        e, w = [], []
        if not _txt(d.get("codigo")):
            e.append("Código vacío")
        if not _txt(d.get("descripcion")):
            e.append("Descripción vacía")
        try:
            _num(d.get("ppb"))
        except ValueError as ex:
            e.append(f"PPB: {ex}")
        if _txt(d.get("tipo")) and _txt(d.get("tipo")) not in ("PT", "GR", "MP"):
            w.append(f"TIPO '{_txt(d.get('tipo'))}' no reconocido")
        r.errores, r.advertencias = e, w
    _dup_check(activos, lambda d: (_txt(d["codigo"]),), "código")
    b.resumen = {}


def _v_salidas(b):
    s = db()
    ya = _sales_keys(s)
    activos = [r for r in b.rows if not r.excluida]
    for r in activos:
        d = r.datos
        e, w = [], []
        if not _txt(d.get("producto")):
            w.append("Línea sin código de producto (servicio/nota): se omitirá")
        try:
            if _num(d.get("cantidad")) is None:
                e.append("Cantidad vacía")
        except ValueError as ex:
            e.append(str(ex))
        try:
            if _date(d.get("fecha")) is None:
                e.append("Fecha vacía")
        except ValueError as ex:
            e.append(str(ex))
        if _txt(d.get("documento")) and _sales_key(d) in ya:
            w.append("Fila ya cargada en una carga anterior (documento/producto/fecha/cantidad): se omitirá")
        r.errores, r.advertencias = e, w
    b.resumen = {}


# ------------------------------- edición ----------------------------------
def edit_row(b, row, user, nuevos):
    if b.estado in ("CONFIRMADO", "DESCARTADO"):
        raise BusinessError("La carga ya fue cerrada.")
    if b.user_id != user.id and not has_perm(user, SPECS[b.tipo]["confirm_perm"]):
        raise BusinessError("Solo quien carga (o quien autoriza) puede corregir.")
    antes = dict(row.datos)
    datos = dict(row.datos)
    for k, v in nuevos.items():
        if k in SPECS[b.tipo]["cols"]:
            datos[k] = v.strip() if isinstance(v, str) else v
    row.datos = datos
    row.editada = True
    validate_batch(b)
    audit(user, "CARGA_FILA_CORREGIDA", "UploadBatch", b.id, {"fila": row.fila, "antes": antes, "despues": datos})


def exclude_row(b, row, user, motivo, excluir=True):
    if b.estado in ("CONFIRMADO", "DESCARTADO"):
        raise BusinessError("La carga ya fue cerrada.")
    if excluir and not motivo.strip():
        raise BusinessError("Indique el motivo de exclusión.")
    row.excluida = excluir
    row.motivo_exclusion = motivo if excluir else ""
    validate_batch(b)
    audit(user, "CARGA_FILA_EXCLUIDA" if excluir else "CARGA_FILA_INCLUIDA", "UploadBatch", b.id,
          {"fila": row.fila, "motivo": motivo})


def discard(b, user):
    if b.estado == "CONFIRMADO":
        raise BusinessError("La carga ya fue confirmada.")
    b.estado = "DESCARTADO"
    audit(user, "CARGA_DESCARTADA", "UploadBatch", b.id)


# ------------------------------- confirmación -----------------------------
def confirm(b, user, opciones=None):
    spec = SPECS[b.tipo]
    if not has_perm(user, spec["confirm_perm"]):
        raise BusinessError("Su rol no puede autorizar esta carga.")
    if b.estado == "CONFIRMADO":
        raise BusinessError("La carga ya fue confirmada.")
    if b.estado == "DESCARTADO":
        raise BusinessError("La carga fue descartada.")
    validate_batch(b)
    if b.estado == "VALIDADO_CON_ERRORES":
        raise BusinessError("Corrija o excluya las filas con error antes de autorizar la carga.")
    opciones = opciones or {}
    b.opciones = opciones
    result = {"STOCK": _c_stock, "BOM": _c_bom, "OV": _c_ov, "LIQUIDACION": _c_liq, "PRODUCTOS": _c_prod,
              "SALIDAS": _c_salidas}[b.tipo](b, user, opciones)
    b.estado = "CONFIRMADO"
    b.confirmado_ts = now()
    b.confirmado_por = user.id
    res = dict(b.resumen or {})
    res["resultado"] = result
    b.resumen = res
    audit(user, "CARGA_CONFIRMADA", "UploadBatch", b.id, {"tipo": b.tipo, "resultado": result})
    if b.tipo == "STOCK":
        corr = res.get("editadas", 0) + res.get("excluidas", 0)
        if corr or res.get("lotes_modificados") or res.get("reservas_afectadas") or res.get("lotes_ausentes_quedaran_en_0") \
                or (res.get("resultado") or {}).get("advertencias"):
            alerts.emit("STOCK_CARGA_DIFERENCIAS", f"Carga de stock #{b.id} confirmada con diferencias/correcciones",
                        f"Archivo {b.archivo}. Filas corregidas {res.get('editadas', 0)}, excluidas {res.get('excluidas', 0)}, "
                        f"con advertencia {res.get('con_advertencia', 0)}; lotes modificados {res.get('lotes_modificados', 0)}, "
                        f"ausentes {res.get('lotes_ausentes_quedaran_en_0', 0)}, reservas afectadas {res.get('reservas_afectadas', 0)}.",
                        "UploadBatch", b.id, creado_por=user)
    return result


def _ensure_product(code, desc, tipo=None, um=None, peso=None):
    s = db()
    p = s.get(Product, code)
    if p is None:
        pres, kg = parse_presentation(desc)
        p = Product(code=code, descripcion=desc or code, tipo=tipo or "", presentacion=pres,
                    peso_kg=peso if peso is not None else kg, producto_base=product_base(desc), um=um or "UNIDAD",
                    origen="CARGA")
        s.add(p)
        s.flush()
    return p


def _ensure_wh(code):
    s = db()
    w = s.get(Warehouse, code)
    if w is None:
        w = Warehouse(code=code, nombre=code, uso="OTRO", prioridad=99, reservable=False,
                      observacion="Creado desde carga SAP; uso pendiente de definir.")
        s.add(w)
        s.flush()
    return w


def _c_stock(b, user, opciones):
    from .stock import _release, evaluate_order, reevaluate_pending, reserved_qty, ORDEN_ESTADOS_CERRADOS
    s = db()
    primera = s.execute(select(UploadBatch).where(UploadBatch.tipo == "STOCK", UploadBatch.estado == "CONFIRMADO")).first() is None
    estado_nuevo = "DISPONIBLE" if primera else (param("estado_lote_nuevo_carga", "DISPONIBLE") or "DISPONIBLE")
    verif = bool(param("carga_autorizada_verifica", True))
    sap = defaultdict(float)
    meta = {}
    for r in b.rows:
        if r.excluida:
            continue
        d = r.datos
        k = (_txt(d["articulo"]), _txt(d.get("lote")) or "SIN LOTE", _txt(d["almacen"]))
        sap[k] += _num(d["cantidad"]) or 0
        meta[k] = d
    file_whs = {k[2] for k in sap}
    for code in file_whs:
        _ensure_wh(code)
    for k, d in meta.items():
        _ensure_product(k[0], _txt(d.get("descripcion")))
    segs = defaultdict(list)
    for lot in s.execute(select(StockLot).where(StockLot.warehouse_code.in_(file_whs))).scalars():
        segs[(lot.product_code, lot.lote, lot.warehouse_code)].append(lot)
    afectadas = set()
    nuevos = actualizados = puestos_cero = 0
    advertencias = []
    for k in set(sap) | set(segs):
        total = sap.get(k, 0.0)
        lots = segs.get(k, [])
        otros = sum(l.cantidad for l in lots if l.estado != "DISPONIBLE")
        disp = next((l for l in lots if l.estado == "DISPONIBLE"), None)
        objetivo = total - otros
        if objetivo < -1e-9:
            advertencias.append(f"{k}: SAP {total:g} < segmentos en otros estados {otros:g}")
            objetivo = 0
        d = meta.get(k, {})
        if disp is None and not lots and total > 0:
            lot = StockLot(product_code=k[0], lote=k[1], warehouse_code=k[2], cantidad=total, estado=estado_nuevo,
                           fecha_venc=_date(d.get("fecha_venc")), fecha_fab=_date(d.get("fecha_fab")),
                           fecha_admision=_date(d.get("fecha_admision")), verificado=verif,
                           fecha_verificacion=now() if verif else None, origen="CARGA_SAP", upload_id=b.id)
            s.add(lot)
            nuevos += 1
            continue
        if disp is None:
            if objetivo > 1e-9:
                disp = StockLot(product_code=k[0], lote=k[1], warehouse_code=k[2], cantidad=0, estado="DISPONIBLE",
                                fecha_venc=_date(d.get("fecha_venc")) or (lots[0].fecha_venc if lots else None),
                                fecha_fab=_date(d.get("fecha_fab")), verificado=verif, origen="CARGA_SAP", upload_id=b.id)
                s.add(disp)
            else:
                continue
        if abs(disp.cantidad - objetivo) > 1e-9:
            actualizados += 1
            if objetivo <= 1e-9:
                puestos_cero += 1
        disp.cantidad = objetivo
        disp.cantidad_verificada = None  # la carga autorizada reemplaza la verificación previa
        disp.verificado = verif or disp.verificado
        if verif:
            disp.fecha_verificacion = now()
        if d:
            disp.fecha_venc = _date(d.get("fecha_venc")) or disp.fecha_venc
            disp.fecha_fab = _date(d.get("fecha_fab")) or disp.fecha_fab
            disp.fecha_admision = _date(d.get("fecha_admision")) or disp.fecha_admision
        disp.upload_id = b.id
        s.flush()
        exceso = reserved_qty(disp) - disp.cantidad
        if exceso > 1e-9:
            for res in sorted([r for r in disp.reservations if r.estado == "ACTIVA"], key=lambda r: r.id, reverse=True):
                if exceso <= 1e-9:
                    break
                afectadas.add(res.line.order)
                if res.cantidad <= exceso + 1e-9:
                    exceso -= res.cantidad
                    _release(res, f"Carga de stock #{b.id} redujo el lote")
                else:
                    res.cantidad -= exceso
                    exceso = 0
    s.flush()
    for o in afectadas:
        alerts.emit("RESERVA_AFECTADA", f"Reserva reducida en OV {o.numero} por carga de stock #{b.id}",
                    "El stock cargado es menor a lo reservado; se reevalúa la OV.", "SalesOrder", o.id, creado_por=user)
        if o.estado not in ORDEN_ESTADOS_CERRADOS:
            evaluate_order(o, user, "REEVALUACION")
    reev = reevaluate_pending(user) if param("reevaluar_tras_carga", True) else 0
    return {"lotes_nuevos": nuevos, "lotes_actualizados": actualizados, "puestos_en_cero": puestos_cero,
            "ov_afectadas": len(afectadas), "ov_reevaluadas": reev, "advertencias": advertencias[:50],
            "base_inicial": primera}


def _c_bom(b, user, opciones):
    s = db()
    incl = opciones.get("incluye_merma", "NO_VERIFICADO")
    if incl not in ("SI", "NO", "NO_VERIFICADO"):
        raise BusinessError("Indique si el Excel ya incluye merma (SI / NO / NO_VERIFICADO).")
    prev = s.execute(select(BOMVersion).where(BOMVersion.activa.is_(True))).scalars().first()
    last = s.execute(select(BOMVersion).order_by(BOMVersion.version.desc())).scalars().first()
    v = BOMVersion(version=(last.version + 1) if last else 1, archivo=b.archivo, upload_id=b.id, user_id=user.id,
                   incluye_merma=incl, cambios=(b.resumen or {}).get("comparacion", {}))
    for r in b.rows:
        if r.excluida:
            continue
        d = r.datos
        pc = _txt(d["padre"])
        p = _ensure_product(pc, _txt(d["padre_desc"]))
        pres, kg = parse_presentation(_txt(d["padre_desc"]))
        if not p.presentacion and pres:
            p.presentacion = pres
        if p.peso_kg is None and kg:
            p.peso_kg = kg
        _ensure_product(_txt(d["componente"]), _txt(d["componente_desc"]))
        orden = d.get("orden")
        v.lines.append(BOMLine(padre=pc, padre_desc=_txt(d["padre_desc"]), tipo_bom=_txt(d.get("tipo_bom")) or "P",
                               cantidad_bom=_num(d["cantidad_bom"]) or 1, almacen_pt=_txt(d.get("almacen_pt")),
                               orden=int(float(orden)) if _txt(orden) else None, componente=_txt(d["componente"]),
                               componente_desc=_txt(d["componente_desc"]), cantidad=_num(d["cantidad"]),
                               almacen_comp=_txt(d["almacen_comp"]), um=_txt(d.get("um")), metodo_emision=_txt(d.get("metodo"))))
    if prev:
        prev.activa = False
    s.add(v)
    s.flush()
    return {"version": v.version, "lineas": len(v.lines), "incluye_merma": incl,
            "version_anterior": prev.version if prev else None}


def _c_ov(b, user, opciones):
    from .orders import _apply_changes, create_order
    s = db()
    grupos = defaultdict(list)
    for r in b.rows:
        if r.excluida or (_txt(r.datos.get("estado")) or "").lower() == "cerrado":
            continue
        grupos[_txt(r.datos["numero"])].append(r.datos)
    creadas = modificadas = sin_cambio = 0
    errores = []
    for numero, items in grupos.items():
        d0 = items[0]
        c = s.execute(select(Customer).where(Customer.ruc == _txt(d0["ruc"]))).scalars().first()
        if c is None:
            c = Customer(ruc=_txt(d0["ruc"]), nombre=_txt(d0["cliente"]), grupo=_txt(d0.get("grupo")),
                         direccion=_txt(d0.get("direccion")), distrito=_txt(d0.get("distrito")),
                         provincia=_txt(d0.get("provincia")))
            s.add(c)
            s.flush()
        cons = {}
        for d in items:
            um = "KILOGRAMO" if _txt(d.get("um")).upper().startswith("KILO") else "UNIDAD"
            tipo = _txt(d.get("tipo"))
            _ensure_product(_txt(d["producto"]), _txt(d.get("descripcion")), tipo if tipo in ("PT", "GR", "MP") else None, um,
                            _num(d.get("peso")))
            ln = cons.setdefault(_txt(d["producto"]), {"product_code": _txt(d["producto"]), "cantidad": 0.0,
                                                        "precio_unitario": _num(d.get("precio")) or 0})
            ln["cantidad"] += (_num(d["solicitado"]) or 0) - (_num(d.get("entregado")) or 0)
        lines = [l for l in cons.values() if l["cantidad"] > 0]
        if not lines:
            continue
        vend = s.execute(select(User).where(User.nombre == _txt(d0.get("vendedor")))).scalars().first()
        o = s.execute(select(SalesOrder).where(SalesOrder.numero == numero)).scalars().first()
        try:
            if o is None:
                o = create_order(user, numero, c.id, _date(d0["fecha_compromiso"]), lines, comentarios=_txt(d0.get("comentarios")),
                                 vendedor_id=vend.id if vend else None, fecha_emision=_date(d0.get("fecha_emision")),
                                 origen="CARGA_SAP")
                if not vend:
                    o.vendedor_nombre = _txt(d0.get("vendedor"))
                creadas += 1
            else:
                actual = {l.product_code: l.cantidad for l in o.lines if l.activa}
                nuevo = {l["product_code"]: l["cantidad"] for l in lines}
                if actual == nuevo and o.fecha_compromiso == _date(d0["fecha_compromiso"]):
                    sin_cambio += 1
                    continue
                _apply_changes(o, user, {"fecha_compromiso": _date(d0["fecha_compromiso"]), "lines": lines})
                modificadas += 1
        except BusinessError as ex:
            errores.append(f"OV {numero}: {ex}")
    return {"ov_creadas": creadas, "ov_modificadas": modificadas, "sin_cambio": sin_cambio, "errores": errores}


def _c_liq(b, user, opciones):
    from .production import register_liquidation
    n = 0
    for r in b.rows:
        if r.excluida:
            continue
        d = r.datos
        register_liquidation(user, _txt(d["numero"]), _txt(d["producto"]), _num(d["completada"]), _txt(d["lote"]),
                             fecha_fab=_date(d.get("fecha_fab")), fecha_venc=_date(d.get("fecha_venc")),
                             cantidad_planificada=_num(d.get("planificada")), origen=opciones.get("origen", "SAP"),
                             almacen=_txt(d.get("almacen")), valorizado=_num(d.get("valorizado")),
                             realizado_por=_txt(d.get("realizado_por")), fecha=_date(d.get("fecha")))
        n += 1
    return {"liquidaciones": n}


def _c_prod(b, user, opciones):
    s = db()
    n_new = n_upd = 0
    for r in b.rows:
        if r.excluida:
            continue
        d = r.datos
        code = _txt(d["codigo"])
        p = s.get(Product, code)
        pres, kg = parse_presentation(_txt(d["descripcion"]))
        if p is None:
            p = Product(code=code, descripcion=_txt(d["descripcion"]), origen="CARGA_DATA")
            s.add(p)
            n_new += 1
        else:
            n_upd += 1
        p.descripcion = _txt(d["descripcion"])
        p.peso_kg = _num(d.get("ppb")) if _num(d.get("ppb")) is not None else (p.peso_kg or kg)
        p.maquina = _txt(d.get("maq")) or p.maquina
        p.tipo = _txt(d.get("tipo")) or p.tipo
        p.calidad = _txt(d.get("calidad")) or p.calidad
        p.presentacion = p.presentacion or pres
        p.producto_base = p.producto_base or product_base(p.descripcion)
        if p.tipo in ("GR", "MP"):
            p.um = "KILOGRAMO"
    return {"nuevos": n_new, "actualizados": n_upd}


def _c_salidas(b, user, opciones):
    s = db()
    ya = _sales_keys(s)  # solo filas de cargas ANTERIORES (dentro de un archivo, líneas repetidas son legítimas)
    n = omit_dup = omit_sin = 0
    for r in b.rows:
        if r.excluida:
            continue
        d = r.datos
        if not _txt(d.get("producto")):
            omit_sin += 1
            continue
        if _txt(d.get("documento")) and _sales_key(d) in ya:
            omit_dup += 1
            continue
        s.add(SalesHistory(product_code=_txt(d["producto"]), fecha=_date(d["fecha"]), cantidad=_num(d["cantidad"]),
                           documento=_txt(d.get("documento")), origen="CARGA", upload_id=b.id))
        n += 1
    return {"registros": n, "omitidos_ya_cargados": omit_dup, "omitidos_sin_producto": omit_sin}


def _sales_key(d):
    return (_txt(d.get("documento")), _txt(d.get("producto")), _date(d.get("fecha")), round(_num(d.get("cantidad")) or 0, 6))


def _sales_keys(s):
    return {(doc, p, f, round(q or 0, 6)) for doc, p, f, q in s.execute(
        select(SalesHistory.documento, SalesHistory.product_code, SalesHistory.fecha, SalesHistory.cantidad)
        .where(SalesHistory.origen == "CARGA")).all()}
