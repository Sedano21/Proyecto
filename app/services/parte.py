"""Parte de producción en línea (ronda 23-oct-2026), como el control de producción de cada línea:
iniciar (con EAN y lote de MP) → pausar / reanudar con causa de una lista fija y foto opcional → finalizar con pallets,
defectos y colaboradores. El sistema calcula total de unidades, kg de MP, peso promedio, desviación de peso, % de
rendimiento y desmedro, horas efectivas, velocidad y el OEE (modelo OEE de Producción):

  TD = tiempo planificado – paros planificados (refrigerio, mantenimiento preventivo, reuniones)
  TP = TD – tiempo muerto;  CP = TP × velocidad nominal;  PR = producción real;  PB = PR – defectuosos
  Disponibilidad = TP / TD;  Rendimiento = PR / CP;  Calidad = PB / PR;  OEE = D × R × C

Desviación en selección (Cimbria, Truer): % rendimiento + % desmedro – 100. En envasado: (peso promedio – presentación) /
presentación y los kg entregados de más."""
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

from sqlalchemy import select

from ..db import db
from ..models import Product, ProductGTIN, ProductionReport, now
from . import alerts
from .core import BusinessError, audit, has_perm, param, user_scope

ESTADOS_PARTE = ["PROGRAMADO", "EN PROCESO", "PAUSADO", "PENDIENTE", "LIQUIDADO", "LIQ. EXT. PENDIENTE", "LIQ. EXT. LIQUIDADO",
                 "REPROCESO", "REENVASADO", "BLOQUEADO", "ANULADO"]
LINEAS_ENVASADO = ("R1", "R2", "R3", "ENV_MANUAL")


def tipo_linea(linea):
    return "ENVASADO" if linea in LINEAS_ENVASADO else "SELECCION"


def lines_for(user):
    """Líneas que ve el usuario: las de su alcance (asistentes) o todas."""
    from ..catalogs import PRODUCTION_LINES
    todas = [c for c, _, _ in PRODUCTION_LINES]
    propias = user_scope(user).get("lineas")
    return [l for l in todas if l in propias] if propias else todas


def can_see_line(user, linea):
    return linea in lines_for(user)


def _check_line(user, r):
    if not can_see_line(user, r.linea):
        raise BusinessError(f"La línea {r.linea} no está en su alcance.")


def causes(tipo):
    c = param("causas_parada", {}) or {}
    return list(c.get(tipo) or []) + [x for x in (param("causas_paro_planificado", []) or []) if x not in (c.get(tipo) or [])]


def _min(a, b):
    return max((b - a).total_seconds() / 60.0, 0)


def _dt(fecha, hhmm):
    if not hhmm:
        return None
    try:
        h, m = [int(x) for x in str(hhmm)[:5].split(":")]
        return datetime.combine(fecha, datetime.min.time()).replace(hour=h, minute=m)
    except (ValueError, TypeError):
        return None


# ------------------------------------------------------------------ horas en formato 12 h o 24 h
def parse_hora(txt, campo="Hora"):
    """Acepta «15:30», «7:05», «3:30 pm», «3 pm», «03:30 p. m.» y devuelve «HH:MM» (24 h). Vacío → ""."""
    t = " ".join(str(txt or "").strip().lower().replace(".", "").split())
    if not t:
        return ""
    sufijo = None
    for suf in ("am", "pm", "a m", "p m"):
        if t.endswith(suf):
            sufijo, t = suf.replace(" ", ""), t[: -len(suf)].strip()
            break
    try:
        h, m = (t.split(":") + ["0"])[:2] if ":" in t else (t, "0")
        h, m = int(h), int(m)
    except ValueError:
        raise BusinessError(f"{campo}: «{txt}» no es una hora válida (use 15:30 o 3:30 pm).")
    if sufijo:
        if not 1 <= h <= 12:
            raise BusinessError(f"{campo}: en formato de 12 horas la hora va de 1 a 12.")
        h = (0 if h == 12 else h) + (12 if sufijo == "pm" else 0)
    if not (0 <= h <= 23 and 0 <= m <= 59):
        raise BusinessError(f"{campo}: «{txt}» no es una hora válida.")
    return f"{h:02d}:{m:02d}"


def fmt_hora(hhmm):
    """Muestra la hora según el parámetro formato_hora (24H o 12H)."""
    if not hhmm:
        return ""
    try:
        h, m = [int(x) for x in str(hhmm)[:5].split(":")]
    except ValueError:
        return str(hhmm)
    if (param("formato_hora", "24H") or "24H").upper() == "12H":
        return f"{(h % 12) or 12}:{m:02d} {'p. m.' if h >= 12 else 'a. m.'}"
    return f"{h:02d}:{m:02d}"


# ------------------------------------------------------------------ validaciones por producto y paquete
def check_product_for_line(linea, product, presentacion_kg=None, und_paquete=None):
    """Envasado: el producto es PT, la presentación coincide con la del artículo y las unidades por paquete son un entero
    positivo. Devuelve la presentación a usar."""
    if tipo_linea(linea) == "ENVASADO":
        if product.tipo and product.tipo != "PT":
            raise BusinessError(f"En envasado se registra producto terminado; {product.code} es {product.tipo}.")
        if product.peso_kg and presentacion_kg and abs(float(presentacion_kg) - float(product.peso_kg)) > 1e-6:
            raise BusinessError(f"La presentación {float(presentacion_kg):g} kg no corresponde a {product.code} "
                                f"({product.peso_kg:g} kg).")
        presentacion_kg = presentacion_kg or product.peso_kg
    if und_paquete not in (None, ""):
        u = float(und_paquete)
        if u <= 0 or u != int(u):
            raise BusinessError("Las unidades por paquete deben ser un número entero mayor que cero.")
    return presentacion_kg


def product_info(code):
    """Datos para validar en pantalla (tiempo real): presentación, tipo, GTIN registrados y tolerancia."""
    p = db().get(Product, (code or "").strip().upper().split(" ")[0])
    if p is None:
        return None
    gt = [g for g in db().execute(select(ProductGTIN.gtin).where(ProductGTIN.product_code == p.code)).scalars()]
    return {"code": p.code, "descripcion": p.descripcion, "tipo": p.tipo, "presentacion_kg": p.peso_kg,
            "und_bulto": p.unidades_por_bulto, "gtins": gt}


# ------------------------------------------------------------------ colaboradores registrados
def validate_workers(colaboradores, operarios=None, fecha=None):
    """Solo colaboradores registrados y activos; sin repetir; no más que los operarios; horas válidas."""
    from ..models import Worker
    activos = {w.id: w for w in db().execute(select(Worker).where(Worker.activo.is_(True))).scalars()}
    por_nombre = {w.nombre.upper(): w for w in activos.values()}
    out, vistos = [], set()
    for c in colaboradores or []:
        wid, nombre = c.get("worker_id"), (c.get("nombre") or "").strip()
        if not wid and not nombre:
            continue
        w = activos.get(int(wid)) if str(wid or "").isdigit() else por_nombre.get(nombre.upper())
        if w is None:
            raise BusinessError(f"«{nombre or wid}» no es un colaborador registrado y activo. Regístrelo primero en Asistencia.")
        if w.id in vistos:
            raise BusinessError(f"{w.nombre} está repetido.")
        vistos.add(w.id)
        ini, fin = parse_hora(c.get("inicio"), f"Inicio de {w.nombre}"), parse_hora(c.get("fin"), f"Fin de {w.nombre}")
        if ini and fin and fin < ini:
            raise BusinessError(f"{w.nombre}: la hora de fin no puede ser menor que la de inicio.")
        kg = c.get("kg")
        x = {"worker_id": w.id, "nombre": w.nombre, "inicio": ini, "fin": fin,
             "refrigerio_min": c.get("refrigerio_min") or "", "kg": kg if kg not in (None, "") else ""}
        for k in ("merma_kg", "remanente_kg"):  # V12: líneas manuales, por persona
            v = _f(c.get(k))
            if v is not None:
                if v < 0:
                    raise BusinessError(f"{w.nombre}: {k.replace('_kg', '')} no puede ser negativo.")
                x[k] = v
        out.append(x)
    if operarios and len(out) > int(operarios):
        raise BusinessError(f"Registró {len(out)} colaboradores para {int(operarios)} operarios.")
    return out


# ------------------------------------------------------------------ EAN
def check_ean(r, user):
    """Compara el EAN escaneado con la tabla GS1 ↔ SAP. ALERTAR (avisa y registra) o BLOQUEAR."""
    ean = (r.ean or "").strip()
    if not ean:
        r.ean_ok = None
        return None
    g = db().get(ProductGTIN, ean)
    ok = g is not None and g.product_code == r.product_code
    r.ean_ok = ok
    if not ok:
        detalle = (f"el EAN {ean} corresponde a {g.product_code} ({g.descripcion})" if g is not None
                   else f"el EAN {ean} no está en la tabla GS1 ↔ SAP")
        if (param("ean_validacion", "ALERTAR") or "").upper() == "BLOQUEAR":
            raise BusinessError(f"No coincide: {detalle}, y se está produciendo {r.product_code}.")
        alerts.emit("EAN_NO_COINCIDE", f"EAN no coincide en {r.linea}: {r.product_code}",
                    f"Al iniciar el parte, {detalle}. Verificar la bobina / etiqueta antes de continuar.", "ProductionReport",
                    r.id, dedup_key=f"EAN:{r.id}:{ean}", creado_por=user)
    return ok


def tolerance_for(linea, presentacion_kg=None):
    """Regla de rendimiento más específica para la línea y presentación (balance de masa)."""
    reglas = param("tolerancia_rendimiento", []) or []

    def match_linea(p):
        p = (p or "*").upper()
        return p == "*" or p == linea or (p.endswith("*") and linea.startswith(p[:-1]))

    def match_pres(p):
        return p in (None, "") or (presentacion_kg is not None and abs(float(p) - float(presentacion_kg)) < 1e-6)
    candidatas = [r for r in reglas if match_linea(r.get("linea")) and match_pres(r.get("presentacion_kg"))]
    if not candidatas:
        return None

    def score(r):
        l = (r.get("linea") or "*")
        return (2 if l not in ("*",) and not l.endswith("*") else 1 if l != "*" else 0) + (2 if r.get("presentacion_kg") not in (None, "") else 0)
    return max(candidatas, key=score)


def ideal_output(linea, mp_kg, presentacion_kg):
    """Producción ideal (100 %) y mínima aceptable según la tolerancia, al digitar el lote y los kg de MP."""
    if not mp_kg or not presentacion_kg:
        return None
    ideal = mp_kg / presentacion_kg
    t = tolerance_for(linea, presentacion_kg) or {}
    minimo = ideal * float(t["min_pct"]) / 100 if t.get("min_pct") else None
    return {"ideal": round(ideal, 1), "minimo": round(minimo, 1) if minimo else None, "regla": t}


# ------------------------------------------------------------------ ciclo del parte
def start(user, datos, r=None):
    """Inicia un parte (nuevo o uno PROGRAMADO)."""
    if not has_perm(user, "produccion.liquidar"):
        raise BusinessError("Registra el parte Producción (asistentes de producción).")
    from ..catalogs import PRODUCTION_LINES
    if r is None:
        linea = datos.get("linea")
        if linea not in {c for c, _, _ in PRODUCTION_LINES}:
            raise BusinessError("Línea de producción no válida.")
        code = (datos.get("product_code") or "").strip().upper().split(" ")[0]
        if not code or db().get(Product, code) is None:
            raise BusinessError("Producto no existe.")
        r = ProductionReport(linea=linea, tipo=tipo_linea(linea), fecha=datos.get("fecha") or date.today(), product_code=code,
                             peso_entrada_kg=0, user_id=user.id)
        db().add(r)
    elif r.estado != "PROGRAMADO":
        raise BusinessError("Solo se inicia un parte programado.")
    _check_line(user, r)
    for k in ("of", "op", "turno", "cliente", "mp_codigo", "lote", "maquina", "maquinista", "ean", "lote_pt"):
        if datos.get(k) not in (None, ""):
            setattr(r, k, str(datos[k]).strip().upper() if k in ("mp_codigo", "maquina") else str(datos[k]).strip())
    for k in ("presentacion_kg", "und_paquete", "sacos", "peso_saco_kg", "saldo_mp_kg", "peso_entrada_kg"):
        if datos.get(k) not in (None, ""):
            setattr(r, k, float(datos[k]))
    for k in ("fv_mp", "fv_pt"):
        if datos.get(k):
            setattr(r, k, datos[k] if isinstance(datos[k], date) else date.fromisoformat(str(datos[k])))
    if r.sacos and r.peso_saco_kg:
        r.peso_entrada_kg = round(r.sacos * r.peso_saco_kg + (r.saldo_mp_kg or 0), 3)
    if not r.peso_entrada_kg or r.peso_entrada_kg <= 0:
        raise BusinessError("Indique los kg de MP (peso por saco y número de sacos, o el total).")
    prod = db().get(Product, r.product_code)
    if prod is not None:
        r.presentacion_kg = check_product_for_line(r.linea, prod, r.presentacion_kg, r.und_paquete)
    if r.tipo == "ENVASADO" and not r.presentacion_kg:
        raise BusinessError("Indique la presentación (kg) para envasado.")
    r.maquina = r.maquina or r.linea
    r.estado = "EN PROCESO"
    r.inicio_ts = now()
    r.hora_inicio = f"{r.inicio_ts:%H:%M}"
    r.motivo_estado = ""
    db().flush()
    check_ean(r, user)
    audit(user, "PARTE_INICIADO", "ProductionReport", r.id, {"linea": r.linea, "of": r.of, "ean": r.ean, "ean_ok": r.ean_ok})
    return r


def pause(user, r, causa, foto=""):
    _check_line(user, r)
    if r.estado != "EN PROCESO":
        raise BusinessError("Solo se pausa un parte en proceso.")
    if causa not in causes(r.tipo):
        raise BusinessError("Elija la causa de la lista.")
    r.pausas = list(r.pausas or []) + [{"inicio": now().isoformat(timespec="seconds"), "fin": None, "causa": causa,
                                        "foto": foto or "", "user": user.nombre}]
    r.estado = "PAUSADO"
    audit(user, "PARTE_PAUSADO", "ProductionReport", r.id, {"causa": causa})
    return r


def resume(user, r, observacion=""):
    _check_line(user, r)
    if r.estado != "PAUSADO":
        raise BusinessError("El parte no está pausado.")
    ps = list(r.pausas or [])
    if ps and not ps[-1].get("fin"):
        ps[-1] = dict(ps[-1], fin=now().isoformat(timespec="seconds"), observacion=(observacion or "").strip())
    r.pausas = ps
    r.estado = "EN PROCESO"
    audit(user, "PARTE_REANUDADO", "ProductionReport", r.id, {})
    return r


def _total_pallets(r):
    return sum(float(p.get("cantidad") or 0) for p in (r.pallets or []))


def finish(user, r, datos):
    """Registra los resultados y calcula los indicadores. datos: hora_fin, refrigerio_min, operarios, pallets
    [{destino, cantidad}], unidades_sueltas, unidades (si no hay pallets), defectos {causa: und}, colaboradores, granel_kg,
    mp_kg, cal3_kg, desmedro_kg, sobrante_kg, observacion."""
    _check_line(user, r)
    if r.estado not in ("EN PROCESO", "PAUSADO"):
        raise BusinessError("Solo se finaliza un parte en proceso.")
    if r.estado == "PAUSADO":
        resume(user, r)
    r.fin_ts = now()
    # la hora de fin es la del cierre: no puede ser menor que la hora actual ni que la de inicio
    hf = parse_hora(datos.get("hora_fin"), "Hora de fin") if datos.get("hora_fin") else ""
    if hf and hf < f"{r.fin_ts:%H:%M}" and r.inicio_ts and r.inicio_ts.date() == r.fin_ts.date():
        raise BusinessError("La hora de fin no puede ser menor que la hora actual.")
    r.hora_fin = f"{r.fin_ts:%H:%M}"
    if r.inicio_ts and r.fin_ts < r.inicio_ts:
        raise BusinessError("La hora de fin no puede ser menor que la de inicio.")
    if datos.get("refrigerio_min") not in (None, ""):
        r.refrigerio_min = int(float(datos["refrigerio_min"]))
    if datos.get("operarios") not in (None, ""):
        r.operarios = int(float(datos["operarios"]))
    if not r.operarios:
        raise BusinessError("Indique el número de operarios (SAP lo exige para liquidar).")
    r.pallets = [p for p in (datos.get("pallets") or []) if float(p.get("cantidad") or 0) > 0]
    defectos = {k: float(v) for k, v in (datos.get("defectos") or {}).items() if v not in (None, "") and float(v) > 0}
    malas = [k for k in defectos if k not in (param("causas_defecto", []) or [])]
    if malas:
        raise BusinessError("Causa de defecto no válida: " + ", ".join(malas))
    r.defectos = defectos
    r.colaboradores = validate_workers(datos.get("colaboradores"), r.operarios)
    for k in ("granel_kg", "mp_kg", "cal3_kg", "desmedro_kg", "sobrante_kg"):
        if datos.get(k) not in (None, ""):
            setattr(r, k, float(datos[k]))
    if r.tipo == "ENVASADO":
        sueltas = float(datos.get("unidades_sueltas") or 0)
        if r.und_paquete and r.pallets:
            if sueltas >= r.und_paquete:
                raise BusinessError(f"Las unidades sueltas ({sueltas:g}) completan un paquete de {r.und_paquete:g}: "
                                    "regístrelo en el pallet.")
            if any(float(p.get("cantidad") or 0) != int(float(p.get("cantidad") or 0)) for p in r.pallets):
                raise BusinessError("En envasado cada pallet se registra en paquetes enteros.")
        if r.pallets:
            r.unidades = _total_pallets(r) * (r.und_paquete or 1) + sueltas
        elif datos.get("unidades") not in (None, ""):
            r.unidades = float(datos["unidades"])
        if r.unidades is None:
            raise BusinessError("Registre los pallets (paquetes por pallet) o las unidades producidas.")
    elif r.pallets and not datos.get("granel_kg"):
        r.granel_kg = _total_pallets(r)
    r.observacion = ((r.observacion + "\n") if r.observacion else "") + (datos.get("observacion") or "").strip()
    compute(r)
    r.estado = "PENDIENTE"
    r.motivo_estado = "liquidación de la OF en SAP"
    db().flush()
    _alert_merma(r, user)
    from .quality_v13 import open_merma_validation  # V13: Calidad valida la merma de cada proceso terminado
    open_merma_validation(r, user)
    audit(user, "PARTE_FINALIZADO", "ProductionReport", r.id, {"oee": (r.indicadores or {}).get("oee"),
                                                               "rendimiento": r.rendimiento_pct})
    return r


def _alert_merma(r, user):
    umbral = param("umbral_merma_pct")
    if umbral is not None and r.desmedro_pct is not None and r.desmedro_pct > float(umbral):
        alerts.emit("MERMA_FUERA_PARAMETRO", f"Desmedro {r.desmedro_pct}% en {r.linea} (OF {r.of or '-'})",
                    f"Umbral configurado {umbral}%. Producto {r.product_code}, lote {r.lote}. Rendimiento {r.rendimiento_pct}%.",
                    "ProductionRequest", "", creado_por=user)
    # V12: la merma y el remanente de MP de cada parte necesitan un destino: deciden Calidad, Operaciones y Logística
    merma = r.desmedro_kg or ((r.desmedro_pct or 0) * (r.peso_entrada_kg or 0) / 100)
    rem = r.sobrante_kg or 0
    if merma > 0 or rem > 0:
        por_persona = [f"{c['nombre']}: merma {c.get('merma_kg', 0):g} kg, remanente {c.get('remanente_kg', 0):g} kg"
                       for c in (r.colaboradores or []) if c.get("merma_kg") or c.get("remanente_kg")]
        alerts.emit("MERMA_REMANENTE_DESTINO", f"Definir destino: {merma:,.1f} kg de merma y {rem:,.1f} kg de remanente "
                                               f"de MP en {r.linea} (OF {r.of or '-'})",
                    f"Producto {r.product_code}, MP {r.mp_codigo or '-'}, lote {r.lote or '-'}. "
                    + ("Por persona: " + "; ".join(por_persona) + ". " if por_persona else "")
                    + "Decidir si se reprocesa, reingresa al almacén o se usa en el próximo lote.",
                    "ProductionReport", r.id, creado_por=user, dedup_key=f"MERMA_REM:{r.id}")


def velocity(r):
    v = param("velocidad_nominal", {}) or {}
    maq = (r.maquina or r.linea or "").upper()
    if r.presentacion_kg:
        for k in (f"{maq}|{r.presentacion_kg:g}", f"{r.linea}|{r.presentacion_kg:g}"):
            if k in v:
                return float(v[k])
    for k in (maq, r.linea):
        if k in v:
            return float(v[k])
    return None


def rating(oee_pct):
    for x in sorted(param("oee_calificacion", []) or [], key=lambda x: -float(x["min"])):
        if oee_pct >= float(x["min"]):
            return x
    return None


def compute(r):
    """Fórmulas del control de producción y del modelo OEE."""
    ind = {}
    ini = r.inicio_ts or _dt(r.fecha, r.hora_inicio)
    if r.fin_ts and r.hora_fin == f"{r.fin_ts:%H:%M}":
        fin = r.fin_ts
    else:
        fin = _dt(ini.date() if ini else r.fecha, r.hora_fin)
    if ini and fin and fin < ini:
        fin += timedelta(days=1)
    planificado = _min(ini, fin) if ini and fin else None
    planificadas = set(param("causas_paro_planificado", []) or [])
    paro_plan = float(r.refrigerio_min or 0)
    muerto = 0.0
    por_causa = Counter()
    for p in r.pausas or []:
        if not p.get("fin"):
            continue
        m = _min(datetime.fromisoformat(p["inicio"]), datetime.fromisoformat(p["fin"]))
        por_causa[p["causa"]] += m
        if p["causa"] in planificadas:
            paro_plan += m
        else:
            muerto += m
    ind.update({"planificado_min": round(planificado, 1) if planificado is not None else None,
                "paro_planificado_min": round(paro_plan, 1), "tiempo_muerto_min": round(muerto, 1),
                "pausas_por_causa": {k: round(v, 1) for k, v in por_causa.items()}})
    r.parada_min = int(round(muerto))
    r.motivo_parada = ", ".join(k for k, _ in por_causa.most_common(3) if k not in planificadas)
    if planificado is not None:
        r.horas = round((planificado - float(r.refrigerio_min or 0)) / 60.0, 2)
        ind["horas_efectivas"] = round(max(planificado - paro_plan - muerto, 0) / 60.0, 2)
        ind["horas_hombre"] = round(r.horas * (r.operarios or 0), 2)
    entrada = r.peso_entrada_kg or 0
    if r.tipo == "ENVASADO":
        pres = r.presentacion_kg or 0
        und = r.unidades or 0
        r.granel_kg = round(und * pres, 3)
        r.rendimiento_pct = round(100 * und * pres / entrada, 2) if entrada and pres else None
        perdida = max(entrada - und * pres - (r.sobrante_kg or 0), 0)
        r.desmedro_kg = round(perdida, 3)
        r.diferencia_kg = round(perdida - (r.cal3_kg or 0), 3)
        consumido = entrada - (r.sobrante_kg or 0)
        if und and pres:
            prom = consumido / und
            ind["peso_promedio_kg"] = round(prom, 4)
            ind["desviacion_peso_pct"] = round(100 * (prom - pres) / pres, 2)
            ind["kg_sobrepeso"] = round(max(prom - pres, 0) * und, 2)
        if ind.get("horas_efectivas"):
            ind["velocidad"] = round(und / (ind["horas_efectivas"] * 60), 2)
            ind["velocidad_um"] = "und/min"
    else:
        total = (r.granel_kg or 0) + (r.mp_kg or 0) + (r.cal3_kg or 0) + (r.desmedro_kg or 0)
        r.rendimiento_pct = round(100 * (r.granel_kg or 0) / entrada, 2) if entrada else None
        r.diferencia_kg = round(entrada - total, 3)
        if ind.get("horas_efectivas"):
            ind["velocidad"] = round((r.granel_kg or 0) / ind["horas_efectivas"], 1)
            ind["velocidad_um"] = "kg/h"
    perdida = r.desmedro_kg if r.tipo == "SELECCION" else max(entrada - (r.granel_kg or 0), 0)
    r.desmedro_pct = round(100 * (perdida or 0) / entrada, 2) if entrada else None
    if r.tipo == "SELECCION" and r.rendimiento_pct is not None and r.desmedro_pct is not None:
        ind["desviacion_pct"] = round(r.rendimiento_pct + r.desmedro_pct - 100, 2)
    if r.horas:
        r.ratio_kg_hora = round((r.granel_kg or 0) / r.horas, 1)
    # OEE
    vel = velocity(r)
    defect = sum((r.defectos or {}).values())
    pr = ((r.unidades or 0) + defect) if r.tipo == "ENVASADO" else (r.granel_kg or 0) + defect
    ind["defectuosos"] = defect
    if planificado is not None and vel:
        td = planificado - paro_plan
        tp = max(td - muerto, 0)
        cp = tp * (vel if r.tipo == "ENVASADO" else vel / 60.0)
        disp = tp / td if td > 0 else None
        rend = pr / cp if cp > 0 else None
        cal = (pr - defect) / pr if pr > 0 else None
        ind.update({"td_min": round(td, 1), "tp_min": round(tp, 1), "velocidad_nominal": vel, "cp": round(cp, 1), "pr": pr,
                    "pb": pr - defect})
        if None not in (disp, rend, cal):
            oee = disp * rend * cal * 100
            cal_txt = rating(oee)
            ind.update({"disponibilidad": round(disp * 100, 1), "rendimiento_oee": round(rend * 100, 1),
                        "calidad": round(cal * 100, 1), "oee": round(oee, 1),
                        "calificacion": cal_txt["nombre"] if cal_txt else "", "plan_accion": cal_txt.get("plan") if cal_txt else ""})
    elif not vel:
        ind["oee_nota"] = "Sin velocidad nominal configurada para la máquina / presentación."
    r.indicadores = ind
    return r


# ------------------------------------------------------------------ estados, Calidad
def set_state(user, r, estado, motivo=""):
    _check_line(user, r)
    if estado not in ESTADOS_PARTE:
        raise BusinessError("Estado no válido.")
    if estado == "BLOQUEADO" or r.estado == "BLOQUEADO":
        raise BusinessError("El bloqueo y la liberación los registra Calidad.")
    if not has_perm(user, "produccion.liquidar"):
        raise BusinessError("Actualiza el estado Producción.")
    if estado == "PENDIENTE" and not (motivo or "").strip():
        raise BusinessError("Indique de qué está pendiente el parte (p. ej. «material de empaque», «liquidación en SAP»).")
    antes = r.estado
    r.estado = estado
    r.motivo_estado = (motivo or "").strip()[:200]
    audit(user, "PARTE_ESTADO", "ProductionReport", r.id, {"antes": antes, "despues": estado, "motivo": motivo})
    return r


# ------------------------------------------------------------------ programación (formato único: Registrar / Programar)
NO_EDITABLES = ("LIQUIDADO", "LIQ. EXT. LIQUIDADO", "ANULADO")


def program(user, datos):
    """Programa un parte a futuro con el mismo formato del registro: línea, producto, fecha, kg previstos y horario."""
    from ..catalogs import PRODUCTION_LINES
    if not has_perm(user, "produccion.liquidar"):
        raise BusinessError("Programa los partes Producción.")
    linea = datos.get("linea")
    if linea not in {c for c, _, _ in PRODUCTION_LINES}:
        raise BusinessError("Línea de producción no válida.")
    if not can_see_line(user, linea):
        raise BusinessError(f"La línea {linea} no está en su alcance.")
    code = (datos.get("product_code") or "").strip().upper().split(" ")[0]
    prod = db().get(Product, code) if code else None
    if prod is None:
        raise BusinessError("Producto no existe.")
    fecha = datos.get("fecha")
    if isinstance(fecha, str):
        fecha = date.fromisoformat(fecha) if fecha else None
    if fecha is None or fecha < date.today():
        raise BusinessError("Programe para hoy o una fecha futura.")
    kg = _kg(datos)
    if not kg or kg <= 0:
        raise BusinessError("Indique los kg de MP previstos.")
    pres = check_product_for_line(linea, prod, _f(datos.get("presentacion_kg")), datos.get("und_paquete"))
    ini, fin = _horario(fecha, datos.get("hora_inicio"), datos.get("hora_fin"))
    r = ProductionReport(linea=linea, tipo=tipo_linea(linea), fecha=fecha, product_code=code, peso_entrada_kg=kg,
                         presentacion_kg=pres, und_paquete=_f(datos.get("und_paquete")), estado="PROGRAMADO",
                         turno=(datos.get("turno") or "").strip(), of=(datos.get("of") or "").strip(),
                         cliente=(datos.get("cliente") or "").strip(), maquina=(datos.get("maquina") or linea).strip().upper(),
                         programa={"hora_inicio": ini, "hora_fin": fin, "kg": kg, "fecha": fecha.isoformat()},
                         user_id=user.id)
    db().add(r)
    db().flush()
    audit(user, "PARTE_PROGRAMADO", "ProductionReport", r.id, {"linea": linea, "fecha": fecha.isoformat(), "kg": kg})
    return r


def _f(v):
    v = str(v if v is not None else "").replace(",", ".").strip()
    return float(v) if v else None


def _kg(datos):
    sacos, peso = _f(datos.get("sacos")), _f(datos.get("peso_saco_kg"))
    if sacos and peso:
        return round(sacos * peso + (_f(datos.get("saldo_mp_kg")) or 0), 3)
    return _f(datos.get("peso_entrada_kg"))


def _horario(fecha, hora_inicio, hora_fin):
    ini, fin = parse_hora(hora_inicio, "Hora de inicio"), parse_hora(hora_fin, "Hora de fin")
    if ini and fin and fin <= ini:
        raise BusinessError("La hora de fin no puede ser menor ni igual a la de inicio.")
    if fin and fecha == date.today() and fin < f"{now():%H:%M}":
        raise BusinessError("La hora de fin no puede ser menor que la hora actual.")
    return ini, fin


def edit_program(user, r, fecha=None, cantidad=None, motivo="", hora_inicio=None, hora_fin=None):
    """Editar fecha y cantidad (kg de MP) de un parte mientras no esté liquidado o anulado."""
    _check_line(user, r)
    if not has_perm(user, "produccion.liquidar"):
        raise BusinessError("Modifica la programación Producción.")
    if r.estado in NO_EDITABLES:
        raise BusinessError(f"El parte está {r.estado.lower()}: ya no se modifica.")
    if not (motivo or "").strip():
        raise BusinessError("Indique el motivo del cambio.")
    if isinstance(fecha, str):
        fecha = date.fromisoformat(fecha) if fecha else None
    antes = {"fecha": r.fecha.isoformat(), "kg": r.peso_entrada_kg}
    if fecha is not None and fecha != r.fecha:
        if r.estado == "PROGRAMADO" and fecha < date.today():
            raise BusinessError("La nueva fecha no puede ser pasada.")
        r.fecha = fecha
    if cantidad is not None:
        if cantidad <= 0:
            raise BusinessError("La cantidad debe ser mayor que cero.")
        r.peso_entrada_kg = cantidad
    if r.estado == "PROGRAMADO":
        ini, fin = _horario(r.fecha, hora_inicio if hora_inicio is not None else (r.programa or {}).get("hora_inicio"),
                            hora_fin if hora_fin is not None else (r.programa or {}).get("hora_fin"))
        r.programa = dict(r.programa or {}, fecha=r.fecha.isoformat(), kg=r.peso_entrada_kg, hora_inicio=ini, hora_fin=fin)
    elif r.estado not in ("EN PROCESO", "PAUSADO"):
        compute(r)  # con resultados: se recalculan rendimiento y desmedro
    db().flush()
    audit(user, "PROGRAMACION_EDITADA", "ProductionReport", r.id,
          {"antes": antes, "despues": {"fecha": r.fecha.isoformat(), "kg": r.peso_entrada_kg}, "motivo": motivo.strip()})
    return r


def add_progress(user, r, destino, cantidad):
    """Avance en línea: registra un pallet mientras produce; el monitoreo muestra rendimiento y desmedro parciales."""
    _check_line(user, r)
    if r.estado not in ("EN PROCESO", "PAUSADO"):
        raise BusinessError("El avance se registra con el parte en proceso.")
    q = _f(cantidad)
    if not q or q <= 0:
        raise BusinessError("Indique la cantidad del pallet.")
    if r.tipo == "ENVASADO" and q != int(q):
        raise BusinessError("En envasado el pallet se registra en paquetes enteros.")
    r.pallets = list(r.pallets or []) + [{"destino": destino or "DESPACHO", "cantidad": q,
                                          "ts": now().isoformat(timespec="minutes")}]
    audit(user, "PARTE_AVANCE", "ProductionReport", r.id, {"destino": destino, "cantidad": q})
    return r


def live(r):
    """Rendimiento y desmedro parciales (en curso) o finales."""
    if r.estado not in ("EN PROCESO", "PAUSADO"):
        return {"rend": r.rendimiento_pct, "desm": r.desmedro_pct, "salida_kg": r.granel_kg, "final": True}
    entrada = r.peso_entrada_kg or 0
    paq = _total_pallets(r)
    salida = paq * (r.und_paquete or 1) * (r.presentacion_kg or 0) if r.tipo == "ENVASADO" else paq
    rend = round(100 * salida / entrada, 1) if entrada else None
    t = tolerance_for(r.linea, r.presentacion_kg) or {}
    minutos = _min(r.inicio_ts, now()) if r.inicio_ts else 0
    return {"rend": rend, "desm": round(100 - rend, 1) if rend is not None else None, "salida_kg": round(salida, 1),
            "final": False, "min_pct": t.get("min_pct"), "minutos": round(minutos),
            "alerta": t.get("min_pct") is not None and rend is not None and minutos > 30 and rend < float(t["min_pct"]) * 0.5}


def monitor(user):
    """Partes del día y en curso de las líneas del usuario para el monitoreo en tiempo real."""
    vis = lines_for(user)
    rows = [r for r in db().execute(select(ProductionReport).where(
        (ProductionReport.estado.in_(["EN PROCESO", "PAUSADO"])) | (ProductionReport.fecha == date.today()))
        .order_by(ProductionReport.linea)).scalars() if r.linea in vis]
    return [(r, live(r)) for r in rows]


def block(user, r, motivo):
    if not has_perm(user, "calidad.bloquear_parte"):
        raise BusinessError("Bloquea un parte Calidad.")
    if not (motivo or "").strip():
        raise BusinessError("Indique el problema encontrado.")
    r.bloqueo = {"motivo": motivo.strip(), "user": user.nombre, "ts": now().isoformat(timespec="minutes"),
                 "estado_anterior": r.estado}
    r.estado = "BLOQUEADO"
    alerts.emit("PARTE_BLOQUEADO", f"Parte bloqueado por Calidad: {r.linea} {r.product_code} (OF {r.of or '-'})",
                f"Lote {r.lote_pt or r.lote or '-'}. Motivo: {motivo.strip()}. Se libera tras fumigación o reproceso.",
                "ProductionReport", r.id, creado_por=user)
    audit(user, "PARTE_BLOQUEADO", "ProductionReport", r.id, {"motivo": motivo})
    return r


def release(user, r, motivo, nuevo_estado="REPROCESO"):
    if not has_perm(user, "calidad.bloquear_parte"):
        raise BusinessError("Libera un parte Calidad.")
    if r.estado != "BLOQUEADO":
        raise BusinessError("El parte no está bloqueado.")
    if nuevo_estado not in ("REPROCESO", "PENDIENTE", "EN PROCESO", "ANULADO"):
        raise BusinessError("Estado de liberación no válido.")
    if not (motivo or "").strip():
        raise BusinessError("Indique la acción realizada (fumigación, reproceso…).")
    r.bloqueo = dict(r.bloqueo or {}, liberado={"motivo": motivo.strip(), "user": user.nombre,
                                                "ts": now().isoformat(timespec="minutes")})
    r.estado = nuevo_estado
    audit(user, "PARTE_LIBERADO", "ProductionReport", r.id, {"motivo": motivo, "estado": nuevo_estado})
    return r


# ------------------------------------------------------------------ tablero de Producción
def dashboard(user, desde=None, hasta=None, linea="", turno="", maquina=""):
    hasta = hasta or date.today()
    desde = desde or hasta - timedelta(days=29)
    q = select(ProductionReport).where(ProductionReport.fecha >= desde, ProductionReport.fecha <= hasta)
    rows = [r for r in db().execute(q).scalars() if can_see_line(user, r.linea)]
    if linea:
        rows = [r for r in rows if r.linea == linea]
    if turno:
        rows = [r for r in rows if r.turno == turno]
    if maquina:
        rows = [r for r in rows if (r.maquina or r.linea) == maquina]
    con_oee = [r for r in rows if (r.indicadores or {}).get("oee") is not None]

    def prom(lst, k):
        vals = [(r.indicadores or {}).get(k) for r in lst if (r.indicadores or {}).get(k) is not None]
        return round(sum(vals) / len(vals), 1) if vals else None
    por_linea = defaultdict(list)
    for r in rows:
        por_linea[r.linea].append(r)
    lineas = []
    for l, lst in sorted(por_linea.items()):
        entrada = sum(x.peso_entrada_kg or 0 for x in lst)
        perdida = sum(x.desmedro_kg or 0 for x in lst)
        lineas.append({"linea": l, "partes": len(lst), "oee": prom([x for x in lst if x in con_oee], "oee"),
                       "disponibilidad": prom(lst, "disponibilidad"), "rendimiento": prom(lst, "rendimiento_oee"),
                       "calidad": prom(lst, "calidad"), "merma_pct": round(100 * perdida / entrada, 2) if entrada else None,
                       "muerto_min": round(sum((x.indicadores or {}).get("tiempo_muerto_min") or 0 for x in lst), 1),
                       "kg": round(sum(x.granel_kg or 0 for x in lst), 1)})
    causas = Counter()
    defectos = Counter()
    for r in rows:
        for k, v in ((r.indicadores or {}).get("pausas_por_causa") or {}).items():
            causas[k] += v
        for k, v in (r.defectos or {}).items():
            defectos[k] += v
    semanas = defaultdict(lambda: defaultdict(float))
    productos = defaultdict(float)
    for r in rows:
        lunes = r.fecha - timedelta(days=r.fecha.weekday())
        semanas[lunes][r.linea] += r.granel_kg or 0
        productos[r.product_code] += r.granel_kg or 0
    serie_oee = defaultdict(list)
    for r in con_oee:
        serie_oee[r.fecha].append(r.indicadores["oee"])
    cal = rating(prom(con_oee, "oee")) if con_oee else None
    return {"desde": desde, "hasta": hasta, "partes": len(rows), "oee": prom(con_oee, "oee"),
            "disponibilidad": prom(con_oee, "disponibilidad"), "rendimiento": prom(con_oee, "rendimiento_oee"),
            "calidad": prom(con_oee, "calidad"), "calificacion": cal, "lineas": lineas,
            "causas": [(k, round(v, 1)) for k, v in causas.most_common(12)], "defectos": defectos.most_common(10),
            "semanas": sorted(semanas.items()), "productos": sorted(productos.items(), key=lambda x: -x[1])[:15],
            "serie_oee": sorted((d, round(sum(v) / len(v), 1)) for d, v in serie_oee.items()),
            "sin_velocidad": len([r for r in rows if (r.indicadores or {}).get("oee_nota")]),
            "kg_sobrepeso": round(sum((r.indicadores or {}).get("kg_sobrepeso") or 0 for r in rows), 1)}


def productivity(user, desde=None, hasta=None):
    """Productividad por persona en selección manual: horas (fin – inicio – refrigerio) y kg por colaborador."""
    hasta = hasta or date.today()
    desde = desde or hasta - timedelta(days=29)
    acc = defaultdict(lambda: {"horas": 0.0, "kg": 0.0, "partes": 0})
    for r in db().execute(select(ProductionReport).where(ProductionReport.fecha >= desde,
                                                         ProductionReport.fecha <= hasta)).scalars():
        if not can_see_line(user, r.linea):
            continue
        for c in r.colaboradores or []:
            ini, fin = _dt(r.fecha, c.get("inicio")), _dt(r.fecha, c.get("fin"))
            h = 0.0
            if ini and fin:
                if fin < ini:
                    fin += timedelta(days=1)
                h = max(_min(ini, fin) - float(c.get("refrigerio_min") or 0), 0) / 60.0
            a = acc[c["nombre"].strip().upper()]
            a["horas"] += h
            a["kg"] += float(c.get("kg") or 0)
            a["partes"] += 1
    filas = [{"nombre": n, "horas": round(a["horas"], 2), "kg": round(a["kg"], 1), "partes": a["partes"],
              "kg_hora": round(a["kg"] / a["horas"], 1) if a["horas"] else None} for n, a in acc.items()]
    filas.sort(key=lambda f: -(f["kg_hora"] or 0))
    return {"desde": desde, "hasta": hasta, "filas": filas}


# ------------------------------------------------------------------ GS1
def _gs1_perm(user):
    if not has_perm(user, "comercial.configurar"):
        raise BusinessError("La tabla GS1 ↔ SAP la cargan y modifican el Jefe Comercial y el Gerente Comercial.")


def save_gtin(user, datos):
    """Alta o actualización de un GTIN desde Configuración comercial."""
    _gs1_perm(user)
    g = str(datos.get("gtin") or "").strip()
    code = str(datos.get("product_code") or "").strip().upper()
    if not g.isdigit() or len(g) not in (8, 12, 13, 14):
        raise BusinessError("GTIN inválido: debe tener 8, 12, 13 o 14 dígitos.")
    if not code:
        raise BusinessError("Indique el código SAP del producto.")
    x = db().get(ProductGTIN, g) or ProductGTIN(gtin=g)
    x.product_code, x.tipo = code, f"EAN{len(g)}" if len(g) in (13, 14) else f"GTIN{len(g)}"
    for f in ("marca", "descripcion", "contenido", "unidad"):
        setattr(x, f, str(datos.get(f) or "").strip()[:250])
    db().add(x)
    db().flush()
    audit(user, "GS1_GUARDADO", "ProductGTIN", g, {"codigo": code})
    return x


def delete_gtin(user, gtin):
    _gs1_perm(user)
    x = db().get(ProductGTIN, str(gtin or "").strip())
    if x is None:
        raise BusinessError("GTIN no encontrado.")
    audit(user, "GS1_RETIRADO", "ProductGTIN", x.gtin, {"codigo": x.product_code})
    db().delete(x)
    db().flush()


def import_gtins(user, path):
    """Carga la tabla GS1 ↔ SAP desde Excel: columnas GTIN (o EAN), tipo, código SAP, marca, descripción, envase,
    contenido, unidad. Reemplaza las filas con el mismo GTIN."""
    from openpyxl import load_workbook
    _gs1_perm(user)
    ws = load_workbook(path, data_only=True, read_only=True).active
    filas = list(ws.iter_rows(values_only=True))
    alias = {"gtin": ("gtin", "ean", "codigo gtin", "código gtin", "gtin13", "gtin14", "ean13", "ean14", "codigo de barras",
                      "código de barras"),
             "tipo": ("tipo", "tipo gtin"), "product_code": ("codigo sap", "código sap", "sap", "codigo", "código", "articulo",
                                                             "artículo"),
             "marca": ("marca",), "descripcion": ("descripcion", "descripción"), "envase": ("envase",),
             "contenido": ("contenido",), "unidad": ("unidad", "um")}

    def norm(x):
        return " ".join(str(x or "").strip().lower().split())
    hdr = None
    for i, r in enumerate(filas[:15]):
        cols = [norm(v) for v in r]
        if any(c in alias["gtin"] for c in cols):
            hdr = i
            break
    if hdr is None:
        raise BusinessError("No se encontró la columna GTIN / EAN.")
    cols = [norm(v) for v in filas[hdr]]
    idx = {k: next((j for j, c in enumerate(cols) if c in a), None) for k, a in alias.items()}
    if idx["product_code"] is None:
        raise BusinessError("Falta la columna con el código SAP del producto.")
    n, errores = 0, []
    for k, r in enumerate(filas[hdr + 1:], start=hdr + 2):
        g = str(r[idx["gtin"]] or "").strip().split(".")[0] if idx["gtin"] is not None else ""
        code = str(r[idx["product_code"]] or "").strip().upper()
        if not g and not code:
            continue
        if not g.isdigit() or len(g) not in (8, 12, 13, 14):
            errores.append(f"Fila {k}: GTIN inválido «{g}».")
            continue
        if not code:
            errores.append(f"Fila {k}: falta el código SAP.")
            continue
        x = db().get(ProductGTIN, g) or ProductGTIN(gtin=g)
        x.product_code = code
        x.tipo = (str(r[idx["tipo"]] or "").strip().upper() if idx["tipo"] is not None else "") or f"EAN{len(g)}"
        for f in ("marca", "descripcion", "envase", "contenido", "unidad"):
            if idx[f] is not None:
                setattr(x, f, str(r[idx[f]] or "").strip()[:250])
        db().add(x)
        n += 1
    audit(user, "GS1_CARGADO", "ProductGTIN", "", {"filas": n, "errores": len(errores)})
    return n, errores


# ------------------------------------------------------------------ Control de producción (formato de Producción)
def control_report(periodo="DIA", fecha=None):
    """Reporte diario, semanal o mensual con el formato enviado por Producción: producción real vs plan y capacidad por
    proceso y línea, horas efectivas, productividad (kg/h), asistencia, hallazgos, situación operativa y acciones."""
    from ..models import Attendance, ProductionPlan
    from .calendar import day_factor, holidays
    fecha = fecha or date.today()
    periodo = periodo if periodo in ("DIA", "SEMANA", "MES") else "DIA"
    if periodo == "SEMANA":
        desde = fecha - timedelta(days=fecha.weekday())
        hasta = desde + timedelta(days=6)
    elif periodo == "MES":
        desde = fecha.replace(day=1)
        hasta = (desde + timedelta(days=32)).replace(day=1) - timedelta(days=1)
    else:
        desde = hasta = fecha
    grupos = param("capacidad_lineas", []) or []
    jornada = {int(k): float(v) for k, v in (param("dias_operacion", {}) or {}).items()}
    feriados = holidays()
    dias = [desde + timedelta(days=i) for i in range((hasta - desde).days + 1)]
    factor = sum(day_factor(d, feriados, jornada) for d in dias)
    linea_grupo = {l: g["grupo"] for g in grupos for l in g.get("lineas", [])}
    partes = [r for r in db().execute(select(ProductionReport).where(ProductionReport.fecha >= desde,
                                                                     ProductionReport.fecha <= hasta)).scalars()
              if r.estado not in ("PROGRAMADO", "ANULADO")]
    programados = [r for r in db().execute(select(ProductionReport).where(ProductionReport.fecha >= desde,
                                                                          ProductionReport.fecha <= hasta,
                                                                          ProductionReport.estado == "PROGRAMADO")).scalars()]
    plan = defaultdict(float)
    for pl in db().execute(select(ProductionPlan).where(ProductionPlan.periodo == "DIA", ProductionPlan.fecha >= desde,
                                                        ProductionPlan.fecha <= hasta)).scalars():
        p = db().get(Product, pl.product_code)
        kg = pl.cantidad * ((p.peso_kg if p and p.peso_kg else (1 if p and (p.um or "").upper().startswith("K") else 0)))
        plan[linea_grupo.get(pl.line_code, pl.line_code or "")] += kg
    for r in programados:
        plan[linea_grupo.get(r.linea, r.linea)] += r.peso_entrada_kg or 0
    real, horas, paradas = defaultdict(float), defaultdict(float), defaultdict(list)
    for r in partes:
        g = linea_grupo.get(r.linea, r.linea)
        real[g] += (r.granel_kg if r.tipo == "SELECCION" else (r.granel_kg or 0)) or 0
        horas[g] += (r.indicadores or {}).get("horas_efectivas") or r.horas or 0
        for p in r.pausas or []:
            if p.get("fin"):
                m = _min(datetime.fromisoformat(p["inicio"]), datetime.fromisoformat(p["fin"]))
                paradas[g].append((p["causa"], round(m)))
        if r.parada_min and not r.pausas:
            paradas[g].append((r.motivo_parada or "Parada", r.parada_min))
    filas = []
    procesos = {"SELECCION": "Selección", "ENVASADO": "Envasado"}
    for g in grupos:
        cap = float(g.get("kg_dia") or 0) * factor
        pl = plan.get(g["grupo"], 0.0) or 0.0
        base = pl or cap
        rl = real.get(g["grupo"], 0.0)
        h = horas.get(g["grupo"], 0.0)
        proceso = "Empaque manual" if g["lineas"] == ["ENV_MANUAL"] else procesos.get(g.get("proceso"), g.get("proceso", ""))
        filas.append({"proceso": proceso, "grupo": g["grupo"], "capacidad": round(cap, 1), "plan": round(pl, 1),
                      "plan_es_capacidad": not pl, "real": round(rl, 2), "cumplimiento": round(100 * rl / base, 2) if base else None,
                      "horas": round(h, 1), "kg_h": round(rl / h) if h else 0, "paradas": paradas.get(g["grupo"], [])})
    tot_plan = sum(f["plan"] or f["capacidad"] for f in filas)
    tot_real = sum(f["real"] for f in filas)
    tot_h = sum(f["horas"] for f in filas)
    asis = Counter(a.estado for a in db().execute(select(Attendance).where(Attendance.fecha >= desde,
                                                                           Attendance.fecha <= hasta)).scalars())
    presentes = asis.get("ASISTENCIA", 0) + asis.get("TARDANZA", 0)
    total_asis = sum(asis.values())
    hallazgos, situacion, acciones = [], [], []
    cumpl = round(100 * tot_real / tot_plan, 1) if tot_plan else None
    if cumpl is not None:
        hallazgos.append(f"Cumplimiento general de producción: {cumpl}%.")
    con_plan = [f for f in filas if f["plan"] or f["capacidad"]]
    bajos = sorted([f for f in con_plan if f["cumplimiento"] is not None and f["cumplimiento"] < 90], key=lambda f: f["cumplimiento"])
    if bajos:
        hallazgos.append("Líneas con mayor desviación: " + ", ".join(f"{f['grupo']} ({f['cumplimiento']}%)" for f in bajos[:3]) + ".")
    prod_ord = sorted([f for f in filas if f["kg_h"]], key=lambda f: -f["kg_h"])
    if prod_ord:
        hallazgos.append("Mayor productividad: " + ", ".join(f"{f['grupo']} ({f['kg_h']:,} kg/h)" for f in prod_ord[:2]) + ".")
        if len(prod_ord) > 2:
            hallazgos.append("Baja productividad: " + ", ".join(f"{f['grupo']} ({f['kg_h']:,} kg/h)" for f in prod_ord[-2:]) + ".")
    for f in filas:
        if (f["plan"] or f["capacidad"]) and f["real"] == 0:
            situacion.append(f"{f['grupo']}: sin producción registrada.")
            acciones.append(f"Priorizar la recuperación de {f['grupo']}.")
        for causa, m in f["paradas"][:2]:
            situacion.append(f"{f['grupo']}: parada de {m} min por {causa.lower()}.")
    for f in bajos:
        if f["real"]:
            situacion.append(f"{f['grupo']}: bajo rendimiento, {f['kg_h']:,} kg/h.")
            acciones.append(f"Mejorar el desempeño de {f['grupo']}.")
    altos = [f for f in filas if f["cumplimiento"] is not None and f["cumplimiento"] >= 100]
    if altos:
        acciones.insert(0, "Mantener el ritmo en " + " y ".join(f["grupo"] for f in altos[:3]) + ".")
    acciones.append("Seguimiento diario al cumplimiento del plan.")
    return {"periodo": periodo, "fecha": fecha, "desde": desde, "hasta": hasta, "filas": filas, "total_plan": round(tot_plan, 1),
            "total_real": round(tot_real, 2), "cumplimiento": cumpl, "horas": round(tot_h, 1),
            "productividad": round(tot_real / tot_h) if tot_h else 0, "asistencia": dict(asis), "presentes": presentes,
            "asistencia_pct": round(100 * presentes / total_asis) if total_asis else None, "hallazgos": hallazgos,
            "situacion": situacion or ["Sin paradas registradas en el periodo."], "acciones": acciones}
