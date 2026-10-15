"""Precios de mercado por origen y análisis en dos niveles (observaciones al demo v0.6, Jefe Comercial).

- Registro desde tres lugares: Importaciones (precios de importación), Compras (compras nacionales) y Comercial
  (competencia directa, productos sustitutos y referencias de mercado). Cada origen lo registra el área responsable.
- Importación de informes en Excel o PDF: el sistema extrae las filas cuyo producto es compatible con lo que ofrecemos
  (competencia directa = mismo producto base; sustituto = producto registrado por Comercial en «Productos sustitutos»,
  uno a uno o por Excel) y las muestra para revisión; nada se registra sin confirmación.
- Análisis: nivel 1 «a cuánto compramos» (compras nacionales e importación: evolución y variación) y nivel 2 «a cuánto
  vendemos» (nuestro precio por kg y nuestro volumen frente a competencia y sustitutos).
- El costo de reposición queda pendiente: el usuario indicará su definición.
"""
import json
import os
import re
import uuid
from collections import defaultdict
from datetime import date, timedelta

from sqlalchemy import select

from ..db import db
from ..models import MarketPrice, Product, ProductSubstitute, SalesOrder
from .core import BusinessError, audit, has_perm, param

CATEGORIAS = {
    "COMPRA_NACIONAL": ("Compra nacional", ("compras.gestionar",)),
    "IMPORTACION": ("Importación", ("importaciones.gestionar",)),
    "COMPETENCIA": ("Competencia directa", ("costos.cargar",)),
    "SUSTITUTO": ("Producto sustituto", ("costos.cargar",)),
    "MERCADO": ("Referencia de mercado (mayorista / minorista)", ("costos.cargar",)),
}
NIVEL_COMPRA = ("COMPRA_NACIONAL", "IMPORTACION")
NIVEL_VENTA = ("COMPETENCIA", "SUSTITUTO")


def categorias_para(user):
    return {k: v[0] for k, v in CATEGORIAS.items() if any(has_perm(user, p) for p in v[1])}


def _check(user, categoria):
    if categoria not in CATEGORIAS:
        raise BusinessError("Categoría de precio no válida.")
    if not any(has_perm(user, p) for p in CATEGORIAS[categoria][1]):
        raise BusinessError(f"El precio de «{CATEGORIAS[categoria][0]}» lo registra su área responsable.")


def bases():
    """Productos base que ofrecemos (p. ej. LENTEJA, FRIJOL CANARIO) -> códigos."""
    out = defaultdict(list)
    for p in db().execute(select(Product).where(Product.tipo.in_(("PT", "MP", "GR")))).scalars():
        b = ((p.producto_base or p.descripcion or "").upper().split("*")[0]).strip()
        if b:
            out[b].append(p.code)
    return out


def classify(texto):
    """(categoría, producto base, código compatible) para una línea de un informe, o (None, None, None)."""
    t = " ".join(str(texto or "").upper().split())
    bs = bases()
    directos = [b for b in bs if b and b in t]
    if directos:
        b = max(directos, key=len)
        return "COMPETENCIA", b, bs[b][0]
    for s in db().execute(select(ProductSubstitute).where(ProductSubstitute.activo.is_(True))).scalars():
        if s.sustituto and s.sustituto.upper() in t:
            return "SUSTITUTO", s.producto_base.upper(), (bs.get(s.producto_base.upper()) or [""])[0]
    return None, None, None


def _norm(x):
    return " ".join(str(x or "").upper().split())


def add_substitute(user, producto_base, sustituto, observacion="", origen="Alta manual"):
    """Registra un sustituto (Comercial). El producto base debe ser uno que ofrecemos."""
    if not has_perm(user, "costos.cargar"):
        raise BusinessError("Los productos sustitutos los registra Comercial.")
    base, sus = _norm(producto_base), _norm(sustituto)
    if not base or not sus:
        raise BusinessError("Indique el producto base y el sustituto.")
    if base not in bases():
        raise BusinessError(f"'{base}' no es un producto base que ofrezcamos (revise el producto base en el maestro).")
    if base == sus:
        raise BusinessError("El sustituto no puede ser el mismo producto base.")
    s = db()
    ex = s.execute(select(ProductSubstitute).where(ProductSubstitute.producto_base == base,
                                                   ProductSubstitute.sustituto == sus)).scalars().first()
    if ex is not None:
        ex.activo, ex.observacion = True, (observacion or ex.observacion).strip()
        return ex, False
    x = ProductSubstitute(producto_base=base, sustituto=sus, observacion=(observacion or "").strip(), origen=origen,
                          user_id=user.id)
    s.add(x)
    s.flush()
    audit(user, "SUSTITUTO_REGISTRADO", "ProductSubstitute", x.id, {"base": base, "sustituto": sus, "origen": origen})
    return x, True


def remove_substitute(user, x):
    if not has_perm(user, "costos.cargar"):
        raise BusinessError("Los productos sustitutos los registra Comercial.")
    x.activo = False
    audit(user, "SUSTITUTO_DESACTIVADO", "ProductSubstitute", x.id, {"base": x.producto_base, "sustituto": x.sustituto})


def import_substitutes(user, path, nombre):
    """Excel con columnas «Producto base» y «Sustituto» (y opcional «Observación»). Devuelve (agregados, errores)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        filas = list(wb.worksheets[0].iter_rows(values_only=True))
    except Exception as e:  # archivo dañado o que no es Excel
        raise BusinessError(f"No se pudo leer el Excel: {e}") from None
    if not filas:
        raise BusinessError("El archivo está vacío.")
    cab = [_norm(c) for c in filas[0]]
    try:
        ib = next(i for i, c in enumerate(cab) if "BASE" in c)
        isus = next(i for i, c in enumerate(cab) if "SUSTITUTO" in c)
    except StopIteration:
        raise BusinessError("El archivo debe tener las columnas «Producto base» y «Sustituto».") from None
    iob = next((i for i, c in enumerate(cab) if "OBSERV" in c), None)
    n, errores = 0, []
    for k, r in enumerate(filas[1:], start=2):
        if not r or all(c in (None, "") for c in r):
            continue
        try:
            _, nuevo = add_substitute(user, r[ib], r[isus], r[iob] if iob is not None and iob < len(r) else "", f"Archivo {nombre}")
            n += nuevo
        except BusinessError as e:
            errores.append(f"Fila {k}: {e}")
    return n, errores


_NUM = re.compile(r"(?:S/\.?\s*)?(\d{1,6}(?:[.,]\d{1,4})?)")
_PRES = re.compile(r"(\d+(?:[.,]\d+)?)\s*(KG|KILOS?|G|GR|GRAMOS)\b")


def _kg_de(texto):
    m = _PRES.search(str(texto or "").upper())
    if not m:
        return None
    v = float(m.group(1).replace(",", "."))
    return v if m.group(2).startswith("K") else v / 1000.0


def _row(texto, precio, fuente):
    cat, base, code = classify(texto)
    if cat is None or not precio:
        return None
    kg = _kg_de(texto)
    return {"texto": str(texto)[:200], "categoria": cat, "producto": base, "product_code": code,
            "precio": round(precio, 4), "presentacion_kg": kg,
            "precio_kg": round(precio / kg, 4) if kg else round(precio, 4),
            "unidad_supuesta": "" if kg else "Sin presentación en el texto: se asume precio por kg (revise)", "fuente": fuente}


def parse_report(path, nombre):
    """Extrae del informe (Excel o PDF) las filas compatibles con nuestros productos."""
    ext = os.path.splitext(nombre or path)[1].lower()
    filas = []
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                textos = [str(c) for c in r if isinstance(c, str) and c.strip()]
                nums = [float(c) for c in r if isinstance(c, (int, float)) and not isinstance(c, bool) and c > 0]
                if textos and nums:
                    x = _row(" ".join(textos), nums[-1], f"{nombre} / {ws.title}")
                    if x:
                        filas.append(x)
    elif ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:  # pragma: no cover
            raise BusinessError("No está instalada la lectura de PDF (pypdf).") from e
        reader = PdfReader(path)
        for i, page in enumerate(reader.pages, 1):
            for linea in (page.extract_text() or "").splitlines():
                nums = [float(x.replace(",", ".")) for x in _NUM.findall(linea)]
                nums = [n for n in nums if n > 0]
                if nums:
                    x = _row(linea, nums[-1], f"{nombre} p.{i}")
                    if x:
                        filas.append(x)
    else:
        raise BusinessError("Formato no admitido: cargue un Excel (.xlsx) o un PDF.")
    return filas


def save_preview(upload_dir, filas, meta):
    os.makedirs(upload_dir, exist_ok=True)
    token = uuid.uuid4().hex
    with open(os.path.join(upload_dir, f"mercado_{token}.json"), "w", encoding="utf-8") as fh:
        json.dump({"filas": filas, "meta": meta}, fh, ensure_ascii=False, default=str)
    return token


def load_preview(upload_dir, token):
    if not re.fullmatch(r"[0-9a-f]{32}", token or ""):
        raise BusinessError("Vista previa no válida.")
    p = os.path.join(upload_dir, f"mercado_{token}.json")
    if not os.path.exists(p):
        raise BusinessError("La vista previa ya no existe: vuelva a cargar el informe.")
    with open(p, encoding="utf-8") as fh:
        return json.load(fh)


def add(user, categoria, rows, origen):
    """Registra precios de una categoría (con el permiso del área que corresponde)."""
    _check(user, categoria)
    s = db()
    n = 0
    for r in rows:
        prod = (r.get("producto") or "").strip().upper()
        if not prod or not r.get("precio_kg") or not r.get("fecha"):
            continue
        precio = float(r["precio_kg"])
        if precio <= 0:
            raise BusinessError("El precio por kg debe ser mayor que cero.")
        tipo = "IMPORTACION" if categoria == "IMPORTACION" else (r.get("tipo_precio") or "MAYORISTA").upper()
        s.add(MarketPrice(producto=prod, variedad=(r.get("variedad") or "").upper(), precio_kg=precio, tipo_precio=tipo,
                          mercado=r.get("mercado") or "", fecha=r["fecha"], fuente=r.get("fuente") or "",
                          url=r.get("url") or "", origen=origen, user_id=user.id, categoria=categoria,
                          competidor=(r.get("competidor") or "").strip(), presentacion=(r.get("presentacion") or "").strip(),
                          product_code=(r.get("product_code") or "").strip()))
        n += 1
    s.flush()
    audit(user, "PRECIOS_MERCADO_REGISTRADOS", "MarketPrice", origen, {"categoria": categoria, "filas": n})
    return n


def _serie(rows):
    rows = sorted(rows, key=lambda m: m.fecha)
    return [(m.fecha, m.precio_kg, m) for m in rows]


def _var(serie):
    if len(serie) < 2:
        return None
    a, b = serie[0][1], serie[-1][1]
    return round((b - a) / a * 100, 1) if a else None


def _venta_kg(desde, hasta):
    """Nuestro precio de venta promedio por kg y volumen (kg) por producto base, en dos mitades del periodo."""
    medio = desde + (hasta - desde) / 2
    acc = defaultdict(lambda: {"soles": 0.0, "kg": 0.0, "kg1": 0.0, "kg2": 0.0})
    for o in db().execute(select(SalesOrder).where(SalesOrder.fecha_emision >= desde, SalesOrder.fecha_emision <= hasta,
                                                   SalesOrder.estado != "CANCELADA")).scalars():
        for l in o.lines:
            p = l.product
            if not l.activa or p is None:
                continue
            kg_u = p.peso_kg or (1.0 if (p.um or "").upper().startswith("K") else None)
            if not kg_u:
                continue
            base = ((p.producto_base or p.descripcion or "").upper().split("*")[0]).strip()
            kg = l.cantidad * kg_u
            a = acc[base]
            a["kg"] += kg
            a["soles"] += l.cantidad * (l.precio_unitario or 0)
            a["kg1" if o.fecha_emision <= medio else "kg2"] += kg
    return acc


def analysis(dias=180):
    hasta = date.today()
    desde = hasta - timedelta(days=dias)
    s = db()
    mp = list(s.execute(select(MarketPrice).where(MarketPrice.fecha >= desde)).scalars())
    por = defaultdict(lambda: defaultdict(list))
    for m in mp:
        por[m.categoria or "MERCADO"][m.producto].append(m)
    compra = []
    for prod in sorted(set(por["COMPRA_NACIONAL"]) | set(por["IMPORTACION"])):
        nac, imp = _serie(por["COMPRA_NACIONAL"].get(prod, [])), _serie(por["IMPORTACION"].get(prod, []))
        fila = {"producto": prod, "nacional": nac[-1][1] if nac else None, "importacion": imp[-1][1] if imp else None,
                "var_nacional": _var(nac), "var_importacion": _var(imp),
                "serie_nacional": [(f"{f:%d/%m}", v) for f, v, _ in nac], "serie_importacion": [(f"{f:%d/%m}", v) for f, v, _ in imp]}
        if fila["nacional"] and fila["importacion"]:
            fila["conviene"] = "importar" if fila["importacion"] < fila["nacional"] else "comprar en el país"
        compra.append(fila)
    ventas = _venta_kg(desde, hasta)
    venta = []
    for prod in sorted(set(ventas) | set(por["COMPETENCIA"]) | set(por["SUSTITUTO"])):
        v = ventas.get(prod)
        nuestro = round(v["soles"] / v["kg"], 4) if v and v["kg"] else None
        comp = sorted(por["COMPETENCIA"].get(prod, []), key=lambda m: m.fecha)
        sus = sorted(por["SUSTITUTO"].get(prod, []), key=lambda m: m.fecha)
        ult_comp = {}
        for m in comp:  # último precio de cada competidor
            ult_comp[m.competidor or m.fuente or "competencia"] = m
        prom_comp = round(sum(m.precio_kg for m in ult_comp.values()) / len(ult_comp), 4) if ult_comp else None
        vol_var = round((v["kg2"] - v["kg1"]) / v["kg1"] * 100, 1) if v and v["kg1"] else None
        fila = {"producto": prod, "nuestro_kg": nuestro, "volumen_kg": v["kg"] if v else 0, "var_volumen": vol_var,
                "competidores": [(k, m.precio_kg, m.fecha) for k, m in ult_comp.items()], "prom_competencia": prom_comp,
                "sustitutos": [(m.competidor or m.variedad or m.fuente, m.precio_kg, m.fecha) for m in sus[-3:]],
                "dif_pct": round((nuestro - prom_comp) / prom_comp * 100, 1) if nuestro and prom_comp else None}
        lectura = []
        if fila["dif_pct"] is not None:
            lectura.append(f"vendemos {abs(fila['dif_pct']):.1f} % {'más caro' if fila['dif_pct'] > 0 else 'más barato'} que la competencia")
        if vol_var is not None:
            lectura.append(f"nuestro volumen {'crece' if vol_var > 0 else 'cae'} {abs(vol_var):.1f} %")
        if fila["dif_pct"] is not None and vol_var is not None:
            if fila["dif_pct"] > 0 and vol_var < 0:
                lectura.append("la competencia estaría ganando mercado")
            elif vol_var > 0:
                lectura.append("estamos ganando volumen")
        fila["lectura"] = "; ".join(lectura) or "Sin datos suficientes para comparar"
        venta.append(fila)
    return {"desde": desde, "hasta": hasta, "compra": compra, "venta": venta}
