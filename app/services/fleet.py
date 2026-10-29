"""Seguimiento de flota (observaciones al demo v0.6, Coordinador de Distribución).

SOAT, revisión técnica, impuesto vehicular, multas y otros documentos por vehículo con alerta de vencimiento; consumo
de combustible (rendimiento km por galón entre cargas); mantenimiento con alerta por kilometraje o fecha; y, ante una
multa o incidencia, qué conductor tenía el vehículo y qué ruta siguió ese día.
"""
from datetime import date, timedelta

from sqlalchemy import select

from ..db import db
from ..models import Driver, FuelLog, Route, Vehicle, VehicleDocument, VehicleMaintenance
from . import alerts
from .core import BusinessError, audit, has_perm, param

TIPOS_DOC = {"SOAT": "SOAT", "TARJETA_CIRCULACION": "Tarjeta de circulación", "REVISION_TECNICA": "Revisión técnica",
             "IMPUESTO_VEHICULAR": "Impuesto vehicular", "MULTA": "Multa / papeleta", "OTRO": "Otro"}


def _req(user):
    if not has_perm(user, "distribucion.flota"):
        raise BusinessError("La flota la mantiene Distribución.")


def _km(vehicle, km):
    if km is not None and (vehicle.km_actual is None or km > vehicle.km_actual):
        vehicle.km_actual = km


def who_had(vehicle, fecha):
    """Rutas del vehículo en la fecha: conductor y pedidos/zonas recorridas."""
    rutas = list(db().execute(select(Route).where(Route.vehicle_id == vehicle.id, Route.fecha == fecha,
                                                  Route.estado != "ANULADA")).scalars())
    return [{"route": r, "conductor": r.driver.nombre if r.driver else "-", "zona": r.zona,
             "paradas": [f"OV {st.order.numero} · {st.order.customer.nombre} · {st.order.distrito}" for st in r.stops],
             "salida": r.salida_ts, "retorno": r.retorno_ts} for r in rutas]


def add_document(user, vehicle, tipo, numero="", fecha=None, vencimiento=None, monto=None, estado="", observacion="",
                 archivo="", hora=""):
    _req(user)
    if tipo not in TIPOS_DOC:
        raise BusinessError("Tipo de documento no válido.")
    if tipo == "REVISION_TECNICA" and vencimiento is None and fecha is not None:
        vencimiento = fecha + timedelta(days=365)  # V12: la revisión técnica es anual
    if tipo != "MULTA" and vencimiento is None:
        raise BusinessError("Indique la fecha de vencimiento.")
    if tipo == "MULTA" and fecha is None:
        raise BusinessError("Indique la fecha de la multa.")
    d = VehicleDocument(vehicle_id=vehicle.id, tipo=tipo, numero=(numero or "").strip(), fecha=fecha, hora=(hora or "").strip(),
                        vencimiento=vencimiento, monto=monto, estado=estado or ("PENDIENTE" if tipo == "MULTA" else "VIGENTE"),
                        observacion=(observacion or "").strip(), archivo=archivo or "", user_id=user.id)
    if tipo == "MULTA":  # quién tenía el vehículo y qué ruta siguió
        rutas = who_had(vehicle, fecha)
        if rutas:
            d.route_id = rutas[0]["route"].id
            d.driver_id = rutas[0]["route"].driver_id
    db().add(d)
    db().flush()
    audit(user, "FLOTA_DOCUMENTO", "Vehicle", vehicle.placa, {"tipo": tipo, "numero": d.numero, "vencimiento": str(vencimiento)})
    return d


def add_fuel(user, vehicle, fecha, km, galones, monto=None, route_id=None, driver_id=None):
    if not (has_perm(user, "distribucion.flota") or has_perm(user, "distribucion.entrega")):
        raise BusinessError("El consumo de combustible lo registra Distribución o el conductor.")
    if not km or not galones or galones <= 0:
        raise BusinessError("Indique kilometraje y galones.")
    if vehicle.km_actual and km < vehicle.km_actual - 1e-9:
        raise BusinessError(f"El kilometraje ({km:g}) es menor al último registrado ({vehicle.km_actual:g}).")
    f = FuelLog(vehicle_id=vehicle.id, fecha=fecha or date.today(), km=km, galones=galones, monto=monto,
                route_id=route_id, driver_id=driver_id, user_id=user.id)
    db().add(f)
    _km(vehicle, km)
    db().flush()
    audit(user, "FLOTA_COMBUSTIBLE", "Vehicle", vehicle.placa, {"km": km, "galones": galones, "monto": monto})
    return f


def add_maintenance(user, vehicle, fecha, tipo, km=None, descripcion="", costo=None, proximo_km=None, proxima_fecha=None):
    _req(user)
    if tipo not in ("PREVENTIVO", "CORRECTIVO"):
        raise BusinessError("Tipo de mantenimiento no válido.")
    m = VehicleMaintenance(vehicle_id=vehicle.id, fecha=fecha or date.today(), tipo=tipo, km=km, descripcion=descripcion,
                           costo=costo, proximo_km=proximo_km, proxima_fecha=proxima_fecha, user_id=user.id)
    db().add(m)
    _km(vehicle, km)
    db().flush()
    alerts.close_by_system(f"FLOTA_MANT:{vehicle.id}", "Mantenimiento registrado")
    audit(user, "FLOTA_MANTENIMIENTO", "Vehicle", vehicle.placa, {"tipo": tipo, "km": km, "proximo_km": proximo_km})
    return m


def fuel_efficiency(vehicle):
    """Rendimiento (km/galón) entre cargas consecutivas: galones de la carga posterior cubren los km recorridos."""
    logs = list(db().execute(select(FuelLog).where(FuelLog.vehicle_id == vehicle.id).order_by(FuelLog.km)).scalars())
    out = []
    for a, b in zip(logs, logs[1:]):
        km = b.km - a.km
        out.append({"desde": a.fecha, "hasta": b.fecha, "km": km, "galones": b.galones,
                    "km_galon": round(km / b.galones, 2) if b.galones else None,
                    "costo_km": round(b.monto / km, 3) if b.monto and km else None})
    return out


def status(vehicle, hoy=None):
    hoy = hoy or date.today()
    s = db()
    docs = list(s.execute(select(VehicleDocument).where(VehicleDocument.vehicle_id == vehicle.id)
                          .order_by(VehicleDocument.id.desc())).scalars())  # V12: rige el último registrado de cada tipo
    ult = {}
    for d in docs:
        if d.tipo != "MULTA" and d.tipo not in ult:
            ult[d.tipo] = d
    mant = s.execute(select(VehicleMaintenance).where(VehicleMaintenance.vehicle_id == vehicle.id)
                     .order_by(VehicleMaintenance.fecha.desc(), VehicleMaintenance.id.desc())).scalars().first()
    ef = fuel_efficiency(vehicle)
    return {"documentos": ult, "multas": [d for d in docs if d.tipo == "MULTA"], "mantenimiento": mant,
            "proximo_km": next_maintenance_km(vehicle, mant if mant and mant.tipo == "PREVENTIVO" else None),
            "km_restantes": (next_maintenance_km(vehicle, mant if mant and mant.tipo == "PREVENTIVO" else None) - vehicle.km_actual)
            if vehicle.km_actual is not None and next_maintenance_km(vehicle, mant if mant and mant.tipo == "PREVENTIVO" else None) else None,
            "rendimiento": ef[-1]["km_galon"] if ef else None, "consumos": ef[-10:],
            "vencidos": [d for d in ult.values() if d.vencimiento and d.vencimiento < hoy]}


def _aviso_dias(tipo):
    cfg = param("aviso_documentos_vehiculo_dias", {}) or {}
    return int(cfg.get(tipo, cfg.get("*", param("dias_aviso_documentos_vehiculo", 15) or 15)))


def next_maintenance_km(vehicle, ultimo=None):
    """V12: próximo mantenimiento preventivo por kilometraje (cada «km_mantenimiento_preventivo» km)."""
    cada = float(param("km_mantenimiento_preventivo", 5000) or 0)
    if ultimo is not None and ultimo.proximo_km:
        return ultimo.proximo_km
    if not cada:
        return None
    if ultimo is not None and ultimo.km is not None:
        return ultimo.km + cada
    if vehicle.km_actual is not None:
        return (int(vehicle.km_actual // cada) + 1) * cada
    return None


def check_alerts(hoy=None):
    """Tarea programada (V12): documentos con aviso por tipo (SOAT 2 semanas, tarjeta de circulación y revisión técnica
    1 mes), recordatorio un día antes y el mismo día; mantenimiento cada 5 000 km con aviso 500 km antes."""
    hoy = hoy or date.today()
    margen = float(param("margen_km_mantenimiento", 500) or 500)
    n = 0
    for v in db().execute(select(Vehicle)).scalars():
        st = status(v, hoy)
        for d in st["documentos"].values():
            if not d.vencimiento:
                continue
            falta = (d.vencimiento - hoy).days
            if falta < 0:
                etapa, txt, prio = "V", "vencido", "ALTA"
            elif falta == 0:
                etapa, txt, prio = "D0", "vence hoy", "ALTA"
            elif falta == 1:
                etapa, txt, prio = "D1", "vence mañana", "ALTA"
            elif falta <= _aviso_dias(d.tipo):
                etapa, txt, prio = "P", "por vencer", None
            else:
                continue
            alerts.emit("FLOTA_DOCUMENTO", f"{TIPOS_DOC.get(d.tipo, d.tipo)} del vehículo {v.placa} {txt} ({d.vencimiento:%d/%m/%Y})",
                        f"Documento {d.numero or '-'}.", "Vehicle", v.id, prioridad=prio,
                        dedup_key=f"FLOTA_DOC:{d.id}:{etapa}")
            n += 1
        m = st["mantenimiento"]
        prox = next_maintenance_km(v, m if m and m.tipo == "PREVENTIVO" else None)
        km_rest = prox - v.km_actual if prox is not None and v.km_actual is not None else None
        if (km_rest is not None and km_rest <= margen) or (m and m.proxima_fecha and m.proxima_fecha <= hoy + timedelta(days=_aviso_dias("*"))):
            alerts.emit("FLOTA_MANTENIMIENTO", f"Mantenimiento {'vencido' if km_rest is not None and km_rest <= 0 else 'próximo'} "
                                               f"del vehículo {v.placa}",
                        f"Km actual {v.km_actual or '-'}; mantenimiento a los {prox or '-'} km"
                        f"{' o el ' + m.proxima_fecha.strftime('%d/%m/%Y') if m and m.proxima_fecha else ''}.", "Vehicle", v.id,
                        prioridad="ALTA" if km_rest is not None and km_rest <= 0 else None,
                        dedup_key=f"FLOTA_MANT:{v.id}:{int(prox or 0)}")
            n += 1
    return n


def soat_status(vehicle, fecha=None):
    """Estado de la placa por su SOAT: (vigente, texto). Sin SOAT registrado o vencido a la fecha → no vigente."""
    fecha = fecha or date.today()
    d = db().execute(select(VehicleDocument).where(VehicleDocument.vehicle_id == vehicle.id, VehicleDocument.tipo == "SOAT")
                     .order_by(VehicleDocument.vencimiento.desc().nullslast())).scalars().first()
    if d is None or d.vencimiento is None:
        return False, "sin SOAT registrado"
    if d.vencimiento < fecha:
        return False, f"SOAT vencido el {d.vencimiento:%d/%m/%Y}"
    return True, f"SOAT vigente hasta el {d.vencimiento:%d/%m/%Y}"


def check_soat(vehicle, fecha=None):
    """Impide asignar una unidad sin SOAT vigente (parámetro exigir_soat_vigente)."""
    if not param("exigir_soat_vigente", True):
        return
    ok, txt = soat_status(vehicle, fecha)
    if not ok:
        alerts.emit("FLOTA_DOCUMENTO", f"Unidad {vehicle.placa} sin SOAT vigente: no se asigna a rutas",
                    f"{txt[0].upper() + txt[1:]}. Registre el SOAT en Flota → {vehicle.placa}.", "Vehicle", vehicle.id,
                    prioridad="ALTA", dedup_key=f"SOAT_BLOQUEO:{vehicle.id}")
        raise BusinessError(f"La unidad {vehicle.placa} no tiene SOAT vigente ({txt}): no se puede asignar.")


def drivers():
    return list(db().execute(select(Driver).order_by(Driver.nombre)).scalars())


# ------------------------------------------------------------------ V8: carga de unidades (formato GEOSATELITAL)
COLS_UNIDADES = {"placa": ("placa",), "nombre": ("nombre",), "conductor": ("conductor",), "flota": ("flota",),
                 "marca": ("marca",), "modelo": ("modelo",), "anio": ("año", "ano", "anio"), "tipo": ("tipo",),
                 "odometro": ("odometro", "odómetro", "kilometraje", "km"), "capacidad": ("capacidad (kg)", "capacidad kg",
                                                                                     "capacidad")}
SIN_DATO = {"", "SIN MARCA", "SIN MODELO", "N/A", "NA", "-", "NONE"}


def _limpio(v):
    t = " ".join(str(v if v is not None else "").split())
    return "" if t.upper() in SIN_DATO else t


def import_vehicles(user, path):
    """Unidades desde el Excel o CSV que exporta GEOSATELITAL (Nombre, Conductor, Flota, Placa, Marca, Modelo, Año, Tipo,
    Odómetro) u otro con esas columnas. Se identifica por la placa: si existe se actualiza, si no se crea. La capacidad
    no viene en GEOSATELITAL: si falta, la unidad queda «no disponible» hasta registrarla."""
    import csv
    import unicodedata
    _req(user)
    if str(path).lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as fh:
            filas = list(csv.reader(fh))
    else:
        from openpyxl import load_workbook
        filas = [list(r) for r in load_workbook(path, read_only=True, data_only=True).active.iter_rows(values_only=True)]
    if not filas:
        raise BusinessError("El archivo está vacío.")

    def norm(x):
        x = unicodedata.normalize("NFKD", str(x or "")).encode("ascii", "ignore").decode().lower()
        return " ".join(x.split())
    cab = [norm(c) for c in filas[0]]
    idx = {}
    for k, alias in COLS_UNIDADES.items():
        for a in alias:
            if norm(a) in cab:
                idx[k] = cab.index(norm(a))
                break
    if "placa" not in idx:
        raise BusinessError("Falta la columna «Placa».")
    nuevas = actualizadas = 0
    errores = []
    for n, r in enumerate(filas[1:], start=2):
        def val(k):
            return r[idx[k]] if k in idx and idx[k] < len(r) else None
        placa = _limpio(val("placa")).upper().replace(" ", "")
        if not placa:
            continue
        if not (5 <= len(placa.replace("-", "")) <= 8):
            errores.append(f"Fila {n}: placa «{placa}» inválida.")
            continue
        v = db().execute(select(Vehicle).where(Vehicle.placa == placa)).scalars().first()
        if v is None:
            v = Vehicle(placa=placa, capacidad_kg=0, propio=True, disponible=False)
            db().add(v)
            nuevas += 1
        else:
            actualizadas += 1
        for campo, k in (("nombre_geosat", "nombre"), ("flota", "flota"), ("marca", "marca"), ("modelo", "modelo"),
                         ("tipo_unidad", "tipo")):
            t = _limpio(val(k))
            if t:
                setattr(v, campo, t[:60].upper() if campo in ("marca", "modelo") else t[:60])
        try:
            a = _limpio(val("anio"))
            v.anio = int(float(a)) if a else v.anio
            km = _limpio(val("odometro"))
            if km and float(km) > 0 and (v.km_actual is None or float(km) >= v.km_actual):
                v.km_actual = float(km)
            cap = _limpio(val("capacidad"))
            if cap and float(cap) > 0:
                v.capacidad_kg = float(cap)
        except ValueError:
            errores.append(f"Fila {n}: año, odómetro o capacidad no numérico.")
        if not v.descripcion:
            v.descripcion = " ".join(x for x in (v.tipo_unidad, v.marca, v.modelo) if x)[:80]
        if v.capacidad_kg <= 0:
            v.disponible = False
    db().flush()
    audit(user, "FLOTA_CARGADA", "Vehicle", "", {"nuevas": nuevas, "actualizadas": actualizadas, "errores": len(errores)})
    return {"nuevas": nuevas, "actualizadas": actualizadas, "errores": errores}
