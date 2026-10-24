"""Reempaque (IF §9): Almacén solicita -> Calidad evalúa -> notifica a Almacén y Producción -> Producción ejecuta.
Q12 (12-oct-2026): el sistema puede generar una pre-solicitud; el Coordinador de Almacén la aprueba y la envía a Calidad."""
from ..db import db
from ..models import Product, RepackRequest, StockLot, now
from . import alerts
from .catalogs_helpers import repack_type_name
from .core import BusinessError, audit, check_warehouse_scope, has_perm
from sqlalchemy import select

from .stock import add_or_increase_lot, change_lot_state, free_qty


def _segment(lot, estado):
    if lot.estado == estado:
        return lot
    return db().execute(select(StockLot).where(StockLot.product_code == lot.product_code,
                                               StockLot.warehouse_code == lot.warehouse_code, StockLot.lote == lot.lote,
                                               StockLot.estado == estado)).scalars().first() or lot

TIPOS = ("ACTUALIZACION_FECHA", "CAMBIO_PRESENTACION", "CONVERSION_CALIDAD")


def request(user, tipo, lot_id, cantidad, producto_destino=None, cantidad_destino=None, order=None, motivo="",
            presolicitud=False):
    """presolicitud=True: la genera el sistema (user=None) y queda PRE_SOLICITUD hasta que el Coordinador de Almacén
    la apruebe (Q12). Si la crea el propio Coordinador, se envía directamente a Calidad."""
    if not presolicitud and not has_perm(user, "reempaque.solicitar"):
        raise BusinessError("La solicitud de reempaque la inicia Almacén (IF §9.1).")
    if tipo not in TIPOS:
        raise BusinessError("Tipo de reempaque no válido.")
    lot = db().get(StockLot, int(lot_id))
    if lot is None:
        raise BusinessError("Lote no encontrado.")
    if lot.estado not in ("DISPONIBLE", "PENDIENTE_REVISION"):
        raise BusinessError(f"El lote está {lot.estado}; no puede seleccionarse para reempaque.")
    libre = free_qty(lot) if lot.estado == "DISPONIBLE" else lot.cantidad
    if cantidad <= 0 or cantidad > libre + 1e-9:
        raise BusinessError(f"Cantidad inválida (libre en el lote: {libre:g}).")
    if tipo in ("CAMBIO_PRESENTACION", "CONVERSION_CALIDAD"):
        if not producto_destino or db().get(Product, producto_destino) is None:
            raise BusinessError("Indique el producto destino (SKU resultante).")
    if not presolicitud:
        check_warehouse_scope(user, lot.warehouse_code)
    r = RepackRequest(tipo=tipo, lot_id=lot.id, cantidad=cantidad, producto_destino=producto_destino or lot.product_code,
                      cantidad_destino=cantidad_destino, order_id=order.id if order else None, motivo=motivo,
                      solicitado_por=user.id if user else None, estado="PRE_SOLICITUD" if presolicitud else "SOLICITADA")
    db().add(r)
    db().flush()
    if presolicitud:
        alerts.emit("REEMPAQUE_PRESOLICITUD", f"Aprobar pre-solicitud de reempaque #{r.id}: {repack_type_name(tipo)}",
                    f"Lote {lot.lote} ({lot.product_code}) × {cantidad:g} en {lot.warehouse_code}. Destino: {r.producto_destino}. "
                    f"{motivo}", "RepackRequest", r.id, dedup_key=f"REEMPAQUE_PRE:{r.id}", creado_por=user)
        audit(user, "REEMPAQUE_PRESOLICITADO", "RepackRequest", r.id, {"tipo": tipo, "lote": lot.lote, "cantidad": cantidad})
        return r
    _send_to_quality(r, lot, user)
    return r


def approve_presolicitud(r, user, aprobar, comentario=""):
    """Q12: el Coordinador de Almacén aprueba la pre-solicitud y la envía a Calidad (o la desestima)."""
    if not has_perm(user, "reempaque.solicitar"):
        raise BusinessError("La pre-solicitud de reempaque la aprueba el Coordinador de Almacén (Q12).")
    if r.estado != "PRE_SOLICITUD":
        raise BusinessError("La solicitud no es una pre-solicitud pendiente.")
    lot = r.lot
    check_warehouse_scope(user, lot.warehouse_code)
    alerts.close_by_system(f"REEMPAQUE_PRE:{r.id}", "Aprobada" if aprobar else "Desestimada")
    if not aprobar:
        if not comentario.strip():
            raise BusinessError("Indique el motivo para desestimar.")
        r.estado = "DESESTIMADA"
        r.motivo = (r.motivo + "\n" if r.motivo else "") + f"Desestimada: {comentario}"
        audit(user, "REEMPAQUE_PRESOLICITUD_DESESTIMADA", "RepackRequest", r.id, {"comentario": comentario})
        return r
    libre = free_qty(lot) if lot.estado == "DISPONIBLE" else lot.cantidad
    if r.cantidad > libre + 1e-9:
        raise BusinessError(f"El lote ya no tiene la cantidad libre ({libre:g}).")
    r.estado = "SOLICITADA"
    r.solicitado_por = user.id
    if comentario:
        r.motivo = (r.motivo + "\n" if r.motivo else "") + comentario
    _send_to_quality(r, lot, user)
    return r


def _send_to_quality(r, lot, user):
    tipo, cantidad, motivo = r.tipo, r.cantidad, r.motivo
    # El producto seleccionado para reempaque queda 'Pendiente de revisión' (IF §6.1). Si se selecciona todo el
    # lote cambia su estado; si es una parte, esa cantidad queda retenida (no reservable) hasta la evaluación.
    db().flush()
    if lot.estado == "DISPONIBLE" and abs(cantidad - lot.cantidad) < 1e-9:
        change_lot_state(lot, "PENDIENTE_REVISION", user, f"Seleccionado para reempaque ({repack_type_name(tipo)})")
        r.lot_id = _segment(lot, "PENDIENTE_REVISION").id
    alerts.emit("REEMPAQUE_SOLICITUD", f"Evaluar reempaque #{r.id}: {repack_type_name(tipo)}",
                f"Lote {lot.lote} ({lot.product_code}) × {cantidad:g}. Destino: {r.producto_destino}. Motivo: {motivo}",
                "RepackRequest", r.id, dedup_key=f"REEMPAQUE_SOL:{r.id}", creado_por=user)
    audit(user, "REEMPAQUE_SOLICITADO", "RepackRequest", r.id, {"tipo": tipo, "lote": lot.lote, "cantidad": cantidad})
    return r


def evaluate(r, user, aprobar, resultado):
    if not has_perm(user, "reempaque.evaluar"):
        raise BusinessError("Calidad evalúa y aprueba el reempaque (IF §9.1).")
    if r.estado != "SOLICITADA":
        raise BusinessError("La solicitud no está pendiente de evaluación.")
    if not resultado.strip():
        raise BusinessError("Registre el resultado de la evaluación/muestreo.")
    r.estado = "APROBADA" if aprobar else "DESAPROBADA"
    r.evaluado_por = user.id
    r.evaluado_ts = now()
    r.resultado_calidad = resultado
    alerts.close_by_system(f"REEMPAQUE_SOL:{r.id}", f"Evaluado: {r.estado}")
    lot = r.lot
    if aprobar:
        alerts.emit("REEMPAQUE_APROBADO", f"Reempaque #{r.id} aprobado: ejecutar",
                    f"{repack_type_name(r.tipo)}. Lote {lot.lote} × {r.cantidad:g} → {r.producto_destino}. {resultado}",
                    "RepackRequest", r.id, creado_por=user)
    else:
        if lot.estado == "PENDIENTE_REVISION" and lot.estado_anterior:
            change_lot_state(lot, lot.estado_anterior, user, f"Reempaque #{r.id} desaprobado")
        alerts.emit("REEMPAQUE_DESAPROBADO", f"Reempaque #{r.id} desaprobado", resultado, "RepackRequest", r.id,
                    creado_por=user)
    audit(user, f"REEMPAQUE_{r.estado}", "RepackRequest", r.id, {"resultado": resultado})


def start(r, user):
    if not has_perm(user, "reempaque.ejecutar"):
        raise BusinessError("Producción ejecuta el reempaque (IF §9.1).")
    if r.estado != "APROBADA":
        raise BusinessError("Solo se ejecutan reempaques aprobados.")
    r.estado = "EN_EJECUCION"
    r.ejecucion_inicio = now()
    r.ejecutado_por = user.id
    audit(user, "REEMPAQUE_INICIO", "RepackRequest", r.id)


def finish(r, user, lote_resultante, cantidad_resultante, nueva_fecha_fab=None, nueva_fecha_venc=None):
    if not has_perm(user, "reempaque.ejecutar"):
        raise BusinessError("Producción registra el reempaque.")
    if r.estado != "EN_EJECUCION":
        raise BusinessError("El reempaque no está en ejecución.")
    if not lote_resultante or cantidad_resultante is None or cantidad_resultante <= 0:
        raise BusinessError("Indique lote y cantidad resultante.")
    lot = r.lot
    if r.cantidad > lot.cantidad + 1e-9:
        raise BusinessError("El lote origen ya no tiene la cantidad del reempaque.")
    lot.cantidad -= r.cantidad
    if lot.cantidad_verificada is not None:
        lot.cantidad_verificada = max(lot.cantidad_verificada - r.cantidad, 0)
    if lot.estado == "PENDIENTE_REVISION" and lot.estado_anterior and lot.cantidad > 1e-9:
        # saldo remanente vuelve a su estado anterior
        change_lot_state(lot, lot.estado_anterior, user, f"Saldo de reempaque #{r.id}")
    nuevo = add_or_increase_lot(r.producto_destino, lot.warehouse_code, lote_resultante, cantidad_resultante,
                                "PENDIENTE_REVISION", nueva_fecha_venc or lot.fecha_venc, nueva_fecha_fab or lot.fecha_fab,
                                origen=f"REEMPAQUE-{r.id}")
    r.estado = "EJECUTADA"
    r.ejecucion_fin = now()
    r.lote_resultante = lote_resultante
    r.cantidad_destino = cantidad_resultante
    r.nueva_fecha_fab = nueva_fecha_fab
    r.nueva_fecha_venc = nueva_fecha_venc
    r.lot_resultante_id = nuevo.id
    alerts.emit("STOCK_PENDIENTE_REVISION", f"Lote reempacado {lote_resultante} pendiente de liberación",
                f"{r.producto_destino} × {cantidad_resultante:g} (reempaque #{r.id}).", "StockLot", nuevo.id, creado_por=user)
    audit(user, "REEMPAQUE_EJECUTADO", "RepackRequest", r.id, {"lote": lote_resultante, "cantidad": cantidad_resultante})
    return nuevo


def import_requests(user, path):
    """Carga de solicitudes de reempaque desde Excel (ronda UX 24-oct-2026). Columnas: tipo, producto, lote, almacén,
    cantidad y, según el tipo, producto destino; motivo opcional. Cada fila válida se envía a Calidad como una solicitud."""
    from openpyxl import load_workbook
    if not has_perm(user, "reempaque.solicitar"):
        raise BusinessError("La solicitud de reempaque la inicia Almacén (IF §9.1).")
    ws = load_workbook(path, data_only=True, read_only=True).active
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        raise BusinessError("El archivo está vacío.")
    cab = [" ".join(str(c or "").strip().lower().split()) for c in filas[0]]

    def col(*n):
        return next((i for i, c in enumerate(cab) if c in n), None)
    it, ip, il, ia = col("tipo"), col("producto", "codigo", "código", "articulo", "artículo"), col("lote"), col("almacen", "almacén")
    iq, idest, im = col("cantidad"), col("producto destino", "destino", "sku destino"), col("motivo", "observacion", "observación")
    if None in (it, ip, il, iq):
        raise BusinessError("Faltan columnas: tipo, producto, lote y cantidad.")
    alias = {"ACTUALIZACION DE FECHA": "ACTUALIZACION_FECHA", "ACTUALIZACIÓN DE FECHA": "ACTUALIZACION_FECHA",
             "CAMBIO DE PRESENTACION": "CAMBIO_PRESENTACION", "CAMBIO DE PRESENTACIÓN": "CAMBIO_PRESENTACION",
             "CONVERSION DE CALIDAD": "CONVERSION_CALIDAD", "CONVERSIÓN DE CALIDAD": "CONVERSION_CALIDAD"}
    creadas, errores = [], []
    for k, r in enumerate(filas[1:], start=2):
        if not any(r):
            continue
        try:
            tipo = str(r[it] or "").strip().upper()
            tipo = alias.get(tipo, tipo.replace(" ", "_"))
            code, lote = str(r[ip] or "").strip().upper(), str(r[il] or "").strip()
            q = select(StockLot).where(StockLot.product_code == code, StockLot.lote == lote, StockLot.cantidad > 0)
            if ia is not None and r[ia]:
                q = q.where(StockLot.warehouse_code == str(r[ia]).strip().upper())
            lots = list(db().execute(q).scalars())
            if not lots:
                raise BusinessError(f"no hay stock del lote {lote} de {code}")
            if len(lots) > 1:
                raise BusinessError(f"el lote {lote} está en varios almacenes; indique el almacén")
            dest = str(r[idest] or "").strip().upper() if idest is not None else ""
            creadas.append(request(user, tipo, lots[0].id, float(r[iq]), dest or None, None, None,
                                   str(r[im] or "").strip() if im is not None else ""))
        except (BusinessError, TypeError, ValueError) as e:
            errores.append(f"Fila {k}: {e}")
    audit(user, "REEMPAQUE_CARGA", "RepackRequest", "", {"creadas": len(creadas), "errores": len(errores)})
    return creadas, errores
