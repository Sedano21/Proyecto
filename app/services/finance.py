"""Finanzas (ronda 23-oct-2026): condición de pago del cliente aprobada por la Coordinación de Finanzas, facturas y notas
de crédito (carga Excel, PDF consolidado o individual, o registro manual), comprobantes de pago que sube Comercial y
confirma Finanzas, deuda y estado de cuenta en el archivo del cliente, y cuadro de cobranza."""
import os
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta

from sqlalchemy import select

from ..db import db
from ..models import Customer, CustomerInvoice, CustomerPayment, OrderAttachment, SalesOrder, now
from . import alerts
from .core import BusinessError, audit, has_perm, param

CONDICIONES = {"CONTADO": "Contado", "CREDITO": "Crédito"}
ESTADOS_CONDICION = {"SIN_SOLICITAR": "Sin solicitar", "PENDIENTE": "Pendiente de Finanzas", "APROBADA": "Aprobada",
                     "OBSERVADA": "Observada", "RECHAZADA": "Rechazada"}
TIPOS_DOC_OV = {"OC": "Orden de compra", "DEVOLUCION": "Constancia de devolución", "NOTA_CREDITO": "Nota de crédito",
                "OTRO": "Otro documento"}
EXT_DOCS = (".pdf", ".jpg", ".jpeg", ".png", ".webp", ".heic", ".xlsx", ".xls", ".docx", ".doc", ".msg", ".eml")


def save_document(file_storage, upload_dir, prefix, exts=EXT_DOCS):
    """Guarda un documento adjunto (orden de compra, plantilla de pedido, comprobante, factura…)."""
    if file_storage is None or not getattr(file_storage, "filename", ""):
        return ""
    ext = os.path.splitext(file_storage.filename)[1].lower()
    if ext not in exts:
        raise BusinessError(f"Formato no permitido ({ext}). Use: {', '.join(e.strip('.') for e in exts)}.")
    os.makedirs(upload_dir, exist_ok=True)
    name = f"{prefix}_{uuid.uuid4().hex[:10]}{ext}"
    file_storage.save(os.path.join(upload_dir, name))
    return name


# ------------------------------------------------------------------ condición de pago
def credit_rules():
    return param("credito_tramos", {}) or {}


def required_docs(condicion, monto=None, dias=None):
    """Documentos que pide Finanzas: ninguno al contado; por tramo de monto en crédito; todos a 90 o 120 días."""
    if condicion != "CREDITO":
        return []
    r = credit_rules()
    todos = list((r.get("documentos") or {}).keys())
    if dias and int(dias) in [int(x) for x in r.get("plazos_todos_los_documentos") or []]:
        return todos
    for t in r.get("tramos") or []:
        if t.get("hasta") is None or (monto or 0) <= float(t["hasta"]):
            return list(t.get("docs") or [])
    return todos


def max_days(monto):
    for t in credit_rules().get("tramos") or []:
        if t.get("hasta") is None or (monto or 0) <= float(t["hasta"]):
            return t.get("dias_max")
    return None


def request_condition(customer, user, condicion, monto=None, dias=None, archivos=None, comentario=""):
    """Comercial envía la condición de pago a Finanzas con los documentos del tramo."""
    if not has_perm(user, "condicion.solicitar"):
        raise BusinessError("Solicita la condición de pago el área Comercial.")
    condicion = (condicion or "").upper()
    if condicion not in CONDICIONES:
        raise BusinessError("Elija la condición: contado o crédito.")
    if condicion == "CREDITO":
        if not monto or float(monto) <= 0:
            raise BusinessError("Indique el monto de la línea de crédito (S/).")
        if not dias or int(dias) <= 0:
            raise BusinessError("Indique el plazo de crédito en días.")
        dias = int(dias)
        todos = [int(x) for x in credit_rules().get("plazos_todos_los_documentos") or []]
        tope = max_days(float(monto))
        if tope and dias > int(tope) and dias not in todos:
            raise BusinessError(f"Para S/ {float(monto):,.0f} el plazo máximo es {tope} días "
                                f"(los plazos de {', '.join(map(str, todos))} días piden todos los documentos).")
    docs = dict(customer.condicion_docs or {})
    docs.update({k: v for k, v in (archivos or {}).items() if v})
    nombres = credit_rules().get("documentos") or {}
    falta = [nombres.get(d, d) for d in required_docs(condicion, float(monto or 0), dias) if not docs.get(d)]
    if falta:
        raise BusinessError("Faltan documentos para Finanzas: " + "; ".join(falta) + ".")
    customer.condicion_pago = condicion
    customer.credito_monto = float(monto) if condicion == "CREDITO" else None
    customer.credito_dias = int(dias) if condicion == "CREDITO" else None
    customer.condicion_docs = docs
    customer.condicion_estado = "PENDIENTE"
    customer.condicion_historial = list(customer.condicion_historial or []) + [
        {"accion": "SOLICITADA", "user": user.nombre, "ts": now().isoformat(timespec="minutes"), "condicion": condicion,
         "monto": customer.credito_monto, "dias": customer.credito_dias, "comentario": (comentario or "").strip()}]
    alerts.emit("CONDICION_PAGO_SOLICITADA", f"Aprobar condición de pago: {customer.nombre}",
                f"{CONDICIONES[condicion]}" + (f" S/ {customer.credito_monto:,.0f} a {customer.credito_dias} días" if condicion == "CREDITO" else "")
                + f". Solicitado por {user.nombre}. {comentario or ''}".strip(), "Customer", customer.id,
                dedup_key=f"CONDICION:{customer.id}:{len(customer.condicion_historial)}", creado_por=user)
    audit(user, "CONDICION_PAGO_SOLICITADA", "Customer", customer.ruc, {"condicion": condicion, "monto": monto, "dias": dias})
    return customer


def resolve_condition(customer, user, decision, comentario=""):
    if not has_perm(user, "finanzas.gestionar"):
        raise BusinessError("La condición de pago la aprueba la Coordinación de Finanzas.")
    decision = (decision or "").upper()
    if decision not in ("APROBADA", "OBSERVADA", "RECHAZADA"):
        raise BusinessError("Decisión inválida.")
    if customer.condicion_estado != "PENDIENTE":
        raise BusinessError("La condición de pago no está pendiente de aprobación.")
    if decision != "APROBADA" and not (comentario or "").strip():
        raise BusinessError("Indique el motivo (qué falta o por qué se rechaza).")
    customer.condicion_estado = decision
    customer.condicion_historial = list(customer.condicion_historial or []) + [
        {"accion": decision, "user": user.nombre, "ts": now().isoformat(timespec="minutes"), "comentario": (comentario or "").strip()}]
    alerts.close_by_system(f"CONDICION:{customer.id}:{len(customer.condicion_historial) - 1}", f"Condición {decision.lower()}")
    alerts.emit("CONDICION_PAGO_RESUELTA", f"Condición de pago {ESTADOS_CONDICION[decision].lower()}: {customer.nombre}",
                (comentario or "").strip() or "Ya puede registrar OV para el cliente.", "Customer", customer.id,
                extra_users=[customer.vendedor_id] if customer.vendedor_id else None, creado_por=user)
    audit(user, f"CONDICION_PAGO_{decision}", "Customer", customer.ruc, {"comentario": comentario})
    return customer


# ------------------------------------------------------------------ facturas y notas de crédito
def _order_by_number(numero):
    numero = (str(numero or "")).strip()
    if not numero:
        return None
    return db().execute(select(SalesOrder).where(SalesOrder.numero == numero)).scalars().first()


def _customer_by(ruc="", nombre=""):
    ruc = str(ruc or "").strip().split(".")[0]
    if ruc:
        c = db().execute(select(Customer).where(Customer.ruc == ruc)).scalars().first()
        if c:
            return c
    nombre = (nombre or "").strip()
    if nombre:
        return db().execute(select(Customer).where(Customer.nombre == nombre)).scalars().first()
    return None


def _apply_credit_note(nc):
    if not nc.referencia:
        return
    fac = db().execute(select(CustomerInvoice).where(CustomerInvoice.tipo == "FACTURA",
                                                     CustomerInvoice.numero == nc.referencia)).scalars().first()
    if fac:
        fac.saldo = round(max(fac.saldo - nc.monto, 0), 2)
        _refresh_state(fac)


def _refresh_state(inv):
    if inv.estado == "ANULADA":
        return
    if inv.saldo <= 0.005:
        inv.estado = "PAGADA"
    elif inv.saldo < inv.monto - 0.005:
        inv.estado = "PARCIAL"
    else:
        inv.estado = "PENDIENTE"


def register_invoice(user, datos, archivo="", origen="MANUAL"):
    """datos: tipo, numero, ruc o customer_id, ov, fecha, vencimiento, moneda, monto, saldo, referencia, observacion."""
    if not has_perm(user, "facturas.cargar"):
        raise BusinessError("Cargan facturas Finanzas y Facturación.")
    tipo = (datos.get("tipo") or "FACTURA").upper().replace(" ", "_")
    if tipo in ("NC", "NOTA_DE_CREDITO", "NOTA_CRÉDITO", "NOTA_DE_CRÉDITO"):
        tipo = "NOTA_CREDITO"
    if tipo not in ("FACTURA", "NOTA_CREDITO"):
        raise BusinessError(f"Tipo de documento no válido: {tipo}.")
    numero = str(datos.get("numero") or "").strip()
    if not numero:
        raise BusinessError("Número de documento obligatorio.")
    order = _order_by_number(datos.get("ov"))
    c = db().get(Customer, int(datos["customer_id"])) if str(datos.get("customer_id") or "").isdigit() else \
        _customer_by(datos.get("ruc"), datos.get("cliente"))
    if c is None and order is not None:
        c = order.customer
    if c is None:
        raise BusinessError(f"{numero}: cliente no encontrado (indique RUC, cliente u OV).")

    def _f(k):
        v = datos.get(k)
        if v in (None, ""):
            return None
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        s = str(v).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(s[:10], fmt).date()
            except ValueError:
                pass
        raise BusinessError(f"{numero}: fecha inválida '{s}' (use dd/mm/aaaa).")
    try:
        monto = float(str(datos.get("monto") or "0").replace(",", ""))
    except ValueError:
        raise BusinessError(f"{numero}: monto inválido.")
    if monto <= 0:
        raise BusinessError(f"{numero}: el monto debe ser mayor que cero.")
    saldo_txt = datos.get("saldo")
    saldo = monto if saldo_txt in (None, "") else float(str(saldo_txt).replace(",", ""))
    inv = db().execute(select(CustomerInvoice).where(CustomerInvoice.tipo == tipo, CustomerInvoice.numero == numero)).scalars().first()
    nuevo = inv is None
    if nuevo:
        inv = CustomerInvoice(tipo=tipo, numero=numero, customer_id=c.id, fecha=_f("fecha") or date.today(), monto=monto,
                              saldo=0 if tipo == "NOTA_CREDITO" else saldo)
        db().add(inv)
    inv.customer_id = c.id
    inv.order_id = order.id if order else inv.order_id
    inv.fecha = _f("fecha") or inv.fecha
    inv.vencimiento = _f("vencimiento") or inv.vencimiento
    inv.moneda = (datos.get("moneda") or inv.moneda or "SOLES").upper()
    inv.monto = monto
    if tipo == "FACTURA" and not nuevo and saldo_txt not in (None, ""):
        inv.saldo = saldo
    inv.referencia = str(datos.get("referencia") or inv.referencia or "").strip()
    inv.observacion = (datos.get("observacion") or inv.observacion or "").strip()
    inv.origen = origen
    inv.user_id = user.id
    if archivo:
        inv.archivo = archivo
    db().flush()
    if tipo == "NOTA_CREDITO" and nuevo:
        _apply_credit_note(inv)
    _refresh_state(inv)
    audit(user, "FACTURA_REGISTRADA" if tipo == "FACTURA" else "NOTA_CREDITO_REGISTRADA", "CustomerInvoice", numero,
          {"cliente": c.ruc, "monto": monto, "origen": origen, "nuevo": nuevo})
    return inv


_ALIAS = {"numero": ("numero", "número", "nro", "n°", "factura", "documento", "comprobante", "serie-numero", "serie y numero"),
          "tipo": ("tipo", "tipo documento", "tipo de documento"),
          "ruc": ("ruc", "ruc cliente", "documento cliente"),
          "cliente": ("cliente", "razon social", "razón social", "nombre cliente"),
          "ov": ("ov", "orden de venta", "pedido", "n° ov", "numero ov"),
          "fecha": ("fecha", "fecha emision", "fecha emisión", "fecha de emision", "fecha de emisión"),
          "vencimiento": ("vencimiento", "fecha vencimiento", "fecha de vencimiento"),
          "moneda": ("moneda",),
          "monto": ("monto", "importe", "total", "importe total"),
          "saldo": ("saldo", "saldo pendiente", "pendiente"),
          "referencia": ("referencia", "factura referencia", "documento referencia")}


def _norm(x):
    return " ".join(str(x or "").strip().lower().replace("_", " ").split())


def import_invoices_excel(user, path, nombre=""):
    """Carga de facturas y notas de crédito desde Excel (exportación de SAP o formato propio). Devuelve (ok, errores)."""
    from openpyxl import load_workbook
    if not has_perm(user, "facturas.cargar"):
        raise BusinessError("Cargan facturas Finanzas y Facturación.")
    ws = load_workbook(path, data_only=True, read_only=True).active
    filas = list(ws.iter_rows(values_only=True))
    hdr_i = None
    for i, r in enumerate(filas[:15]):
        cols = [_norm(v) for v in r]
        if any(c in _ALIAS["numero"] for c in cols) and any(c in _ALIAS["monto"] for c in cols):
            hdr_i = i
            break
    if hdr_i is None:
        raise BusinessError("No se encontró la fila de encabezados (se necesitan al menos «número» y «monto»).")
    cols = [_norm(v) for v in filas[hdr_i]]
    idx = {k: next((j for j, c in enumerate(cols) if c in alias), None) for k, alias in _ALIAS.items()}
    ok, errores = [], []
    for n, r in enumerate(filas[hdr_i + 1:], start=hdr_i + 2):
        if not any(v not in (None, "") for v in r):
            continue
        datos = {k: (r[j] if j is not None and j < len(r) else None) for k, j in idx.items()}
        try:
            ok.append(register_invoice(user, datos, origen="EXCEL"))
        except BusinessError as e:
            errores.append(f"Fila {n}: {e}")
    audit(user, "FACTURAS_CARGADAS", "CustomerInvoice", nombre, {"ok": len(ok), "errores": len(errores)})
    return ok, errores


def attach_pdf(user, archivo, numeros):
    """PDF individual (un número) o consolidado (varios números): se adjunta a cada factura indicada."""
    if not has_perm(user, "facturas.cargar"):
        raise BusinessError("Cargan facturas Finanzas y Facturación.")
    nums = [x.strip() for x in str(numeros or "").replace(";", ",").replace("\n", ",").split(",") if x.strip()]
    if not archivo or not nums:
        raise BusinessError("Adjunte el PDF e indique el número (o los números, si es consolidado).")
    hechos, faltan = [], []
    for n in nums:
        inv = db().execute(select(CustomerInvoice).where(CustomerInvoice.numero == n)).scalars().first()
        if inv is None:
            faltan.append(n)
            continue
        inv.archivo = archivo
        hechos.append(n)
    audit(user, "FACTURA_PDF", "CustomerInvoice", ",".join(hechos), {"archivo": archivo, "no_encontradas": faltan})
    return hechos, faltan


def set_invoice_state(inv, user, estado, observacion=""):
    if not has_perm(user, "finanzas.gestionar"):
        raise BusinessError("Actualiza el estado de pago la Coordinación de Finanzas.")
    estado = (estado or "").upper()
    if estado not in ("PENDIENTE", "PAGADA", "ANULADA"):
        raise BusinessError("Estado inválido.")
    antes = inv.estado
    inv.estado = estado
    inv.saldo = 0 if estado in ("PAGADA", "ANULADA") else (inv.saldo or inv.monto)
    if observacion:
        inv.observacion = (inv.observacion + "\n" if inv.observacion else "") + observacion.strip()
    audit(user, "FACTURA_ESTADO", "CustomerInvoice", inv.numero, {"antes": antes, "despues": estado, "obs": observacion})
    return inv


# ------------------------------------------------------------------ pagos
def report_payment(user, customer, fecha, monto, archivo, invoice_id=None, medio="", operacion="", observacion=""):
    """Comercial (vendedor, coordinador o jefe) o Finanzas sube el comprobante; Finanzas lo confirma."""
    if not has_perm(user, "pagos.reportar"):
        raise BusinessError("Suben comprobantes de pago el vendedor, el coordinador, el jefe comercial o Finanzas.")
    if user.role_code == "EJECUTIVO_VENTAS" and customer.vendedor_id not in (None, user.id):
        raise BusinessError("Solo puede subir comprobantes de los clientes que atiende.")
    if not archivo:
        raise BusinessError("Adjunte el comprobante de pago.")
    try:
        monto = float(monto)
    except (TypeError, ValueError):
        raise BusinessError("Monto inválido.")
    if monto <= 0:
        raise BusinessError("El monto debe ser mayor que cero.")
    inv = db().get(CustomerInvoice, int(invoice_id)) if invoice_id else None
    if inv is not None and inv.customer_id != customer.id:
        raise BusinessError("La factura no pertenece al cliente.")
    if isinstance(fecha, str):
        fecha = date.fromisoformat(fecha) if fecha else date.today()
    p = CustomerPayment(customer_id=customer.id, invoice_id=inv.id if inv else None, fecha=fecha or date.today(), monto=monto,
                        medio=(medio or "").strip(), operacion=(operacion or "").strip(), archivo=archivo,
                        observacion=(observacion or "").strip(), user_id=user.id)
    db().add(p)
    db().flush()
    audit(user, "PAGO_REPORTADO", "Customer", customer.ruc, {"monto": monto, "factura": inv.numero if inv else None})
    if has_perm(user, "finanzas.gestionar"):
        confirm_payment(p, user, True, "Registrado por Finanzas")
    else:
        alerts.emit("PAGO_REPORTADO", f"Confirmar pago de {customer.nombre}: S/ {monto:,.2f}",
                    (f"Factura {inv.numero}. " if inv else "") + f"Subido por {user.nombre}. {observacion or ''}".strip(),
                    "CustomerPayment", p.id, dedup_key=f"PAGO:{p.id}", creado_por=user)
    return p


def confirm_payment(p, user, ok, observacion=""):
    if not has_perm(user, "finanzas.gestionar"):
        raise BusinessError("Confirma los pagos la Coordinación de Finanzas.")
    if p.estado != "REPORTADO":
        raise BusinessError("El pago ya fue revisado.")
    if not ok and not (observacion or "").strip():
        raise BusinessError("Indique por qué se observa el comprobante.")
    p.estado = "CONFIRMADO" if ok else "OBSERVADO"
    p.confirmado_por = user.id
    p.confirmado_ts = now()
    if observacion:
        p.observacion = ((p.observacion + "\n") if p.observacion else "") + observacion.strip()
    if ok and p.invoice is not None:
        p.invoice.saldo = round(max((p.invoice.saldo or 0) - p.monto, 0), 2)
        _refresh_state(p.invoice)
    alerts.close_by_system(f"PAGO:{p.id}", "Pago revisado")
    audit(user, "PAGO_CONFIRMADO" if ok else "PAGO_OBSERVADO", "CustomerPayment", p.id, {"obs": observacion})
    return p


# ------------------------------------------------------------------ estado de cuenta y cobranza
def statement(customer, hoy=None):
    """Deuda y estado de cuenta del cliente (todo en su archivo)."""
    hoy = hoy or date.today()
    invs = list(db().execute(select(CustomerInvoice).where(CustomerInvoice.customer_id == customer.id)
                             .order_by(CustomerInvoice.fecha.desc())).scalars())
    facturas = [i for i in invs if i.tipo == "FACTURA" and i.estado != "ANULADA"]
    abiertas = [i for i in facturas if i.saldo > 0.005]
    vencidas = [i for i in abiertas if i.vencimiento and i.vencimiento < hoy]
    pagos = list(db().execute(select(CustomerPayment).where(CustomerPayment.customer_id == customer.id)
                              .order_by(CustomerPayment.fecha.desc())).scalars())
    return {"documentos": invs, "abiertas": abiertas, "vencidas": vencidas,
            "deuda": round(sum(i.saldo for i in abiertas), 2), "deuda_vencida": round(sum(i.saldo for i in vencidas), 2),
            "facturado": round(sum(i.monto for i in facturas) - sum(i.monto for i in invs if i.tipo == "NOTA_CREDITO"), 2),
            "pagos": pagos, "por_confirmar": [p for p in pagos if p.estado == "REPORTADO"],
            "estado_pago": "VENCIDA" if vencidas else ("CON_SALDO" if abiertas else "AL_DIA")}


def collection_board(hoy=None):
    """Cuadro de cobranza de Finanzas: vencidas por antigüedad, por vencer, pagos por confirmar y condiciones por aprobar."""
    hoy = hoy or date.today()
    aviso = int(param("dias_aviso_vencimiento_factura", 3) or 0)
    abiertas = [i for i in db().execute(select(CustomerInvoice).where(CustomerInvoice.tipo == "FACTURA",
                                                                      CustomerInvoice.estado.in_(["PENDIENTE", "PARCIAL"]))
                                        .order_by(CustomerInvoice.vencimiento)).scalars() if i.saldo > 0.005]
    tramos = {"1-15": [], "16-30": [], "31-60": [], "61+": []}
    por_vencer, vigentes = [], []
    for i in abiertas:
        if i.vencimiento and i.vencimiento < hoy:
            d = (hoy - i.vencimiento).days
            tramos["1-15" if d <= 15 else "16-30" if d <= 30 else "31-60" if d <= 60 else "61+"].append((i, d))
        elif i.vencimiento and i.vencimiento <= hoy + timedelta(days=aviso):
            por_vencer.append(i)
        else:
            vigentes.append(i)
    por_cliente = defaultdict(lambda: {"vencido": 0.0, "por_vencer": 0.0, "docs": 0})
    for k, lst in tramos.items():
        for i, _ in lst:
            por_cliente[i.customer]["vencido"] += i.saldo
            por_cliente[i.customer]["docs"] += 1
    for i in por_vencer:
        por_cliente[i.customer]["por_vencer"] += i.saldo
    pagos = list(db().execute(select(CustomerPayment).where(CustomerPayment.estado == "REPORTADO")
                              .order_by(CustomerPayment.ts)).scalars())
    condiciones = list(db().execute(select(Customer).where(Customer.condicion_estado == "PENDIENTE")).scalars())
    return {"tramos": tramos, "por_vencer": por_vencer, "vigentes": vigentes, "pagos": pagos, "condiciones": condiciones,
            "por_cliente": sorted(por_cliente.items(), key=lambda x: -x[1]["vencido"]),
            "total_vencido": round(sum(i.saldo for lst in tramos.values() for i, _ in lst), 2),
            "total_por_vencer": round(sum(i.saldo for i in por_vencer), 2),
            "total_abierto": round(sum(i.saldo for i in abiertas), 2)}


def _sent(dedup_key):
    from ..models import Alert
    return db().execute(select(Alert.id).where(Alert.dedup_key == dedup_key)).first() is not None


def check_overdue(hoy=None):
    """Tarea periódica: factura vencida con saldo → alerta a Finanzas y al vendedor (una vez por factura)."""
    hoy = hoy or date.today()
    n = 0
    for i in db().execute(select(CustomerInvoice).where(CustomerInvoice.tipo == "FACTURA",
                                                        CustomerInvoice.estado.in_(["PENDIENTE", "PARCIAL"]),
                                                        CustomerInvoice.vencimiento < hoy)).scalars():
        if i.saldo <= 0.005 or _sent(f"FACT_VENC:{i.id}"):
            continue
        if alerts.emit("FACTURA_VENCIDA", f"Factura {i.numero} vencida: {i.customer.nombre}",
                       f"Venció el {i.vencimiento:%d/%m/%Y}; saldo S/ {i.saldo:,.2f}.", "CustomerInvoice", i.id,
                       extra_users=[i.customer.vendedor_id] if i.customer.vendedor_id else None,
                       dedup_key=f"FACT_VENC:{i.id}") is not None:
            n += 1
    return n


# ------------------------------------------------------------------ expediente de la OV
def attach_to_order(order, user, tipo, archivo, nombre="", observacion=""):
    tipo = (tipo or "").upper()
    if tipo not in TIPOS_DOC_OV:
        raise BusinessError("Tipo de documento no válido.")
    if tipo == "OC":
        from .orders import attach_oc
        return attach_oc(order, user, archivo, nombre)
    if not archivo:
        raise BusinessError("Adjunte el archivo.")
    a = OrderAttachment(order_id=order.id, tipo=tipo, nombre=(nombre or TIPOS_DOC_OV[tipo]).strip(), archivo=archivo,
                        observacion=(observacion or "").strip(), user_id=user.id)
    db().add(a)
    audit(user, "OV_DOCUMENTO", "SalesOrder", order.numero, {"tipo": tipo, "archivo": archivo})
    return a


def request_credit_note(order, user, motivo, monto=None):
    if not has_perm(user, "nc.solicitar"):
        raise BusinessError("Solicitan notas de crédito el vendedor, el coordinador o el jefe comercial.")
    if not (motivo or "").strip():
        raise BusinessError("Indique el motivo de la nota de crédito.")
    a = alerts.emit("NOTA_CREDITO_SOLICITADA", f"Nota de crédito para la OV {order.numero} ({order.customer.nombre})",
                    f"Motivo: {motivo.strip()}" + (f". Monto estimado S/ {float(monto):,.2f}" if monto else "")
                    + f". Solicitado por {user.nombre}.", "SalesOrder", order.id, creado_por=user)
    audit(user, "NOTA_CREDITO_SOLICITADA", "SalesOrder", order.numero, {"motivo": motivo, "monto": monto})
    return a


def order_documents(order):
    atts = list(db().execute(select(OrderAttachment).where(OrderAttachment.order_id == order.id)
                             .order_by(OrderAttachment.ts)).scalars())
    invs = list(db().execute(select(CustomerInvoice).where(CustomerInvoice.order_id == order.id)
                             .order_by(CustomerInvoice.fecha)).scalars())
    return {"adjuntos": atts, "facturas": invs}
