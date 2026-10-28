"""Carga de usuarios por Excel (V9, 28-oct-2026): correo corporativo, nombre y perfil (rol) del sistema.

- Columnas: «Correo» (o «Correos»), «Apellidos y nombres» (o «Nombre»), «Rol» (código o nombre del perfil, p. ej.
  EJECUTIVO_VENTAS o «Ejecutivo de Ventas (vendedor por canal)») y, opcional, «Almacén». El DNI, si viene, no se guarda.
- No se adivinan perfiles: si «Rol» está vacío o no existe, la fila no se carga y se informa.
- El usuario se identifica por su correo: si ya existe se actualizan su nombre y su perfil. El nombre de usuario es la
  parte del correo antes de «@». Ingresan con Microsoft 365 (correo corporativo); la clave local es aleatoria y, si
  Sistemas asigna una, se cambia en el primer ingreso.
"""
import re
import secrets
import unicodedata

from sqlalchemy import func, select
from werkzeug.security import generate_password_hash

from ..db import db
from ..models import Role, User, Warehouse
from .core import BusinessError, audit, has_perm, param

EMAIL = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)+$")


def _n(x):
    x = unicodedata.normalize("NFKD", str(x or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(x.split())


def _nombre(t):
    t = " ".join(str(t or "").split())
    return t.title() if t.isupper() else t


def import_users(admin, path):
    from openpyxl import load_workbook
    if not has_perm(admin, "usuarios.administrar"):
        raise BusinessError("Solo Sistemas administra usuarios.")
    filas = [list(r) for r in load_workbook(path, read_only=True, data_only=True).active.iter_rows(values_only=True)
             if any(c not in (None, "") for c in r)]
    if not filas:
        raise BusinessError("El archivo está vacío.")
    cab = [_n(c) for c in filas[0]]

    def ix(*nombres):
        return next((cab.index(_n(x)) for x in nombres if _n(x) in cab), None)
    i_mail, i_nom, i_rol, i_alm = ix("correo", "correos", "email"), ix("apellidos y nombres", "nombre", "nombres"), \
        ix("rol", "perfil"), ix("almacen", "almacén")
    if i_mail is None or i_rol is None:
        raise BusinessError("Faltan columnas: «Correo» y «Rol» (perfil del sistema).")
    roles = {r.code: r for r in db().execute(select(Role)).scalars()}
    por_nombre = {_n(r.nombre): r for r in roles.values()}
    dominio = (param("correo_dominio_corporativo", "") or "").strip().lower()
    res = {"creados": [], "actualizados": [], "sin_rol": [], "errores": []}
    vistos = set()
    for k, r in enumerate(filas[1:], start=2):
        mail = str(r[i_mail] or "").strip().lower() if i_mail < len(r) else ""
        nombre = _nombre(r[i_nom]) if i_nom is not None and i_nom < len(r) else ""
        rol_txt = str(r[i_rol] or "").strip() if i_rol < len(r) else ""
        if not EMAIL.match(mail):
            res["errores"].append(f"Fila {k} ({nombre or '-'}): correo «{mail}» inválido.")
            continue
        if dominio and not mail.endswith("@" + dominio):
            res["errores"].append(f"Fila {k} ({nombre or '-'}): el correo no es del dominio {dominio}.")
            continue
        if mail in vistos:
            res["errores"].append(f"Fila {k}: el correo {mail} está repetido en el archivo.")
            continue
        vistos.add(mail)
        rol = roles.get(rol_txt.upper()) or por_nombre.get(_n(rol_txt))
        if rol is None:
            res["sin_rol"].append(f"{nombre or mail} ({mail}): perfil «{rol_txt or 'vacío'}» no existe en el sistema")
            continue
        alm = str(r[i_alm] or "").strip() if i_alm is not None and i_alm < len(r) else ""
        if alm and db().get(Warehouse, alm) is None:
            res["errores"].append(f"Fila {k} ({nombre}): almacén «{alm}» no existe; se carga sin almacén.")
            alm = ""
        u = db().execute(select(User).where(func.lower(User.email) == mail)).scalars().first()
        if u is None:
            base = re.sub(r"[^a-z0-9._-]", "", mail.split("@")[0]) or "usuario"
            un, n = base, 1
            while db().execute(select(User).where(User.username == un)).first():
                n += 1
                un = f"{base}{n}"
            u = User(username=un, nombre=nombre or un, email=mail, role_code=rol.code,
                     password_hash=generate_password_hash(secrets.token_urlsafe(24)), activo=True,
                     debe_cambiar_clave=True, almacen=alm)
            db().add(u)
            res["creados"].append(f"{u.nombre} → {rol.nombre} (usuario {un})")
        else:
            antes = (u.nombre, u.role_code, u.almacen)
            u.nombre, u.role_code = nombre or u.nombre, rol.code
            if alm:
                u.almacen = alm
            if antes != (u.nombre, u.role_code, u.almacen):
                res["actualizados"].append(f"{u.nombre} → {rol.nombre}")
    db().flush()
    audit(admin, "USUARIOS_CARGADOS", "User", "", {k: len(v) for k, v in res.items()})
    return res
