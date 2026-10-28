"""Carga de OV desde las exportaciones de SAP en Excel (V8, formatos del 27-oct-2026).

SAP entrega dos archivos:
1. Lista de pedidos de cliente: una fila por OV (# SAP, Fecha, Cliente, Empleado del departamento de ventas,
   Comentarios, Fecha de vencimiento, «Pedidos de cliente - C<RUC>», Peso Total, Status de documento, Destino, …).
   No trae los productos. Las OV «Abiertos» quedan en espera de su detalle; las «Cerrado» no se importan (si ya están
   abiertas en el módulo y aún no entraron a preparación, se cierran como «Atendida en SAP»).
2. Detalle de UNA OV: una fila por producto (Artículo, Descripción, Cantidad, Precio tras el descuento, Impuesto,
   Almacén, Cantidad pendiente, Peso…). El archivo no trae el número de la OV: se elige en pantalla o se toma del
   nombre del archivo (p. ej. «1636048.xlsx»). Al llegar el detalle se crea la OV con la cantidad PENDIENTE.

Sin duplicados: la cabecera y la OV se identifican por el número de SAP; si se cargan varias veces al día solo se
actualiza lo que cambió.
"""
import re
import unicodedata

from sqlalchemy import select

from ..db import db
from ..models import Customer, SalesOrder, SapOrderHeader, now
from .core import BusinessError, audit, has_perm


def order_by_number(numero):
    from .orders import order_by_number as _f  # import diferido: orders importa servicios que importan este módulo
    return _f(numero)


def _n(x):
    x = unicodedata.normalize("NFKD", str(x or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(x.replace("#", " # ").split())


def _rows(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _txt(v):
    return " ".join(str(v).split()) if v not in (None, "") else ""


def _num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(" ", "")
    if "," in t and "." in t:
        t = t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".")
    return float(t)


def _check(user):
    if not has_perm(user, "ov.importar"):
        raise BusinessError("Su rol no importa OV de SAP.")


def detect(path):
    """CABECERAS, DETALLE o None según las columnas."""
    cab = [_n(c) for c in (_rows(path)[:1] or [[]])[0]]
    if "# sap" in cab and "status de documento" in cab:
        return "CABECERAS"
    if "articulo" in cab and "cantidad" in cab:
        return "DETALLE"
    return None


# ------------------------------------------------------------------ 1. lista de OV
def import_headers(user, path):
    from .uploads import _date
    _check(user)
    filas = _rows(path)
    if not filas:
        raise BusinessError("El archivo está vacío.")
    cab = [_n(c) for c in filas[0]]
    if "# sap" not in cab:
        raise BusinessError("No es la lista de pedidos de SAP: falta la columna «# SAP».")
    pos = {}
    for i, c in enumerate(cab):
        pos.setdefault(c, []).append(i)

    def col(r, nombre, k=0):
        ix = pos.get(_n(nombre), [])
        return r[ix[k]] if len(ix) > k and ix[k] < len(r) else None
    res = {"nuevas": 0, "actualizadas": 0, "sin_cambio": 0, "pendientes_detalle": 0, "cerradas_omitidas": 0,
           "cerradas_en_modulo": 0, "canceladas": 0, "errores": []}
    for n, r in enumerate(filas[1:], start=2):
        numero = _txt(col(r, "# SAP"))
        if not numero:
            continue
        numero = numero[:-2] if numero.endswith(".0") else numero
        estado = _txt(col(r, "Status de documento"))
        cancelado = _txt(col(r, "Cancelado")).lower() in ("si", "sí", "yes", "y")
        ref = _txt(col(r, "Comentarios", 1))
        m = re.search(r"\bC(\d{8}|\d{11})\b", ref) or re.search(r"\b(\d{11})\b", ref)
        try:
            datos = {"fecha": _date(col(r, "Fecha")), "fecha_venc": _date(col(r, "Fecha de vencimiento")),
                     "cliente": _txt(col(r, "Cliente"))[:200],
                     "card_code": ("C" + m.group(1)) if m else "", "ruc": m.group(1) if m else "",
                     "vendedor": _txt(col(r, "Empleado del departamento de ventas"))[:120],
                     "comentario": _txt(col(r, "Comentarios", 0)), "destino": _txt(col(r, "Destino"))[:250],
                     "peso_total": _num(col(r, "Peso Total")), "estado_sap": estado[:20], "cancelado": cancelado}
        except ValueError as e:
            res["errores"].append(f"Fila {n} (OV {numero}): {e}")
            continue
        h = db().execute(select(SapOrderHeader).where(SapOrderHeader.numero == numero)).scalars().first()
        o = order_by_number(numero)  # V10: también la OV del módulo creada en SAP
        cerrada = estado.lower().startswith("cerr") or cancelado
        if cerrada:
            if cancelado:
                res["canceladas"] += 1
            if o is not None and _cerrar_en_modulo(o, user, "cancelada" if cancelado else "cerrada"):
                res["cerradas_en_modulo"] += 1
            elif h is None and o is None:
                res["cerradas_omitidas"] += 1
                continue
        if h is None:
            h = SapOrderHeader(numero=numero, user_id=user.id, **datos)
            db().add(h)
            res["nuevas"] += 1
        else:
            cambio = {k: v for k, v in datos.items() if getattr(h, k) != v}
            if cambio:
                for k, v in cambio.items():
                    setattr(h, k, v)
                h.actualizado_ts = now()
                res["actualizadas"] += 1
            else:
                res["sin_cambio"] += 1
        if o is not None:
            h.order_id = o.id
        elif not cerrada:
            res["pendientes_detalle"] += 1
    db().flush()
    audit(user, "SAP_EXCEL_CABECERAS", "SapOrderHeader", "", {k: v for k, v in res.items() if k != "errores"})
    return res


def _cerrar_en_modulo(o, user, como):
    """Misma regla que la carga de SAP: si la OV aún no entra a preparación se cierra y libera reservas, sin descontar
    stock (lo refleja la próxima carga de stock de SAP)."""
    from .stock import release_order
    if o.estado not in ("BORRADOR", "REGISTRADA", "EVALUADA", "BLOQUEADA"):
        return False
    release_order(o, f"OV {como} en SAP")
    antes = o.estado
    o.estado = "CANCELADA" if como == "cancelada" else "CERRADA"
    o.motivo_detencion = f"{'Cancelada' if como == 'cancelada' else 'Atendida'} en SAP"
    audit(user, "OV_CERRADA_EN_SAP", "SalesOrder", o.numero, {"estado_anterior": antes, "como": como})
    return True


def pending_headers():
    return list(db().execute(select(SapOrderHeader).where(SapOrderHeader.order_id.is_(None),
                                                          SapOrderHeader.cancelado.is_(False),
                                                          ~SapOrderHeader.estado_sap.ilike("cerr%"))
                             .order_by(SapOrderHeader.fecha_venc, SapOrderHeader.numero)).scalars())


# ------------------------------------------------------------------ 2. detalle de una OV
def number_from_filename(nombre):
    """Número de OV en el nombre del archivo, si coincide con una cabecera cargada."""
    for m in re.findall(r"\d{5,10}", nombre or ""):
        if db().execute(select(SapOrderHeader).where(SapOrderHeader.numero == m)).scalars().first():
            return m
    return None


def _customer(h):
    c = None
    if h.ruc:
        c = db().execute(select(Customer).where(Customer.ruc == h.ruc)).scalars().first()
    if c is None and h.card_code:
        c = db().execute(select(Customer).where(Customer.codigo_sap == h.card_code)).scalars().first()
    if c is None and h.cliente:  # la fila no trae «Pedidos de cliente - C<RUC>»: se busca por el nombre exacto
        from .customer_conditions import norm_name
        nn = norm_name(h.cliente)
        c = next((x for x in db().execute(select(Customer)).scalars() if norm_name(x.nombre) == nn), None)
    if c is None:
        if not h.ruc:
            raise BusinessError(f"OV {h.numero}: la lista de SAP no trae el RUC de «{h.cliente}» y no existe un cliente con "
                                "ese nombre en el módulo. Regístrelo en Clientes y vuelva a subir el detalle.")
        c = Customer(ruc=h.ruc, nombre=h.cliente, codigo_sap=h.card_code, direccion=h.destino,
                     condicion_estado="APROBADA")  # cliente que ya compra en SAP
        db().add(c)
        db().flush()
    elif h.card_code and not c.codigo_sap:
        c.codigo_sap = h.card_code
    return c


def import_detail(user, path, numero=None, nombre_archivo=""):
    """Crea la OV del módulo con el detalle exportado de SAP (cantidad pendiente de cada producto)."""
    from .orders import create_order
    from .uploads import _ensure_product
    _check(user)
    numero = (numero or "").strip() or number_from_filename(nombre_archivo)
    if not numero:
        raise BusinessError("Indique a qué OV corresponde el detalle (o nombre el archivo con el número de la OV).")
    h = db().execute(select(SapOrderHeader).where(SapOrderHeader.numero == numero)).scalars().first()
    if h is None:
        raise BusinessError(f"Primero cargue la lista de pedidos de SAP que incluye la OV {numero}.")
    if h.order_id:
        raise BusinessError(f"La OV {numero} ya está en el módulo.")
    if h.estado_sap.lower().startswith("cerr") or h.cancelado:
        raise BusinessError(f"La OV {numero} está cerrada o cancelada en SAP: no se importa.")
    filas = _rows(path)
    cab = [_n(c) for c in (filas[0] if filas else [])]
    if "articulo" not in cab:
        raise BusinessError("No es el detalle de una OV de SAP: falta la columna «Artículo».")

    def col(r, nombre):
        k = _n(nombre)
        return r[cab.index(k)] if k in cab and cab.index(k) < len(r) else None
    lineas, errores = {}, []
    for n, r in enumerate(filas[1:], start=2):
        code = _txt(col(r, "Artículo")).upper()
        if not code:
            continue
        try:
            q = _num(col(r, "Cantidad"))
            pend = _num(col(r, "Cantidad pendiente"))
            precio = _num(col(r, "Precio tras el descuento"))
            if precio is None:
                precio = _num(col(r, "Precio por unidad"))
        except ValueError:
            errores.append(f"Fila {n}: cantidad o precio no numérico.")
            continue
        q = pend if pend is not None else q
        if not q or q <= 0:
            continue  # línea ya atendida en SAP
        desc = _txt(col(r, "Descripción Artículo"))
        um = "KILOGRAMO" if _txt(col(r, "Código de unidad de medida")).upper() in ("KGM", "KG") else "UNIDAD"
        _ensure_product(code, desc, "PT" if code.startswith("PT") else None, um)
        ln = lineas.setdefault(code, {"product_code": code, "cantidad": 0.0, "precio_unitario": precio or 0,
                                      "tax_code": _txt(col(r, "Impuesto")).upper()[:10]})
        ln["cantidad"] += q
    if errores:
        raise BusinessError(" ".join(errores[:5]))
    if not lineas:
        raise BusinessError(f"La OV {numero} no tiene cantidades pendientes: ya está atendida en SAP.")
    c = _customer(h)
    from ..models import User
    vend = db().execute(select(User).where(User.nombre == h.vendedor)).scalars().first()
    entrega = {"site_id": None, "lugar_entrega": "Destino SAP" if h.destino else "", "direccion": h.destino,
               "distrito": "", "provincia": ""}
    o = create_order(user, numero, c.id, h.fecha_venc or h.fecha, list(lineas.values()), comentarios=h.comentario,
                     vendedor_id=vend.id if vend else None, vendedor_nombre=h.vendedor, fecha_emision=h.fecha,
                     origen="CARGA_SAP", entrega=entrega)
    h.order_id = o.id
    db().flush()
    audit(user, "SAP_EXCEL_DETALLE", "SalesOrder", numero, {"lineas": len(lineas)})
    return o
