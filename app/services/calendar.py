"""Calendario de operación: jornada por día de la semana y feriados (decisiones del 15 y 16-oct-2026).

- Sábados hasta mediodía (factor 0,5), domingos no se opera.
- Feriados: la lista se carga cada año (Excel o PDF, o uno por uno). Un feriado puede marcarse «se trabaja» según la
  necesidad; en ese caso cuenta como día normal.
- Parámetro 'feriados': lista de {fecha: AAAA-MM-DD, nombre, se_trabaja}. Se aceptan también fechas sueltas (formato
  anterior), que se leen como feriados no laborables.
"""
from datetime import date

from .core import BusinessError, audit, has_perm, param, set_param


def holidays():
    """{fecha_iso: {fecha, nombre, se_trabaja}} ordenado por fecha."""
    out = {}
    for x in param("feriados", []) or []:
        if isinstance(x, str):
            out[x[:10]] = {"fecha": x[:10], "nombre": "", "se_trabaja": False}
        elif isinstance(x, dict) and x.get("fecha"):
            out[str(x["fecha"])[:10]] = {"fecha": str(x["fecha"])[:10], "nombre": x.get("nombre", ""),
                                         "se_trabaja": bool(x.get("se_trabaja"))}
    return dict(sorted(out.items()))


def day_factor(d, feriados=None, jornada=None):
    """1 jornada completa; 0,5 sábado; 0 domingo o feriado que no se trabaja."""
    feriados = holidays() if feriados is None else feriados
    f = feriados.get(d.isoformat())
    if f and not f["se_trabaja"]:
        return 0.0
    if jornada is None:
        jornada = {int(k): float(v) for k, v in (param("dias_operacion", {}) or {}).items()}
    return jornada.get(d.weekday(), 1.0)


def _guardar(user, datos, motivo):
    set_param("feriados", [datos[k] for k in sorted(datos)], user)
    audit(user, "FERIADOS", "Param", "feriados", {"motivo": motivo, "total": len(datos)})


def _puede_cargar(user):
    if not has_perm(user, "config.editar"):
        raise BusinessError("La lista de feriados la carga Sistemas.")


def add_holiday(user, fecha, nombre="", se_trabaja=False):
    _puede_cargar(user)
    if not isinstance(fecha, date):
        raise BusinessError("Indique la fecha del feriado.")
    datos = holidays()
    datos[fecha.isoformat()] = {"fecha": fecha.isoformat(), "nombre": (nombre or "").strip(), "se_trabaja": bool(se_trabaja)}
    _guardar(user, datos, f"Feriado {fecha.isoformat()}")


def remove_holiday(user, fecha_iso):
    _puede_cargar(user)
    datos = holidays()
    if datos.pop(fecha_iso, None) is None:
        raise BusinessError("Feriado no encontrado.")
    _guardar(user, datos, f"Quitado {fecha_iso}")


def set_worked(user, fecha_iso, se_trabaja):
    """Marcar si un feriado se trabaja según la necesidad (Sistemas o el Jefe de Producción)."""
    if not (has_perm(user, "config.editar") or has_perm(user, "produccion.programar")):
        raise BusinessError("Si un feriado se trabaja lo indica el Jefe de Producción o Sistemas.")
    datos = holidays()
    if fecha_iso not in datos:
        raise BusinessError("Feriado no encontrado.")
    datos[fecha_iso]["se_trabaja"] = bool(se_trabaja)
    _guardar(user, datos, f"{fecha_iso} {'se trabaja' if se_trabaja else 'no se trabaja'}")


def import_holidays(user, path, nombre_archivo):
    """Carga anual: Excel o PDF con columnas «Fecha», «Nombre» (o «Descripción») y opcional «Se trabaja» (SÍ/NO).
    Reemplaza los feriados de los años que trae el archivo. Devuelve (cargados, errores, años)."""
    from .uploads import _date, _pdf_tables, norm
    _puede_cargar(user)
    alias = {norm("Fecha"): "fecha", norm("Nombre"): "nombre", norm("Descripción"): "nombre", norm("Descripcion"): "nombre",
             norm("Feriado"): "nombre", norm("Se trabaja"): "se_trabaja", norm("Laborable"): "se_trabaja"}
    if str(path).lower().endswith(".pdf"):
        filas = _pdf_tables(path, alias, ["fecha"])
    else:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            filas = list(wb.worksheets[0].iter_rows(values_only=True))
        except Exception as e:  # noqa: BLE001
            raise BusinessError(f"No se pudo leer el Excel: {e}") from None
    cab_i = next((i for i, f in enumerate(filas[:15]) if any(alias.get(norm(c)) == "fecha" for c in (f or ()))), None)
    if cab_i is None:
        raise BusinessError("El archivo debe tener la columna «Fecha» (y «Nombre»; opcional «Se trabaja»).")
    cols = {alias.get(norm(c)): j for j, c in enumerate(filas[cab_i]) if alias.get(norm(c))}
    nuevos, errores = {}, []
    for k, f in enumerate(filas[cab_i + 1:], start=cab_i + 2):
        if not f or all(c in (None, "") for c in f):
            continue
        try:
            d = _date(f[cols["fecha"]])
            if d is None:
                raise ValueError("fecha vacía")
        except (ValueError, IndexError) as e:
            errores.append(f"Fila {k}: {e}")
            continue
        st = str(f[cols["se_trabaja"]] if "se_trabaja" in cols and cols["se_trabaja"] < len(f) else "").strip().upper()
        nuevos[d.isoformat()] = {"fecha": d.isoformat(),
                                 "nombre": str(f[cols["nombre"]] or "").strip() if "nombre" in cols and cols["nombre"] < len(f) else "",
                                 "se_trabaja": st in ("SI", "SÍ", "S", "X", "1", "TRUE")}
    if not nuevos:
        raise BusinessError("El archivo no tiene feriados válidos. " + "; ".join(errores[:5]))
    anios = sorted({k[:4] for k in nuevos})
    datos = {k: v for k, v in holidays().items() if k[:4] not in anios}  # se reemplaza el año que trae el archivo
    datos.update(nuevos)
    _guardar(user, datos, f"Carga anual {', '.join(anios)} ({nombre_archivo})")
    return len(nuevos), errores, anios
