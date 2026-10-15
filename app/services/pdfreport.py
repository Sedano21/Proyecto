"""Informes en PDF descargable (obs. demo v0.7: «deben poder descargar en PDF los informes que se requiera»).

Convierte el contenido del informe tal como se ve en pantalla (títulos, indicadores, tablas, gráficos y textos) a un
PDF con fpdf2, sin servicios externos. Se omiten los controles de la pantalla (filtros, botones, formularios). Los
gráficos SVG se insertan como imagen vectorial.

Fuente: DejaVu Sans si está disponible (static/fonts o la del sistema), para tildes y signos; si no, Helvetica con los
caracteres fuera de Latin-1 reemplazados.
"""
import io
import os
import re
from datetime import datetime
from html import unescape
from html.parser import HTMLParser

OMITIR_TAGS = {"script", "style", "form", "button", "select", "input", "textarea", "nav", "datalist", "dialog", "template"}
OMITIR_CLASES = {"noprint", "tabs", "row-actions", "row-menu", "filters", "flash", "demo-strip", "live-note", "skip"}
BLOQUES = {"p", "div", "section", "li", "h1", "h2", "h3", "h4", "table", "ul", "ol", "svg", "details", "summary", "small"}


class Nodo:
    __slots__ = ("tag", "attrs", "hijos", "padre")

    def __init__(self, tag, attrs=None, padre=None):
        self.tag, self.attrs, self.hijos, self.padre = tag, dict(attrs or {}), [], padre

    def clases(self):
        return set((self.attrs.get("class") or "").split())

    def texto(self):
        out = []
        for h in self.hijos:
            if isinstance(h, str):
                out.append(h)
            elif h.tag == "br":
                out.append("\n")
            elif h.tag not in OMITIR_TAGS and not (h.clases() & OMITIR_CLASES) and h.tag != "svg":
                out.append(h.texto())
        return "".join(out)


VACIOS = {"br", "img", "input", "meta", "link", "hr", "use", "line", "rect", "circle", "path", "polyline", "col"}


class _Arbol(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.raiz = Nodo("root")
        self.act = self.raiz

    def handle_starttag(self, tag, attrs):
        n = Nodo(tag, attrs, self.act)
        self.act.hijos.append(n)
        if tag not in VACIOS:
            self.act = n

    def handle_startendtag(self, tag, attrs):
        self.act.hijos.append(Nodo(tag, attrs, self.act))

    def handle_endtag(self, tag):
        n = self.act
        while n is not None and n.tag != tag:
            n = n.padre
        if n is not None and n.padre is not None:
            self.act = n.padre

    def handle_data(self, data):
        self.act.hijos.append(data)


def _buscar(n, pred):
    if not isinstance(n, Nodo):
        return None
    if pred(n):
        return n
    for h in n.hijos:
        r = _buscar(h, pred)
        if r is not None:
            return r
    return None


def _sin_var(v):
    """Las variables CSS (var(--x, #fff)) no existen en el PDF: se usa el valor de respaldo o un gris oscuro."""
    v = re.sub(r"var\(--[\w-]+\s*,\s*([^)]+)\)", r"\1", v)
    v = re.sub(r"var\(--[\w-]+\)", "#1c2621", v).replace('"', "'")
    return "none" if v.strip() == "transparent" else v


def _svg_texto(n):
    """Serializa un SVG del árbol con estilos en línea (las clases CSS no existen en el PDF)."""
    estilos = {"grid": 'stroke="#e2e6e3" stroke-width="1"', "ax": 'font-size="11" fill="#5f6b65" font-family="sans-serif"',
               "lbl": 'font-size="11" fill="#1c2621" font-family="sans-serif"', "val": 'font-size="10" fill="#1c2621" font-family="sans-serif"'}

    def ser(x):
        if isinstance(x, str):
            return x.replace("&", "&amp;").replace("<", "&lt;")
        if x.tag == "title":
            return ""
        attrs = " ".join(f'{k}="{_sin_var(v or "")}"' for k, v in x.attrs.items() if k not in ("class", "role", "aria-label"))
        extra = " ".join(estilos[c] for c in x.clases() if c in estilos)
        if x.tag == "svg" and "xmlns" not in x.attrs:
            attrs += ' xmlns="http://www.w3.org/2000/svg"'
            vb = x.attrs.get("viewbox") or x.attrs.get("viewBox")
            if vb:
                attrs = attrs.replace("viewbox=", "viewBox=")
                p = vb.split()
                if len(p) == 4:
                    attrs += f' width="{p[2]}" height="{p[3]}"'
        cuerpo = "".join(ser(h) for h in x.hijos)
        return f"<{x.tag} {attrs} {extra}>{cuerpo}</{x.tag}>"
    return ser(n)


def _limpio(t):
    t = unescape(t or "")
    t = re.sub(r"[ \t\r\f\v]+", " ", t)
    t = re.sub(r" *\n *", "\n", t)
    return re.sub(r"\n{3,}", "\n\n", t).strip()


class _PDF:
    def __init__(self, titulo, subtitulo=""):
        from fpdf import FPDF
        self.pdf = FPDF(orientation="L", unit="mm", format="A4")
        self.pdf.set_auto_page_break(True, margin=12)
        self.pdf.set_margins(12, 12, 12)
        self.titulo = titulo
        self.unicode = self._fuente()
        pdf = self.pdf
        tit, sub = self.t(titulo), self.t(subtitulo)

        def header():
            pdf.set_font(self.f, "B", 9)
            pdf.set_text_color(0, 122, 61)
            pdf.cell(0, 5, tit, new_x="LMARGIN", new_y="NEXT")
            pdf.set_draw_color(0, 146, 63)
            pdf.line(12, pdf.get_y(), pdf.w - 12, pdf.get_y())
            pdf.ln(3)
            pdf.set_text_color(28, 38, 33)

        def footer():
            pdf.set_y(-9)
            pdf.set_font(self.f, "", 7)
            pdf.set_text_color(95, 107, 101)
            pdf.cell(0, 4, self.t(f"{sub} · Generado {datetime.now():%d/%m/%Y %H:%M} · Página {pdf.page_no()}"), align="R")
        pdf.header, pdf.footer = header, footer
        pdf.add_page()

    def _fuente(self):
        base = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "fonts")
        for d in (base, "/usr/share/fonts/truetype/dejavu"):
            r, b = os.path.join(d, "DejaVuSans.ttf"), os.path.join(d, "DejaVuSans-Bold.ttf")
            if os.path.exists(r) and os.path.exists(b):
                self.pdf.add_font("DejaVu", "", r)
                self.pdf.add_font("DejaVu", "B", b)
                self.pdf.add_font("DejaVu", "I", r)
                self.f = "DejaVu"
                return True
        self.f = "Helvetica"
        return False

    def t(self, s):
        s = str(s or "")
        if self.unicode:
            return s
        rep = {"—": "-", "–": "-", "→": "->", "≥": ">=", "≤": "<=", "…": "...", "‘": "'", "’": "'", "“": '"', "”": '"',
               "•": "-", "⚑": "!", "✔": "OK", "×": "x"}
        for a, b in rep.items():
            s = s.replace(a, b)
        return s.encode("latin-1", "replace").decode("latin-1")

    # --- bloques ---
    def titulo_bloque(self, texto, nivel):
        tam = {1: 15, 2: 12.5, 3: 11, 4: 10}.get(nivel, 10)
        self.pdf.ln(1.5 if nivel > 1 else 0)
        self.pdf.set_font(self.f, "B", tam)
        self.pdf.set_text_color(0, 99, 47) if nivel <= 2 else self.pdf.set_text_color(28, 38, 33)
        self.pdf.multi_cell(0, tam * 0.5, self.t(texto), new_x="LMARGIN", new_y="NEXT")
        self.pdf.set_text_color(28, 38, 33)
        self.pdf.ln(0.8)

    def parrafo(self, texto, tam=9, estilo=""):
        if not texto:
            return
        self.pdf.set_font(self.f, estilo, tam)
        self.pdf.multi_cell(0, tam * 0.48, self.t(texto), new_x="LMARGIN", new_y="NEXT")
        self.pdf.ln(0.8)

    def indicadores(self, kpis):
        if not kpis:
            return
        pdf = self.pdf
        ancho = (pdf.w - 24) / min(len(kpis), 5) - 2
        for i, (valor, etiqueta, nota) in enumerate(kpis):
            if i and i % 5 == 0:
                pdf.ln(20)
            x, y = 12 + (i % 5) * (ancho + 2), pdf.get_y()
            pdf.set_draw_color(207, 213, 209)
            pdf.set_fill_color(247, 248, 247)
            pdf.rect(x, y, ancho, 18, style="DF")
            pdf.set_xy(x + 2, y + 1.5)
            pdf.set_font(self.f, "B", 11)
            pdf.cell(ancho - 4, 6, self.t(valor)[:40])
            pdf.set_xy(x + 2, y + 7.5)
            pdf.set_font(self.f, "", 7.5)
            pdf.multi_cell(ancho - 4, 3.3, self.t((etiqueta + (" · " + nota if nota else "")))[:120])
            pdf.set_xy(12, y)
        pdf.ln(21)

    def tabla(self, filas):
        filas = [f for f in filas if any((c[0] or "").strip() for c in f)]
        if not filas:
            return
        ncol = max(sum(cs for _, cs, _ in f) for f in filas)
        pdf = self.pdf
        pdf.set_font(self.f, "", 7.5 if ncol <= 8 else 6.5)
        pdf.set_draw_color(207, 213, 209)
        # ancho de columnas según el largo del texto (acotado)
        largo = [4] * ncol
        for f in filas:
            j = 0
            for txt, cs, _ in f:
                if cs == 1 and j < ncol:
                    largo[j] = max(largo[j], min(len(max((txt or "").split("\n"), key=len)), 40))
                j += cs
        total = sum(largo)
        anchos = [max(10.0, (pdf.w - 24) * x / total) for x in largo]
        k = (pdf.w - 24) / sum(anchos)
        anchos = [a * k for a in anchos]
        try:
            with pdf.table(col_widths=anchos, text_align="LEFT", line_height=pdf.font_size * 1.35,
                           headings_style=_estilo_cab(self.f), first_row_as_headings=any(c[2] for c in filas[0]),
                           borders_layout="SINGLE_TOP_LINE" if False else "HORIZONTAL_LINES") as t:
                for f in filas:
                    r = t.row()
                    usadas = 0
                    for txt, cs, _ in f:
                        cs = min(cs, ncol - usadas)
                        if cs <= 0:
                            break
                        r.cell(self.t(txt), colspan=cs)
                        usadas += cs
                    while usadas < ncol:
                        r.cell("")
                        usadas += 1
        except Exception:  # noqa: BLE001 - una tabla irregular se imprime como texto
            for f in filas:
                self.parrafo(" | ".join(c[0] for c in f), 7.5)
        pdf.ln(2.5)

    def grafico(self, svg):
        try:
            img = io.BytesIO(svg.encode("utf-8"))
            ancho = min(self.pdf.w - 24, 170)
            if self.pdf.get_y() > self.pdf.h - 80:
                self.pdf.add_page()
            self.pdf.image(img, x=12, w=ancho)
            self.pdf.ln(2)
        except Exception:  # noqa: BLE001 - si el gráfico no se puede dibujar se indica y se sigue
            self.parrafo("(Gráfico disponible en la pantalla del informe; su tabla de datos se incluye en el PDF.)", 7.5, "I")

    def salida(self):
        return bytes(self.pdf.output())


def _estilo_cab(fuente):
    from fpdf.fonts import FontFace
    return FontFace(family=fuente, emphasis="BOLD", fill_color=(235, 246, 239), color=(0, 99, 47))


def _filas_tabla(t):
    filas = []

    def recorrer(n):
        for h in n.hijos:
            if not isinstance(h, Nodo):
                continue
            if h.tag == "tr":
                fila = []
                for c in h.hijos:
                    if isinstance(c, Nodo) and c.tag in ("td", "th") and not (c.clases() & OMITIR_CLASES):
                        try:
                            cs = int(c.attrs.get("colspan") or 1)
                        except ValueError:
                            cs = 1
                        fila.append((_limpio(c.texto()), max(cs, 1), c.tag == "th"))
                if fila:
                    filas.append(fila)
            elif h.tag in ("thead", "tbody", "tfoot"):
                recorrer(h)
    recorrer(t)
    return filas


def _render(pdf, n):
    for h in n.hijos:
        if isinstance(h, str):
            txt = _limpio(h)
            if txt:
                pdf.parrafo(txt)
            continue
        if h.tag in OMITIR_TAGS or (h.clases() & OMITIR_CLASES) or h.attrs.get("hidden") is not None and h.tag != "section":
            continue
        cls = h.clases()
        if h.tag in ("h1", "h2", "h3", "h4"):
            pdf.titulo_bloque(_limpio(h.texto()), int(h.tag[1]))
        elif h.tag == "table":
            pdf.tabla(_filas_tabla(h))
        elif h.tag == "svg":
            if "ic" not in cls:
                pdf.grafico(_svg_texto(h))
        elif "legend" in cls:
            pdf.parrafo("Series: " + " · ".join(_limpio(x.texto()) for x in h.hijos if isinstance(x, Nodo)), 7.5)
        elif "kpis" in cls:
            kpis = []
            for k in h.hijos:
                if isinstance(k, Nodo) and "kpi" in k.clases():
                    b = _buscar(k, lambda x: x.tag == "b")
                    s = _buscar(k, lambda x: x.tag == "span")
                    sm = _buscar(k, lambda x: x.tag == "small")
                    kpis.append((_limpio(b.texto()) if b else "", _limpio(s.texto()) if s else "", _limpio(sm.texto()) if sm else ""))
            pdf.indicadores(kpis)
        elif h.tag in ("ul", "ol"):
            for i, li in enumerate([x for x in h.hijos if isinstance(x, Nodo) and x.tag == "li"], start=1):
                pdf.parrafo(("• " if h.tag == "ul" else f"{i}. ") + _limpio(li.texto()))
        elif h.tag == "p" or ("reco" in cls) or ("note" in cls):
            if any(isinstance(x, Nodo) and x.tag in ("table", "svg") for x in h.hijos):
                _render(pdf, h)
            else:
                pdf.parrafo(_limpio(h.texto()), 8.5 if "small" in cls else 9)
        elif any(isinstance(x, Nodo) and (x.tag in BLOQUES or x.clases() & {"kpis", "card", "grid"}) for x in h.hijos):
            _render(pdf, h)
        else:
            txt = _limpio(h.texto())
            if txt:
                pdf.parrafo(txt, 8.5 if "small" in cls or h.tag == "small" else 9, "B" if h.tag == "b" else "")


def html_to_pdf(html, titulo, subtitulo="Corporación Productiva · Módulo de Operaciones"):
    """PDF del contenido principal (<main class="content">) de una pantalla de informe."""
    p = _Arbol()
    p.feed(html)
    main = _buscar(p.raiz, lambda n: n.tag == "main") or p.raiz
    pdf = _PDF(titulo, subtitulo)
    _render(pdf, main)
    return pdf.salida()


def table_pdf(titulo, columnas, filas, notas=None, subtitulo="Corporación Productiva · Módulo de Operaciones"):
    """PDF de una tabla simple (informe personalizado)."""
    pdf = _PDF(titulo, subtitulo)
    pdf.titulo_bloque(titulo, 1)
    for n in notas or []:
        pdf.parrafo(n, 8.5)
    pdf.tabla([[(c, 1, True) for c in columnas]] + [[("" if v is None else str(v), 1, False) for v in f] for f in filas])
    return pdf.salida()


def html_to_xlsx(html, titulo):
    """Excel con las tablas del informe (una hoja por tabla, con el título de su sección)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    p = _Arbol()
    p.feed(html)
    main = _buscar(p.raiz, lambda n: n.tag == "main") or p.raiz
    wb = Workbook()
    wb.remove(wb.active)
    ultimo = [titulo]
    usados = set()

    def recorrer(n):
        for h in n.hijos:
            if not isinstance(h, Nodo) or h.tag in OMITIR_TAGS or (h.clases() & OMITIR_CLASES):
                continue
            if h.tag in ("h1", "h2", "h3"):
                ultimo[0] = _limpio(h.texto()) or ultimo[0]
            elif h.tag == "table":
                filas = _filas_tabla(h)
                if not filas:
                    continue
                nombre = re.sub(r"[\[\]:*?/\\]", " ", ultimo[0])[:28].strip() or "Tabla"
                base, k = nombre, 2
                while nombre in usados:
                    nombre = f"{base[:25]} {k}"
                    k += 1
                usados.add(nombre)
                ws = wb.create_sheet(nombre)
                ws.append([ultimo[0]])
                ws["A1"].font = Font(bold=True, size=12)
                for f in filas:
                    vals = []
                    for txt, cs, _ in f:
                        vals.append(_num_o_texto(txt))
                        vals += [""] * (cs - 1)
                    ws.append(vals)
                    if f and f[0][2]:
                        for c in ws[ws.max_row]:
                            c.font = Font(bold=True)
                            c.fill = PatternFill("solid", fgColor="EBF6EF")
                for col in ws.columns:
                    ws.column_dimensions[col[0].column_letter].width = min(45, max(10, max(len(str(c.value or "")) for c in col) + 2))
            else:
                recorrer(h)
    recorrer(main)
    if not wb.sheetnames:
        wb.create_sheet("Informe").append([titulo, "Sin tablas en el periodo."])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _num_o_texto(t):
    s = (t or "").strip()
    m = re.fullmatch(r"-?[\d ]+([.,]\d+)?", s.replace(" ", " "))
    if m and any(ch.isdigit() for ch in s):
        try:
            return float(s.replace(" ", "").replace(",", "."))
        except ValueError:
            return s
    return s
