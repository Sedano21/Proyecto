"""Cadena de suministro (ronda UX del 24-oct-2026).

- Aviso automático a Comercial 48 h y 24 h antes de la fecha comprometida cuando la OV corre riesgo de no completarse
  (stock incompleto o compromiso NO CUMPLE).
- El Jefe de Logística avisa, el mismo día o al día siguiente de creada la OV, que no llegará a tiempo (motivo obligatorio).
- Proyección de demanda cargada por Comercial y consultada por Logística frente al stock libre.
- Stock sin rotación: aviso mensual a Comercial, Almacén y Gerencia con los productos sin salidas en el periodo configurado.
- Movimientos punto a punto para el montacarguista: cualquier área los solicita; se ejecutan con foto de inicio y foto de
  validación al terminar.
"""
from collections import defaultdict
from datetime import date, datetime, timedelta

from sqlalchemy import func, select

from ..db import db
from ..models import Alert, DemandForecast, MoveRequest, Product, SalesHistory, SalesOrder, StockLot, now
from . import alerts
from .core import BusinessError, audit, has_perm, param


def _avisado(key):
    """Cada aviso se da una sola vez (aunque ya lo hayan cerrado): las verificaciones corren cada minuto."""
    return db().execute(select(Alert.id).where(Alert.dedup_key == key)).first() is not None


# ------------------------------------------------------------------ OV en riesgo y aviso de atraso
def check_order_risk(ahora=None):
    """OV abiertas con fecha comprometida dentro de 48 h / 24 h (configurable) y riesgo de no completarse."""
    from .stock import ORDEN_ESTADOS_CERRADOS
    ahora = ahora or datetime.now()
    horas = sorted({int(h) for h in (param("horas_aviso_riesgo_ov", [48, 24]) or []) if h}, reverse=True)
    if not horas:
        return 0
    limite = (ahora + timedelta(hours=max(horas))).date()
    n = 0
    q = select(SalesOrder).where(SalesOrder.fecha_compromiso >= ahora.date(), SalesOrder.fecha_compromiso <= limite,
                                 SalesOrder.estado.notin_(ORDEN_ESTADOS_CERRADOS + ("PREPARADA", "EN_RUTA", "BORRADOR")))
    for o in db().execute(q).scalars():
        compromiso = (o.compromiso or {}).get("estado") if isinstance(o.compromiso, dict) else None
        riesgo = []
        if o.estado_stock != "COMPLETA":
            riesgo.append(f"stock {o.estado_stock.replace('_', ' ').lower()} ({o.cobertura_unidades}% cubierto)")
        if compromiso == "NO_CUMPLE":
            riesgo.append("Operaciones respondió NO CUMPLE")
        if o.aviso_atraso:
            riesgo.append(f"Logística avisó atraso: {o.aviso_atraso.get('motivo', '')}")
        if not riesgo:
            continue
        faltan = (datetime.combine(o.fecha_compromiso, datetime.min.time()) - ahora).total_seconds() / 3600
        dentro = [h for h in horas if faltan <= h]
        if not dentro:
            continue
        umbral = min(dentro)  # el aviso de 24 h reemplaza al de 48 h
        if _avisado(f"OV_RIESGO:{o.id}:{umbral}"):
            continue
        alerts.emit("OV_RIESGO_DESPACHO", f"OV {o.numero}: riesgo de no completar ({umbral} h para el {o.fecha_compromiso:%d/%m/%Y})",
                    f"{o.customer.nombre if o.customer else ''}. " + "; ".join(riesgo) + ".", "SalesOrder", o.id,
                    prioridad="ALTA" if umbral <= 24 else None,
                    extra_users=[o.vendedor_id] if o.vendedor_id else None, dedup_key=f"OV_RIESGO:{o.id}:{umbral}")
        n += 1
    return n


def report_delay(user, order, motivo, nueva_fecha=None, hoy=None):
    """El Jefe de Logística avisa que la OV no llegará a tiempo: solo el día de creación o el siguiente, con motivo."""
    hoy = hoy or date.today()
    if not has_perm(user, "ov.reportar_atraso"):
        raise BusinessError("El aviso de atraso lo registra el Jefe de Logística.")
    creada = order.fecha_emision or (order.creado_ts.date() if getattr(order, "creado_ts", None) else hoy)
    if hoy > creada + timedelta(days=1):
        raise BusinessError(f"El aviso se registra el día de creación de la OV o el siguiente (creada el {creada:%d/%m/%Y}).")
    if not (motivo or "").strip():
        raise BusinessError("Indique el motivo por el que la OV no llegará a tiempo.")
    if isinstance(nueva_fecha, str):
        nueva_fecha = date.fromisoformat(nueva_fecha) if nueva_fecha else None
    if nueva_fecha is not None and nueva_fecha <= order.fecha_compromiso:
        raise BusinessError("La fecha estimada debe ser posterior a la fecha comprometida.")
    order.aviso_atraso = {"motivo": motivo.strip(), "user": user.nombre, "ts": now().isoformat(timespec="minutes"),
                          "nueva_fecha": nueva_fecha.isoformat() if nueva_fecha else ""}
    alerts.emit("OV_ATRASO_LOGISTICA", f"Logística: la OV {order.numero} no llegará a tiempo",
                f"Motivo: {motivo.strip()}." + (f" Fecha estimada: {nueva_fecha:%d/%m/%Y}." if nueva_fecha else "")
                + " Coordine con el cliente.", "SalesOrder", order.id, prioridad="ALTA",
                extra_users=[order.vendedor_id] if order.vendedor_id else None, creado_por=user)
    audit(user, "OV_AVISO_ATRASO", "SalesOrder", order.numero, order.aviso_atraso)
    return order


# ------------------------------------------------------------------ proyección de demanda
def _mes(v):
    v = str(v or "").strip()[:7]
    try:
        datetime.strptime(v, "%Y-%m")
    except ValueError:
        raise BusinessError(f"Mes inválido «{v}» (AAAA-MM).")
    return v


def save_forecast(user, mes, product_code, cantidad, canal="", customer_id=None, observacion=""):
    if not has_perm(user, "demanda.cargar"):
        raise BusinessError("La proyección de demanda la carga Comercial.")
    mes = _mes(mes)
    code = (product_code or "").strip().upper().split(" ")[0]
    if db().get(Product, code) is None:
        raise BusinessError(f"Producto {code or '(vacío)'} no existe.")
    try:
        cantidad = float(str(cantidad).replace(",", ""))
    except (TypeError, ValueError):
        raise BusinessError(f"{code}: cantidad inválida.")
    if cantidad < 0:
        raise BusinessError(f"{code}: la cantidad no puede ser negativa.")
    canal = (canal or "").strip().upper()
    q = select(DemandForecast).where(DemandForecast.mes == mes, DemandForecast.product_code == code,
                                     DemandForecast.canal == canal, DemandForecast.customer_id.is_(customer_id) if customer_id
                                     is None else DemandForecast.customer_id == customer_id)
    f = db().execute(q).scalars().first()
    obs = (observacion or "").strip()[:250]
    if f is not None and param("demanda_cambio_requiere_gg", True) and abs(f.cantidad - cantidad) > 1e-9:
        # V7: modificar una proyección ya registrada requiere la autorización de Gerencia General (con justificación)
        if not obs:
            raise BusinessError(f"{code} ({mes}): para modificar una proyección registrada indique la justificación.")
        f.cambio = {"estado": "PENDIENTE", "cantidad": cantidad, "anterior": f.cantidad, "observacion": obs,
                    "user": user.nombre, "user_id": user.id, "ts": now().isoformat(timespec="minutes")}
        db().flush()
        alerts.emit("DEMANDA_CAMBIO", f"Cambio de proyección {code} {mes}: {f.cantidad:g} → {cantidad:g}",
                    f"Solicitado por {user.nombre}. Justificación: {obs}", "DemandForecast", f.id,
                    dedup_key=f"DEMANDA_CAMBIO:{f.id}", creado_por=user)
        audit(user, "DEMANDA_CAMBIO_SOLICITADO", "DemandForecast", f.id, f.cambio)
        return f
    if f is None:
        f = DemandForecast(mes=mes, product_code=code, canal=canal, customer_id=customer_id, cantidad=cantidad)
        db().add(f)
    f.cantidad, f.observacion, f.user_id = cantidad, obs, user.id
    db().flush()
    return f


def save_forecast_group(user, mes, filas):
    """V7: registro por grupo («+» filas) y envío completo de una sola vez. Se valida todo antes de guardar: si una fila
    tiene error no se guarda ninguna. filas: [{product_code, cantidad, canal, customer_id, observacion}]."""
    if not has_perm(user, "demanda.cargar"):
        raise BusinessError("La proyección de demanda la carga Comercial.")
    mes = _mes(mes)
    validas, errores, vistos = [], [], set()
    for i, r in enumerate(filas, start=1):
        code = (r.get("product_code") or "").strip().upper().split(" ")[0]
        if not code and not str(r.get("cantidad") or "").strip():
            continue
        if db().get(Product, code) is None:
            errores.append(f"Fila {i}: producto {code or '(vacío)'} no existe.")
            continue
        try:
            q = float(str(r.get("cantidad")).replace(",", ""))
            if q < 0:
                raise ValueError
        except (TypeError, ValueError):
            errores.append(f"Fila {i}: cantidad inválida.")
            continue
        cid = int(r["customer_id"]) if str(r.get("customer_id") or "").isdigit() else None
        clave = (code, (r.get("canal") or "").strip().upper(), cid)
        if clave in vistos:
            errores.append(f"Fila {i}: {code} repetido con el mismo canal y cliente.")
            continue
        vistos.add(clave)
        validas.append(dict(r, product_code=code, cantidad=q, customer_id=cid))
    if errores:
        raise BusinessError(" ".join(errores[:8]))
    if not validas:
        raise BusinessError("Agregue al menos un producto con su cantidad.")
    res = {"nuevas": 0, "cambios": 0}
    for r in validas:
        f = save_forecast(user, mes, r["product_code"], r["cantidad"], r.get("canal", ""), r["customer_id"],
                          r.get("observacion", ""))
        c = f.cambio or {}
        pendiente = c.get("estado") == "PENDIENTE" and abs(float(c.get("cantidad", -1)) - r["cantidad"]) < 1e-9 \
            and abs(f.cantidad - r["cantidad"]) > 1e-9
        res["cambios" if pendiente else "nuevas"] += 1
    audit(user, "DEMANDA_GRUPO", "DemandForecast", mes, res)
    return res


def decide_change(f, user, aprobar, comentario=""):
    """Gerencia General autoriza (o no) el cambio de una proyección registrada."""
    if not has_perm(user, "demanda.aprobar_cambio"):
        raise BusinessError("Los cambios de la proyección los autoriza Gerencia General.")
    c = dict(f.cambio or {})
    if c.get("estado") != "PENDIENTE":
        raise BusinessError("No hay un cambio pendiente.")
    if not aprobar and not (comentario or "").strip():
        raise BusinessError("Indique el motivo para no autorizar el cambio.")
    if aprobar:
        f.cantidad = float(c["cantidad"])
        f.observacion = c.get("observacion", "")[:250]
    c.update({"estado": "APROBADO" if aprobar else "RECHAZADO", "decidio": user.nombre,
              "decision_ts": now().isoformat(timespec="minutes"), "comentario": (comentario or "").strip()})
    f.cambio = c
    alerts.close_by_system(f"DEMANDA_CAMBIO:{f.id}", f"{'Autorizado' if aprobar else 'No autorizado'} por {user.nombre}")
    audit(user, "DEMANDA_CAMBIO_" + ("APROBADO" if aprobar else "RECHAZADO"), "DemandForecast", f.id, c)


def pending_changes():
    return [f for f in db().execute(select(DemandForecast)).scalars() if (f.cambio or {}).get("estado") == "PENDIENTE"]


def import_forecast(user, path):
    """Excel con columnas: mes (AAAA-MM), producto (código), cantidad y, opcionales, canal y observación."""
    from openpyxl import load_workbook
    if not has_perm(user, "demanda.cargar"):
        raise BusinessError("La proyección de demanda la carga Comercial.")
    ws = load_workbook(path, data_only=True, read_only=True).active
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        raise BusinessError("El archivo está vacío.")
    cab = [" ".join(str(c or "").strip().lower().split()) for c in filas[0]]

    def col(*n):
        return next((i for i, c in enumerate(cab) if c in n), None)
    im, ip, iq = col("mes", "periodo"), col("producto", "codigo", "código", "articulo", "artículo"), col("cantidad", "proyeccion", "proyección")
    ic, io = col("canal"), col("observacion", "observación")
    if None in (im, ip, iq):
        raise BusinessError("Faltan columnas: mes, producto y cantidad.")
    n, errores = 0, []
    for k, r in enumerate(filas[1:], start=2):
        if not any(r):
            continue
        mes = r[im].strftime("%Y-%m") if isinstance(r[im], (date, datetime)) else r[im]
        try:
            save_forecast(user, mes, r[ip], r[iq], r[ic] if ic is not None else "", None, r[io] if io is not None else "")
            n += 1
        except BusinessError as e:
            errores.append(f"Fila {k}: {e}")
    audit(user, "DEMANDA_CARGADA", "DemandForecast", "", {"filas": n, "errores": len(errores)})
    return n, errores


def forecast_view(desde_mes=None, meses=3):
    """Proyección por producto y mes frente al stock libre (para la previsión de Logística)."""
    from .stock import free_qty
    hoy = date.today()
    desde = _mes(desde_mes) if desde_mes else f"{hoy:%Y-%m}"
    y, m = int(desde[:4]), int(desde[5:])
    lista = []
    for i in range(meses):
        mm = (m - 1 + i) % 12 + 1
        lista.append(f"{y + (m - 1 + i) // 12}-{mm:02d}")
    rows = list(db().execute(select(DemandForecast).where(DemandForecast.mes.in_(lista))).scalars())
    por = defaultdict(lambda: defaultdict(float))
    for f in rows:
        por[f.product_code][f.mes] += f.cantidad
    libre = defaultdict(float)
    for lot in db().execute(select(StockLot).where(StockLot.product_code.in_(list(por)), StockLot.cantidad > 0,
                                                   StockLot.estado == "DISPONIBLE")).scalars():
        libre[lot.product_code] += free_qty(lot)
    filas = []
    for code, ms in sorted(por.items()):
        total = sum(ms.values())
        p = db().get(Product, code)
        filas.append({"code": code, "descripcion": p.descripcion if p else "", "meses": ms, "total": total,
                      "libre": round(libre[code], 2), "cobertura": round(100 * libre[code] / total, 1) if total else None,
                      "faltante": round(max(total - libre[code], 0), 2)})
    return {"meses": lista, "filas": filas, "registros": rows}


# ------------------------------------------------------------------ stock sin rotación
def no_rotation(dias=None, hoy=None):
    hoy = hoy or date.today()
    dias = int(dias or param("dias_sin_rotacion", 90) or 90)
    desde = hoy - timedelta(days=dias)
    stock = defaultdict(float)
    for lot in db().execute(select(StockLot).where(StockLot.cantidad > 0)).scalars():
        stock[lot.product_code] += lot.cantidad
    ultima = dict(db().execute(select(SalesHistory.product_code, func.max(SalesHistory.fecha))
                               .group_by(SalesHistory.product_code)).all())
    out = []
    for code, q in sorted(stock.items()):
        u = ultima.get(code)
        if u is None or u < desde:
            p = db().get(Product, code)
            out.append({"code": code, "descripcion": p.descripcion if p else "", "tipo": p.tipo if p else "", "stock": round(q, 2),
                        "ultima_salida": u, "dias": (hoy - u).days if u else None})
    return {"dias": dias, "filas": out}


def check_no_rotation(hoy=None):
    """Aviso mensual (una vez por mes) a Comercial, Almacén y Gerencia."""
    hoy = hoy or date.today()
    if hoy.day < int(param("dia_aviso_sin_rotacion", 1) or 1) or _avisado(f"SIN_ROT:{hoy:%Y-%m}"):
        return 0
    d = no_rotation(hoy=hoy)
    if not d["filas"]:
        return 0
    top = ", ".join(f"{x['code']} ({x['stock']:g})" for x in d["filas"][:10])
    alerts.emit("STOCK_SIN_ROTACION", f"{len(d['filas'])} producto(s) sin rotación en {d['dias']} días ({hoy:%m/%Y})",
                f"Productos con stock y sin salidas: {top}{'…' if len(d['filas']) > 10 else ''}. Revise el informe de rotación.",
                "Product", "", dedup_key=f"SIN_ROT:{hoy:%Y-%m}")
    return len(d["filas"])


# ------------------------------------------------------------------ movimientos punto a punto (montacarguista)
def _area(user):
    from ..models import Role
    r = db().get(Role, user.role_code)
    return r.area if r else ""


EJECUTORES_MOV = ("MONTACARGUISTA", "AUX_ALMACEN", "AUX_DESPACHO")


def move_executors():
    """V7: quién puede ejecutar el movimiento: montacarguista, auxiliar de almacén o auxiliar de distribución."""
    from ..models import User
    return list(db().execute(select(User).where(User.role_code.in_(EJECUTORES_MOV), User.activo.is_(True))
                             .order_by(User.role_code, User.nombre)).scalars())


def move_points(almacen=""):
    """A03: lista configurable de puntos por almacén (zonas, racks, líneas, despacho); siempre admite «otro»."""
    cfg = param("puntos_movimiento", {}) or {}
    return list(dict.fromkeys(list(cfg.get(almacen or "", [])) + list(cfg.get("*", []))))


TIPOS_CARGA = {"PALLET_PT": "Pallet de PT", "MERMA": "Merma", "REMANENTE": "Remanente de MP", "MUESTRA": "Muestra",
               "PALLET": "Pallet (vacío)"}


def move_cargo(datos):
    """V12: valida el tipo de carga y sus datos y arma el detalle legible.
    PALLET_PT: n.° de pallets y código de PT · MERMA: kg · REMANENTE: kg y código de MP ·
    MUESTRA: PT, cantidad y cliente al que se envía · PALLET: cantidad."""
    from ..models import Customer, Product
    from .stock import product_kind
    tipo = (datos.get("tipo") or "").strip().upper()
    if not tipo:
        return {}, ""
    if tipo not in TIPOS_CARGA:
        raise BusinessError("Elija qué se mueve de la lista.")

    def num(k, txt, entero=False):
        try:
            v = float(datos.get(k))
        except (TypeError, ValueError):
            raise BusinessError(f"Indique {txt}.")
        if v <= 0 or (entero and v != int(v)):
            raise BusinessError(f"{txt[0].upper() + txt[1:]}: debe ser {'un entero ' if entero else ''}mayor que cero.")
        return int(v) if entero else v

    def prod(k, tipos, txt):
        code = (datos.get(k) or "").strip().upper().split(" ")[0]
        p = db().get(Product, code) if code else None
        if p is None:
            raise BusinessError(f"Elija el {txt} de la lista.")
        if tipos and product_kind(p) not in tipos:
            raise BusinessError(f"{code} no es {txt}.")
        return p
    c = {"tipo": tipo}
    if tipo == "PALLET_PT":
        c["cantidad"] = num("cantidad", "la cantidad de pallets", True)
        p = prod("producto", ("PT",), "producto terminado")
        c["producto"] = p.code
        det = f"{c['cantidad']} pallet(s) de PT {p.code} {p.descripcion}"
    elif tipo == "MERMA":
        c["kg"] = num("kg", "el peso en kg")
        det = f"Merma {c['kg']:g} kg"
    elif tipo == "REMANENTE":
        c["kg"] = num("kg", "el peso en kg")
        p = prod("mp", ("MP", "GR"), "materia prima")
        c["mp"] = p.code
        det = f"Remanente de MP {p.code} {p.descripcion}: {c['kg']:g} kg"
    elif tipo == "MUESTRA":
        p = prod("producto", ("PT",), "producto terminado")
        c["producto"] = p.code
        c["cantidad"] = num("cantidad", "la cantidad de la muestra")
        cid = datos.get("customer_id")
        cu = db().get(Customer, int(cid)) if str(cid or "").isdigit() else None
        if cu is None:
            raise BusinessError("Indique a qué cliente se enviará la muestra.")
        c["customer_id"] = cu.id
        det = f"Muestra de {p.code} {p.descripcion} × {c['cantidad']:g} para {cu.nombre}"
    else:
        c["cantidad"] = num("cantidad", "la cantidad de pallets", True)
        det = f"{c['cantidad']} pallet(s)"
    return c, det


def request_move(user, origen, destino, detalle="", cantidad="", prioridad="MEDIA", almacen="", observacion="",
                 programado_para=None, asignado_id=None, aviso_min=None, carga=None):
    if not has_perm(user, "movimiento.solicitar"):
        raise BusinessError("Su rol no solicita movimientos al montacarguista.")
    carga_d, det = move_cargo(carga or {})
    if det:
        detalle = det + ((" · " + detalle.strip()) if (detalle or "").strip() else "")
        cantidad = str(carga_d.get("cantidad") or carga_d.get("kg") or cantidad or "")
    if isinstance(programado_para, str):
        try:
            programado_para = datetime.fromisoformat(programado_para) if programado_para.strip() else None
        except ValueError:
            raise BusinessError("Fecha y hora programadas inválidas.")
    if programado_para is not None and programado_para < datetime.now() - timedelta(minutes=1):
        raise BusinessError("La fecha y hora programadas ya pasaron.")
    ejecutor = None
    if asignado_id:
        from ..models import User
        ejecutor = db().get(User, int(asignado_id))
        if ejecutor is None or ejecutor.role_code not in EJECUTORES_MOV:
            raise BusinessError("Asigne el movimiento a un montacarguista, auxiliar de almacén o auxiliar de distribución.")
    origen, destino = (origen or "").strip(), (destino or "").strip()
    if not origen or not destino:
        raise BusinessError("Indique el punto de origen y el de destino.")
    if origen.upper() == destino.upper():
        raise BusinessError("Origen y destino son el mismo punto.")
    prioridad = (prioridad or "MEDIA").upper()
    if prioridad not in ("ALTA", "MEDIA", "BAJA"):
        raise BusinessError("Prioridad no válida.")
    m = MoveRequest(area=_area(user), almacen=(almacen or user.almacen or "").strip(), origen=origen, destino=destino,
                    detalle=(detalle or "").strip(), cantidad=(cantidad or "").strip(), prioridad=prioridad,
                    observacion=(observacion or "").strip(), solicitado_por=user.id,
                    tipo="PROGRAMADO" if programado_para else "INMEDIATO", programado_para=programado_para,
                    asignado_a=ejecutor.id if ejecutor else None, carga=carga_d,
                    aviso_min=_aviso(aviso_min) if programado_para else None)
    db().add(m)
    db().flush()
    cuando = f" Programado para el {programado_para:%d/%m %H:%M}." if programado_para else " Inmediato."
    alerts.emit("MOVIMIENTO_SOLICITADO", f"Movimiento #{m.id}: {origen} → {destino}",
                f"{m.area}: {m.detalle or 'sin detalle'} {m.cantidad}. Prioridad {prioridad.lower()}.{cuando}"
                + (f" Asignado a {ejecutor.nombre}." if ejecutor else ""), "MoveRequest", m.id,
                prioridad="ALTA" if prioridad == "ALTA" else None, extra_users=[ejecutor.id] if ejecutor else None,
                dedup_key=f"MOV:{m.id}", creado_por=user)
    audit(user, "MOVIMIENTO_SOLICITADO", "MoveRequest", m.id, {"origen": origen, "destino": destino})
    return m


def reschedule_move(user, m, fecha, hora, motivo=""):
    """V12: quien solicitó un movimiento programado puede cambiar el día y la hora."""
    if m.solicitado_por != user.id:
        raise BusinessError("Solo quien solicitó el movimiento lo reprograma.")
    if m.estado != "SOLICITADA":
        raise BusinessError("El movimiento ya se inició o terminó.")
    nuevo = _when(fecha, hora)
    if nuevo is None:
        raise BusinessError("Indique el día y la hora.")
    if nuevo < datetime.now() - timedelta(minutes=1):
        raise BusinessError("La nueva fecha y hora ya pasaron.")
    antes = m.programado_para
    m.programado_para, m.tipo, m.avisado_ts = nuevo, "PROGRAMADO", None
    m.historial = list(m.historial or []) + [{"evento": "REPROGRAMADO", "de": antes.isoformat(timespec="minutes") if antes else "",
                                              "a": nuevo.isoformat(timespec="minutes"), "motivo": (motivo or "").strip(),
                                              "user": user.nombre, "ts": datetime.now().isoformat(timespec="minutes")}]
    audit(user, "MOVIMIENTO_REPROGRAMADO", "MoveRequest", m.id, {"de": str(antes), "a": str(nuevo)})
    return m


def _when(fecha, hora):
    """Día (obligatorio) y hora (editable; por defecto la del parámetro) del movimiento programado."""
    if not (fecha or "").strip():
        return None
    try:
        d = date.fromisoformat(fecha.strip()[:10])
        h = (hora or param("hora_movimiento_programado", "08:00") or "08:00").strip()[:5]
        hh, mm = [int(x) for x in h.split(":")]
        return datetime.combine(d, datetime.min.time()).replace(hour=hh, minute=mm)
    except (ValueError, TypeError):
        raise BusinessError("Fecha u hora programadas inválidas.")


def _aviso(v):
    """A04: la anticipación del aviso varía: la elige quien solicita (por defecto el parámetro)."""
    try:
        v = int(v) if v not in (None, "") else int(param("aviso_movimiento_min", 30) or 30)
    except ValueError:
        raise BusinessError("Anticipación del aviso inválida.")
    if not 0 <= v <= 24 * 60:
        raise BusinessError("La anticipación del aviso va de 0 a 1440 minutos.")
    return v


def reassign_move(user, m, nuevo_id, motivo=""):
    """A05: solo quien solicitó reasigna (o deja libre para el primero disponible); queda el rastro del cambio."""
    from ..models import User
    if m.solicitado_por != user.id:
        raise BusinessError("Solo quien solicitó el movimiento puede reasignarlo.")
    if m.estado != "SOLICITADA":
        raise BusinessError("Solo se reasigna un movimiento que aún no empezó.")
    nuevo = db().get(User, int(nuevo_id)) if nuevo_id else None
    if nuevo is not None and nuevo.role_code not in EJECUTORES_MOV:
        raise BusinessError("Asigne a un montacarguista, auxiliar de almacén o auxiliar de distribución.")
    if (nuevo.id if nuevo else None) == m.asignado_a:
        raise BusinessError("Ya está asignado a esa persona.")
    antes = m.operador.nombre if m.operador else "primero disponible"
    m.historial = list(m.historial or []) + [{"ts": now().isoformat(timespec="minutes"), "user": user.nombre, "de": antes,
                                              "a": nuevo.nombre if nuevo else "primero disponible",
                                              "motivo": (motivo or "").strip()}]
    m.asignado_a = nuevo.id if nuevo else None
    alerts.emit("MOVIMIENTO_SOLICITADO", f"Movimiento #{m.id} reasignado: {m.origen} → {m.destino}",
                f"De {antes} a {nuevo.nombre if nuevo else 'el primero disponible'}. {motivo}", "MoveRequest", m.id,
                extra_users=[nuevo.id] if nuevo else None, creado_por=user)
    audit(user, "MOVIMIENTO_REASIGNADO", "MoveRequest", m.id, m.historial[-1])
    return m


def check_scheduled_moves(ahora=None):
    """A04: aviso a quien ejecuta (o a todos los ejecutores si no hay asignado) con la anticipación elegida."""
    ahora = ahora or datetime.now()
    n = 0
    for m in db().execute(select(MoveRequest).where(MoveRequest.estado == "SOLICITADA", MoveRequest.tipo == "PROGRAMADO",
                                                    MoveRequest.avisado_ts.is_(None))).scalars():
        if m.programado_para and m.programado_para - timedelta(minutes=m.aviso_min or 0) <= ahora:
            alerts.emit("MOVIMIENTO_SOLICITADO", f"Movimiento programado #{m.id} a las {m.programado_para:%H:%M}: {m.origen} → {m.destino}",
                        f"{m.detalle or ''} {m.cantidad or ''}".strip() or "Prepárese para ejecutarlo.", "MoveRequest", m.id,
                        prioridad="ALTA", extra_users=[m.asignado_a] if m.asignado_a else None,
                        dedup_key=f"MOV_PROG:{m.id}")
            m.avisado_ts = now()
            n += 1
    db().flush()
    return n


def start_move(user, m, foto):
    if not has_perm(user, "movimiento.ejecutar"):
        raise BusinessError("Ejecuta el movimiento el montacarguista.")
    if m.estado != "SOLICITADA":
        raise BusinessError("El movimiento ya fue tomado o cerrado.")
    if m.asignado_a and m.asignado_a != user.id:
        raise BusinessError(f"El movimiento está asignado a {m.operador.nombre}.")
    if not foto:
        raise BusinessError("Tome la foto del punto de origen antes de iniciar.")
    m.estado, m.asignado_a, m.inicio_ts, m.foto_inicio = "EN_CURSO", user.id, now(), foto
    audit(user, "MOVIMIENTO_INICIADO", "MoveRequest", m.id, {})
    return m


def finish_move(user, m, foto, observacion=""):
    if m.estado != "EN_CURSO" or m.asignado_a != user.id:
        raise BusinessError("Solo quien inició el movimiento lo termina.")
    if not foto:
        raise BusinessError("La foto de validación en el destino es obligatoria.")
    m.estado, m.fin_ts, m.foto_fin = "COMPLETADA", now(), foto
    if observacion:
        m.observacion = ((m.observacion + "\n") if m.observacion else "") + observacion.strip()
    alerts.close_by_system(f"MOV:{m.id}", f"Movimiento completado por {user.nombre}")
    alerts.emit("MOVIMIENTO_COMPLETADO", f"Movimiento #{m.id} completado: {m.origen} → {m.destino}",
                f"Validado con foto por {user.nombre}.", "MoveRequest", m.id,
                extra_users=[m.solicitado_por] if m.solicitado_por else None, creado_por=user)
    audit(user, "MOVIMIENTO_COMPLETADO", "MoveRequest", m.id, {})
    return m


def cancel_move(user, m, motivo):
    if m.solicitado_por != user.id:  # V7: solo quien la creó elimina la solicitud
        raise BusinessError("Solo quien creó la solicitud puede eliminarla.")
    if m.estado == "COMPLETADA":
        raise BusinessError("El movimiento ya se completó.")
    if not (motivo or "").strip():
        raise BusinessError("Indique el motivo de la anulación.")
    m.estado = "ANULADA"
    m.observacion = ((m.observacion + "\n") if m.observacion else "") + f"Anulado: {motivo.strip()}"
    alerts.close_by_system(f"MOV:{m.id}", "Movimiento anulado")
    audit(user, "MOVIMIENTO_ANULADO", "MoveRequest", m.id, {"motivo": motivo})
    return m


def moves(user):
    """Tablero: el montacarguista ve las solicitudes de todas las áreas; los demás, las suyas y las de su área."""
    rows = list(db().execute(select(MoveRequest).where(
        (MoveRequest.estado.in_(["SOLICITADA", "EN_CURSO"])) | (MoveRequest.ts >= datetime.now() - timedelta(days=7)))
        .order_by(MoveRequest.estado != "SOLICITADA", MoveRequest.prioridad != "ALTA", MoveRequest.ts.desc())).scalars())
    if has_perm(user, "movimiento.ejecutar") or has_perm(user, "movimiento.ver_todos"):
        return rows
    area = _area(user)
    return [m for m in rows if m.solicitado_por == user.id or m.area == area]


# ------------------------------------------------------------------ V7: ranking de stock sin rotación
def _duracion(dias):
    if dias is None:
        return "—"
    sem, mes = dias / 7.0, dias / 30.4
    return f"{dias} días · {sem:.1f} semanas · {mes:.1f} meses"


def rotation_ranking(dias=None, hoy=None, top=20, almacen=None):
    """Dos frentes: TIEMPO (más días sin salida) y CANTIDAD (más kg inmovilizados sin salida en el periodo), top 20 de
    cada uno, con días, semanas y meses, y una recomendación de venta por producto."""
    from .costs import cost_per_kg
    hoy = hoy or date.today()
    base = no_rotation(dias, hoy)
    lots = defaultdict(list)
    q = select(StockLot).where(StockLot.cantidad > 0)
    if almacen:  # A02: consolidado, con filtro por almacén
        q = q.where(StockLot.warehouse_code == almacen)
    for lot in db().execute(q).scalars():
        lots[lot.product_code].append(lot)
    compradores = defaultdict(lambda: defaultdict(float))
    for o in db().execute(select(SalesOrder).where(SalesOrder.estado != "CANCELADA",
                                                   SalesOrder.fecha_emision >= hoy - timedelta(days=365))).scalars():
        for l in o.lines:
            if l.activa:
                compradores[l.product_code][o.customer.nombre] += l.cantidad
    filas = []
    for x in base["filas"]:
        p = db().get(Product, x["code"])
        ls = lots.get(x["code"], [])
        if almacen:
            if not ls:
                continue
            x = dict(x, stock=round(sum(l.cantidad for l in ls), 2))
        ingreso = min([l.fecha_admision for l in ls if l.fecha_admision] or [None]) if ls else None
        dias_sin = x["dias"] if x["dias"] is not None else ((hoy - ingreso).days if ingreso else None)
        kg_u = (p.peso_kg or (1.0 if (p.um or "").upper().startswith("K") else 0)) if p else 0
        kg = round(x["stock"] * kg_u, 1)
        venc = min([l.fecha_venc for l in ls if l.fecha_venc] or [None])
        dias_venc = (venc - hoy).days if venc else None
        c = cost_per_kg(p) if p else None
        quienes = sorted(compradores.get(x["code"], {}).items(), key=lambda t: -t[1])[:3]
        if dias_venc is not None and dias_venc <= 60:
            reco = f"Vence en {dias_venc} días: liquidar con precio especial o promoción por volumen esta semana."
        elif quienes:
            reco = "Ofrecer a quienes lo compraron en el último año: " + ", ".join(n for n, _ in quienes) + "."
        elif x["ultima_salida"] is None:
            reco = "Nunca tuvo salida: proponerlo como «producto del mes», muestra comercial o combo con un producto líder."
        else:
            reco = "Incluirlo en la próxima campaña y revisar el precio frente al mercado."
        filas.append(dict(x, dias_sin=dias_sin, duracion=_duracion(dias_sin), kg=kg, vence=venc, dias_venc=dias_venc,
                          valor=round(kg * c, 2) if c and kg else None, compradores=quienes, recomendacion=reco,
                          ingreso=ingreso))
    por_tiempo = sorted(filas, key=lambda f: -(f["dias_sin"] if f["dias_sin"] is not None else 10 ** 6))[:top]
    # A01: la «cantidad» se mide en S/ cuando hay costo cargado; si no, en kg
    por_cantidad = sorted(filas, key=lambda f: (f["valor"] is None, -(f["valor"] if f["valor"] is not None else (f["kg"] or f["stock"]))))[:top]
    for i, f in enumerate(por_tiempo, start=1):
        f["rank_tiempo"] = i
    for i, f in enumerate(por_cantidad, start=1):
        f["rank_cantidad"] = i
    return {"dias": base["dias"], "total": len(filas), "por_tiempo": por_tiempo, "por_cantidad": por_cantidad,
            "kg_total": round(sum(f["kg"] for f in filas), 1),
            "valor_total": round(sum(f["valor"] or 0 for f in filas), 2), "almacen": almacen or ""}
