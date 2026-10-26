"""Observaciones del módulo V7 (26-oct-2026)."""
from datetime import date, timedelta

import pytest

from app.services.core import BusinessError

from .conftest import alerts_of
from .test_web import login

LN = [{"product_code": "PT-905", "cantidad": 5, "precio_unitario": 10}]


def test_regularizacion_con_fecha_no_pasada_queda_en_borrador_y_alerta(s, U, cust):
    from app.services import orders as osvc
    u = U("asist.ventas")
    o = osvc.create_order(u, "RG-1", cust().id, date.today(), LN, regularizacion="Pedido telefónico")
    s.flush()
    assert o.estado == "REGULARIZACION_OBSERVADA" and o.id
    a = alerts_of(s, "OV_REGULARIZACION_FECHA", o.id)
    assert a and u.id in a[0].usuarios
    with pytest.raises(BusinessError, match="anterior a hoy"):
        osvc.regularize_fix_date(o, u, date.today())
    osvc.regularize_fix_date(o, u, date.today() - timedelta(days=1))
    assert o.estado == "REGULARIZADA"


def test_regularizacion_mercaderia_salio_descuenta_fefo(s, U, cust):
    from app.models import SalesHistory
    from app.services import orders as osvc
    from app.services.stock import free_qty, lots_of
    u = U("asist.ventas")
    o = osvc.create_order(u, "RG-2", cust().id, date.today() - timedelta(days=2), LN, regularizacion="Atendido el sábado")
    s.flush()
    antes = sum(l.cantidad for l in lots_of("PT-905"))
    prop = osvc.regularization_exits(o)
    assert prop and prop[0]["propuesta"]
    osvc.regularize_close(o, U("coord.ventas"), "Guía 001-1")
    s.flush()
    assert o.estado == "REGULARIZADA_CERRADA" and o.lines[0].despachado == 5
    assert abs(sum(l.cantidad for l in lots_of("PT-905")) - (antes - 5)) < 1e-6
    assert s.query(SalesHistory).filter_by(documento="RG-2", origen="REGULARIZACION").count() >= 1
    assert o.regularizacion["salidas"]


def test_regularizacion_ya_descontada_en_sap_no_descuenta(s, U, cust):
    from app.services import orders as osvc
    from app.services.stock import lots_of
    o = osvc.create_order(U("asist.ventas"), "RG-3", cust().id, date.today() - timedelta(days=2), LN, regularizacion="x")
    s.flush()
    antes = sum(l.cantidad for l in lots_of("PT-905"))
    osvc.regularize_close(o, U("coord.ventas"), "Ya en SAP", descontar=False)
    assert sum(l.cantidad for l in lots_of("PT-905")) == antes and o.regularizacion["descontado"] is False


def test_fuera_del_corte_queda_borrador_y_se_registra_al_corregir_fecha(s, U, cust):
    from app.services import orders as osvc
    from app.services.core import set_param
    set_param("hora_corte_ov_dia_siguiente", "00:00")
    manana = date.today() + timedelta(days=1)
    o = osvc.create_order(U("coord.ventas"), "CT-1", cust().id, manana, LN, oc_archivo="oc.pdf", evaluar=False)
    s.flush()
    assert o.estado == "BORRADOR"
    assert "Corrija la fecha" in o.motivo_detencion
    osvc.modify_order(o, U("coord.ventas"), {"fecha_compromiso": manana + timedelta(days=2)})
    assert o.estado != "BORRADOR"


def test_precio_bajo_base_en_vivo(s, U, cust):
    from app.models import PriceItem, Product
    from app.services.prices import check_line_price
    c = cust()
    it = s.query(PriceItem).filter(PriceItem.lista == "GENERAL", PriceItem.vigente.is_(True)).first()
    if it is None:
        pytest.skip("demo sin lista general")
    c.canal_precio = it.canal
    p = s.get(Product, it.product_code)
    assert check_line_price(p, c, it.precio * 0.5)["estado"] == "DEBAJO"
    assert check_line_price(p, c, it.precio * 2)["estado"] == "SOBRE"


def test_otra_direccion_se_agrega_como_direccion_n(s, U, cust):
    from app.services import orders as osvc
    c = cust()
    n0 = len(c.sites)
    e = osvc.resolve_delivery(c, None, "", "Jr. Ficticio 456", "Ate", "Lima", guardar=True, user=U("coord.ventas"))
    s.flush()
    s.refresh(c)
    assert len(c.sites) == n0 + 1 and e["site_id"] and e["lugar_entrega"].startswith("Dirección")
    e2 = osvc.resolve_delivery(c, None, "", "jr.  ficticio 456", "ATE", "Lima", guardar=True, user=U("coord.ventas"))
    s.flush()
    s.refresh(c)
    assert len(c.sites) == n0 + 1 and e2["site_id"] == e["site_id"]


def test_pantallas_v7(app):
    c = app.test_client()
    login(c, "coord.ventas")
    assert c.get("/ov/regularizacion").status_code == 200
    r = c.get("/ov/clientes/1")
    assert r.status_code == 200 and "Penalidades del cliente" in r.get_data(as_text=True)
    r = c.get("/ov/api/precio?producto=PT-905&cliente=1&precio=1")
    assert r.status_code == 200 and "estado" in r.get_json()


def test_penalidades_ilimitadas_deteccion_y_resultado(s, U, cust):
    from app.services import orders as osvc
    from app.services import penalties as pen
    c = cust()
    filas = [{"tipo": "FECHA_ENTREGA", "modo": "FIJO", "valor": "150"}, {"tipo": "ENTREGA_PARCIAL", "modo": "PCT_OV", "valor": "10"},
             {"tipo": "OTRO", "nombre": "Etiquetado incorrecto", "modo": "FIJO", "valor": "80"}]
    cond = pen.save_conditions(c, U("coord.ventas"), filas)
    assert "ETIQUETADO_INCORRECTO" in cond.penalidad_tipos and len(cond.penalidad_tipos) == 3
    o = osvc.create_order(U("coord.ventas"), "PN-1", c.id, date.today() + timedelta(days=3), LN, oc_archivo="oc.pdf", evaluar=False)
    s.flush()
    o.lines[0].entregado = 3
    nuevas = pen.detect_for_order(o, U("coord.ventas"), fecha_entrega=o.fecha_compromiso + timedelta(days=1))
    assert {p.tipo for p in nuevas} == {"FECHA_ENTREGA", "ENTREGA_PARCIAL"}
    assert pen.detect_for_order(o, U("coord.ventas"), fecha_entrega=o.fecha_compromiso + timedelta(days=1)) == []
    assert alerts_of(s, "CLIENTE_PENALIDAD", nuevas[0].id)
    with pytest.raises(BusinessError, match="desestim"):
        pen.resolve(nuevas[0], U("coord.ventas"), "DESESTIMADA")
    pen.resolve(nuevas[0], U("coord.ventas"), "PAGADA", 100)
    pen.resolve(nuevas[1], U("coord.ventas"), "DESESTIMADA", comentario="El cliente aceptó la reposición")
    ind = pen.indicator(customer_id=c.id)["total"]
    assert ind["pagada"] == 100 and ind["calculada"] == 150 + round(50 * 0.10, 2) and ind["pendientes"] == 0


def test_codigo_de_cliente_correlativo(s, U):
    from app.models import Customer
    from app.services import customers as csv_
    from app.services.penalties import next_customer_code
    s.query(Customer).first().codigo_sap = "C000120"
    s.flush()
    assert next_customer_code() == "C000121"
    c = csv_.save_customer(U("coord.ventas"), None, {"ruc": "20999999991", "nombre": "CLIENTE FICTICIO V7",
                                                      "direccion": "Av. Ficticia 1", "ubigeo": "150101"})
    assert c.codigo_sap == "C000121"


def test_carga_sap_cierra_ov_atendida_en_sap_sin_duplicar(s, U, tmp_path):
    import os
    import openpyxl
    from app.demo_data import write_ovs
    from app.models import SalesOrder
    from app.services import uploads as up
    from app.services.stock import line_reserved
    p = os.path.join(tmp_path, "ovs.xlsx")
    write_ovs(p, date.today())
    up.confirm(up.create_batch(U("coord.ventas"), "OV", p, "ovs.xlsx"), U("coord.ventas"))
    n = s.query(SalesOrder).count()
    res = up.confirm(up.create_batch(U("coord.ventas"), "OV", p, "ovs.xlsx"), U("coord.ventas"))  # 2.ª carga del día
    assert s.query(SalesOrder).count() == n and res["ov_creadas"] == 0
    o = s.query(SalesOrder).filter_by(numero="SAP-9001").one()
    assert any(line_reserved(l) > 0 for l in o.lines)
    wb = openpyxl.load_workbook(p)
    ws = wb["OVS"]
    hdr = [c.value for c in ws[1]]
    ie, inum = hdr.index("Estado"), hdr.index("Número de documento")
    for row in ws.iter_rows(min_row=2):
        if row[inum].value == "SAP-9001":
            row[ie].value = "Cerrado"
    wb.save(p)
    res = up.confirm(up.create_batch(U("coord.ventas"), "OV", p, "ovs.xlsx"), U("coord.ventas"))
    s.flush()
    assert res["ov_cerradas_en_sap"] == 1 and o.estado == "CERRADA"
    assert all(line_reserved(l) == 0 for l in o.lines)


def test_informe_gestion_comercial_y_exportaciones(app):
    c = app.test_client()
    login(c, "jefe.ventas")
    r = c.get("/informes/gestion")
    assert r.status_code == 200
    h = r.get_data(as_text=True)
    assert "OTIF" in h and "Fill rate" in h and "Embudo comercial" in h and "Desempeño por persona" in h
    assert c.get("/informes/gestion?tipo=GRANEL&modo_canal=excluir&canal=X").status_code == 200
    r = c.get("/informes/gestion?formato=pptx")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    assert c.get("/informes/gestion?formato=pdf").data[:4] == b"%PDF"
    assert c.get("/informes/gestion?formato=xlsx").data[:2] == b"PK"


def test_formulas_nivel_de_servicio(s, U, cust):
    from datetime import datetime
    from app.services import orders as osvc
    from app.services.commercial_mgmt import _nivel
    o1 = osvc.create_order(U("coord.ventas"), "NS-1", cust().id, date.today() + timedelta(days=3), LN, oc_archivo="o", evaluar=False)
    o2 = osvc.create_order(U("coord.ventas"), "NS-2", cust().id, date.today() + timedelta(days=3), LN, oc_archivo="o", evaluar=False)
    s.flush()
    o1.salida_ts = datetime.combine(o1.fecha_compromiso, datetime.min.time())
    o1.lines[0].despachado = 5
    o2.salida_ts = datetime.combine(o2.fecha_compromiso + timedelta(days=1), datetime.min.time())
    o2.lines[0].despachado = 4
    n = _nivel([o1, o2], "")
    assert (n["on_time"], n["in_full"], n["otif"], n["fill_unidades"]) == (50.0, 50.0, 50.0, 90.0)


@pytest.mark.parametrize("usuario", ["coord.ventas", "jefe.ventas", "gerente.general"])
def test_portada_comercial_con_metas(app, usuario):
    c = app.test_client()
    login(c, usuario)
    r = c.get("/")
    assert r.status_code == 200 and "Gestión comercial del mes" in r.get_data(as_text=True)


def test_alerta_meta_atrasada(s, U):
    from app.models import Alert, SalesTarget
    from app.services.commercial_mgmt import check_goal_alerts
    hoy = date.today()
    s.add(SalesTarget(mes=f"{hoy:%Y-%m}", alcance="GENERAL", clave="", soles=10_000_000))
    s.flush()
    if hoy.day < 5:
        pytest.skip("al inicio del mes el avance esperado aún es bajo")
    assert check_goal_alerts(hoy) == 1 and check_goal_alerts(hoy) == 1
    assert s.query(Alert).filter_by(type_code="META_ATRASADA").count() == 1


def test_catalogo_carpetas_destacado_y_alta(s, U, app, tmp_path):
    from app.models import Product
    from app.services import catalog as ct
    u = U("jefe.ventas")
    p = s.get(Product, "PT-905")
    ct.add_document(u, p, "FICHA_TECNICA", "ficha.pdf", "Ficha técnica 2026")
    ct.add_document(u, p, "RECETAS", "receta.pdf")
    docs = ct.documents(p)
    assert len(docs["FICHA_TECNICA"]) == 1 and len(docs["RECETAS"]) == 1
    with pytest.raises(BusinessError):
        ct.add_document(U("asist.ventas"), p, "FOTOS", "x.jpg")
    ct.set_featured(u, p, "MES", "Campaña de invierno", date.today() + timedelta(days=10))
    assert p in ct.featured()
    nuevo = ct.add_product(u, {"code": "PT-9999", "descripcion": "producto ficticio v7", "tipo": "PT", "peso_kg": "1"})
    assert nuevo.descripcion == "PRODUCTO FICTICIO V7"
    c = app.test_client()
    login(c, "jefe.ventas")
    assert c.get("/ov/catalogo/PT-905").status_code == 200
    assert "Destacados" in c.get("/ov/catalogo").get_data(as_text=True)


def test_oportunidad_varios_productos_y_cotizaciones(s, U, cust, app):
    from werkzeug.datastructures import MultiDict
    from app.services import opportunities as op
    o = op.save_opportunity(U("coord.ventas"), cust(), MultiDict([("titulo", "Ingreso a cadena"),
                                                                  ("productos", "PT-905 · algo"), ("productos", "PT-902"),
                                                                  ("productos", "Quinua orgánica")]))
    assert o.productos == "PT-905, PT-902, Quinua orgánica"
    q = op.save_quote(U("coord.ventas"), cust(), [{"product_code": "PT-905", "cantidad": 2, "precio": 10}], opp=o)
    s.commit()
    c = app.test_client()
    login(c, "coord.ventas")
    assert c.get("/ov/cotizaciones").status_code == 200
    r = c.get(f"/ov/cotizacion/{q.id}?formato=pdf")
    assert r.data[:4] == b"%PDF"


def test_cotizaciones_sap_lectura_y_escritura_bloqueada(s, U, cust):
    from app.services import sap_sl
    c = cust()
    c.codigo_sap = "C0001"
    s.flush()

    class FakeSL:
        def get_all(self, path, params=None):
            if path == "Quotations":
                return [{"DocNum": 501, "DocDate": "2026-10-20", "DocDueDate": "2026-11-20", "CardCode": "C0001",
                         "CardName": c.nombre, "DocTotal": 100, "DocumentStatus": "bost_Open",
                         "DocumentLines": [{"ItemCode": "PT-905", "Quantity": 10, "Price": 10}]}]
            return []
    r = sap_sl.pull_quotes(U("coord.ventas"), FakeSL())
    assert r["nuevas"] == 1
    assert sap_sl.pull_quotes(U("coord.ventas"), FakeSL())["actualizadas"] == 1
    from app.models import Quote
    q = s.query(Quote).filter_by(origen="SAP").one()
    with pytest.raises(BusinessError, match="no está autorizada"):
        sap_sl.push_quote(q, U("coord.ventas"))


def test_metas_ampliadas_y_proyeccion_por_grupo(s, U, cust):
    from app.services import margin, supply
    from app.services.commercial_mgmt import new_customers_and_skus
    mes = f"{date.today():%Y-%m}"
    t = margin.save_target(U("coord.ventas"), mes, "GENERAL", "", 1000, 500, 2, 5)
    assert t.nuevos_clientes == 2 and t.nuevos_sku == 5
    with pytest.raises(BusinessError):
        margin.save_target(U("asist.ventas"), mes, "GENERAL", "", 1)
    assert "clientes" in new_customers_and_skus(date.today().replace(day=1), date.today(), [])
    filas = [{"product_code": "PT-905", "cantidad": 100, "canal": "", "customer_id": "", "observacion": ""},
             {"product_code": "PT-902 · x", "cantidad": 50, "canal": "", "customer_id": str(cust().id), "observacion": "Cliente nuevo"}]
    assert supply.save_forecast_group(U("jefe.ventas"), mes, filas) == {"nuevas": 2, "cambios": 0}
    with pytest.raises(BusinessError, match="no existe"):
        supply.save_forecast_group(U("jefe.ventas"), mes, [{"product_code": "NO-EXISTE", "cantidad": 1}])
    filas[0].update(cantidad=150, observacion="Campaña")
    assert supply.save_forecast_group(U("jefe.ventas"), mes, filas[:1]) == {"nuevas": 0, "cambios": 1}
    f = supply.pending_changes()[0]
    with pytest.raises(BusinessError):
        supply.decide_change(f, U("jefe.ventas"), True)
    with pytest.raises(BusinessError, match="motivo"):
        supply.decide_change(f, U("gerente.general"), False)
    supply.decide_change(f, U("gerente.general"), True)
    assert f.cantidad == 150 and not supply.pending_changes()


def test_rotacion_top20_y_turnos_por_area_con_tercero(s, U, app):
    from app.services import shifts, supply
    d = supply.rotation_ranking(30)
    assert len(d["por_tiempo"]) <= 20 and len(d["por_cantidad"]) <= 20
    for f in d["por_tiempo"]:
        assert f["recomendacion"] and ("semanas" in f["duracion"] or f["dias_sin"] is None)
    jp = U("jefe.produccion")
    otro = s.query(type(jp)).filter_by(username="coord.ventas").one()
    with pytest.raises(BusinessError, match="su área"):
        shifts.add_reinforcement(jp, otro.id, date.today(), None, "08:00", "12:00", "Emergencia")
    with pytest.raises(BusinessError, match="costo por hora"):
        shifts.add_reinforcement(jp, None, date.today(), None, "08:00", "12:00", "Emergencia", {"nombre": "Persona Ficticia"})
    us = shifts.add_reinforcement(jp, None, date.today(), date.today() + timedelta(days=1), "08:00", "12:00", "Pedido urgente",
                                  {"nombre": "Persona Ficticia", "empresa": "Servicios Demo SAC"}, 15)
    c = shifts.reinforcement_costs(date.today(), date.today() + timedelta(days=5), shifts.managed_areas(jp))
    assert any(r["us"].id == us.id and r["horas"] == 8 and r["costo"] == 120 for r in c["filas"])
    s.commit()
    cl = app.test_client()
    login(cl, "jefe.produccion")
    h = cl.get("/configuracion/turnos").get_data(as_text=True)
    assert "Costo de horas extra" in h and "Persona Ficticia" in h
    login(cl, "coord.almacen")
    assert cl.get("/suministro/rotacion?frente=cantidad").status_code == 200


def test_alertas_por_prioridad(s, U):
    from app.services import alerts
    u = U("coord.ventas")
    rows = alerts.by_priority(u, alerts.inbox(u))
    claves = [(a.estado in alerts.CERRADOS, alerts.nature_for(u, a) != "TAREA", alerts.PRIORIDAD_ORDEN.get(a.prioridad, 9)) for a in rows]
    assert claves == sorted(claves)
