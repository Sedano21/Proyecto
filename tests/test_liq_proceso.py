"""Liquidación de proceso CP-OPE-004 (30-oct-2026): solo Producción."""
import io
from datetime import date

import pytest
from werkzeug.datastructures import MultiDict

from app.services.core import BusinessError

from .test_web import login


def _parte(s, linea="TRUER", kg=1000.0, granel=900.0, desmedro=60.0, sobrante=40.0, lote="L-DEMO-1"):
    from app.models import ProductionReport
    r = ProductionReport(linea=linea, tipo="SELECCION", fecha=date.today(), product_code="PT-905", peso_entrada_kg=kg,
                         granel_kg=granel, desmedro_kg=desmedro, sobrante_kg=sobrante, lote=lote, horas=8, operarios=4,
                         estado="PENDIENTE")
    s.add(r)
    s.flush()
    return r


def _form(liq, **extra):
    """Simula el formulario de edición con los datos actuales del formato."""
    d = MultiDict()
    for x in liq.servicio:
        d.add("s_etq", x["etiqueta"]); d.add("s_val", str(x["valor"]))
    for x in liq.producto:
        d.add("p_etq", x["etiqueta"]); d.add("p_val", str(x["valor"]))
    for i, p in enumerate(liq.procesos):
        d.add(f"t{i}", p["titulo"])
        d.add(f"h{i}", "" if p.get("horas") is None else str(p["horas"]))
        d.add(f"o{i}", "" if p.get("operarios") is None else str(p["operarios"]))
        for f in p["filas"]:
            d.add(f"f{i}_proceso", f["proceso"]); d.add(f"f{i}_calidad", f["calidad"])
            for k, n in (("sacos", "sacos"), ("peso_saco", "peso"), ("saldo", "saldo")):
                d.add(f"f{i}_{n}", "" if f[k] is None else str(f[k]))
    for k, v in extra.items():
        d.add(k, v)
    return d


def test_solo_produccion(s, U):
    from app.services import process_liq as pl
    r = _parte(s)
    with pytest.raises(BusinessError):
        pl.create(U("asist.calidad"), [r.id])
    with pytest.raises(BusinessError):
        pl.create(U("coord.almacen"), [r.id])
    assert pl.create(U("asist.produccion"), [r.id]).numero == 1


def test_prellenado_desde_partes_y_calculos(s, U):
    from app.services import process_liq as pl
    a = _parte(s, "TRUER")
    b = _parte(s, "CIMBRIA", kg=500, granel=450, desmedro=50, sobrante=0)
    liq = pl.create(U("asist.produccion"), [a.id, b.id])
    assert liq.numero == 1 and liq.lote == "L-DEMO-1" and liq.linea == "TRUER / CIMBRIA" and len(liq.procesos) == 2
    assert liq.procesos[0]["titulo"] == "1ER PROCESO : TRUER" and liq.procesos[1]["titulo"] == "2DO PROCESO : CIMBRIA"
    etq = [x["etiqueta"] for x in liq.servicio]
    assert etq[:2] == ["Proceso", "Proceso N°1"] and "Fecha del reporte" in etq
    assert [x["etiqueta"] for x in liq.producto][:2] == ["Descripción", "Marca"]
    c = pl.calc(liq)["procesos"][0]
    assert c["ingreso"] == 1000 and c["filas"][1]["pct"] == 90.0
    assert c["desmedro_kg"] == 60 and c["desmedro_pct"] == 6.0
    assert c["desviacion"] == 0  # 1000 = 900 + 60 + 40
    assert c["horas"] == 8 and c["operarios"] == 4
    with pytest.raises(BusinessError):  # partes de distinto SKU
        from app.models import ProductionReport
        x = ProductionReport(linea="R1", tipo="ENVASADO", fecha=date.today(), product_code="PT-903", peso_entrada_kg=1, estado="PENDIENTE")
        s.add(x); s.flush()
        pl.create(U("asist.produccion"), [a.id, x.id])
    assert pl.create(U("asist.produccion"), [a.id]).numero == 2


def test_campos_filas_y_procesos_se_agregan_o_quitan(s, U):
    from app.services import process_liq as pl
    liq = pl.create(U("jefe.produccion"), [_parte(s).id])
    f = _form(liq, agregar_proceso="1", observaciones="Proceso sin novedades")
    # quitar el campo «Proceso N°3», agregar uno nuevo y una fila por sacos
    i = f.getlist("s_etq").index("Proceso N°3")
    etq, val = f.getlist("s_etq"), f.getlist("s_val")
    etq[i], val[i] = "", ""
    f.setlist("s_etq", etq + ["Contenedor"]); f.setlist("s_val", val + ["CTNR-DEMO-01"])
    f.setlist("f0_proceso", f.getlist("f0_proceso") + ["SALIDA - DESMEDRO"])
    f.setlist("f0_calidad", f.getlist("f0_calidad") + ["POLVO"])
    f.setlist("f0_sacos", f.getlist("f0_sacos") + ["2"]); f.setlist("f0_peso", f.getlist("f0_peso") + ["50"])
    f.setlist("f0_saldo", f.getlist("f0_saldo") + [""])
    pl.update(U("asist.produccion"), liq, f)
    etqs = [x["etiqueta"] for x in liq.servicio]
    assert "Proceso N°3" not in etqs and {"etiqueta": "Contenedor", "valor": "CTNR-DEMO-01"} in liq.servicio
    assert len(liq.procesos) == 2 and liq.procesos[1]["filas"][0]["proceso"] == "INGRESO"
    c = pl.calc(liq)["procesos"][0]
    assert c["desmedro_kg"] == 160 and c["filas"][-1]["total"] == 100
    assert liq.observaciones == "Proceso sin novedades"
    pl.update(U("asist.produccion"), liq, _form(liq, quitar_1="1"))
    assert len(liq.procesos) == 1
    f = _form(liq)
    f.setlist("s_etq", f.getlist("s_etq") + [""]); f.setlist("s_val", f.getlist("s_val") + ["sin nombre"])
    with pytest.raises(BusinessError):
        pl.update(U("asist.produccion"), liq, f)


def test_variante_pf_con_costos(s, U):
    from app.services import process_liq as pl
    liq = pl.create(U("asist.produccion"), [_parte(s, kg=1000, granel=880, desmedro=120, sobrante=0).id], "PF")
    assert [f["calidad"] for f in liq.procesos[0]["filas"]] == ["GRANEL", "PROCESADO", "DESMEDRO"]
    pl.update(U("asist.produccion"), liq, _form(liq, mo_seleccion="10", mo_reproceso="2", costo_hora="5",
                                                costo_unidad="0.1", costo_kg_desmedro="2"))
    k = pl.calc(liq)["costos"]
    assert k["horas"] == 12 and k["mo"] == 60 and k["cif"] == 100 and k["total"] == 160 and k["unitario"] == 0.16
    assert k["desmedro_kg"] == 120 and k["desmedro_costo"] == 240


def test_cerrar_reabrir_y_fotos(s, U):
    from app.services import process_liq as pl
    liq = pl.create(U("asist.produccion"), [_parte(s).id])
    pl.add_photos(U("asist.produccion"), liq, ["a.jpg", "", "b.jpg"])
    assert liq.fotos == ["a.jpg", "b.jpg"]
    pl.update(U("asist.produccion"), liq, _form(liq, quitar_foto="a.jpg"))
    assert liq.fotos == ["b.jpg"]
    pl.close(U("asist.produccion"), liq)
    with pytest.raises(BusinessError):
        pl.update(U("asist.produccion"), liq, _form(liq))
    with pytest.raises(BusinessError):
        pl.reopen(U("asist.produccion"), liq)
    pl.reopen(U("jefe.produccion"), liq)
    assert liq.estado == "BORRADOR"


def test_excel_con_el_formato(s, U, app):
    from openpyxl import load_workbook
    from app.services import process_liq as pl
    liq = pl.create(U("asist.produccion"), [_parte(s).id])
    data = pl.export_excel(liq, app.static_folder, app.config["UPLOAD_DIR"])
    ws = load_workbook(io.BytesIO(data)).active
    vals = {c.value for row in ws.iter_rows() for c in row if c.value is not None}
    for t in ("LIQUIDACION DE PROCESO ", "CP- OPE - 004", "OPERACIONES", " N° 00001", "DATOS DEL SERVICIO", "DATOS DEL PRODUCTO",
              "PROCESO", "CALIDAD", "SACOS (UNI)", "PESO POR SACO (KG)", "SALDO (KG)", "TOTAL (KG)", "MOVIMIENTO (KG)",
              "DESVIACIÓN KG", "TOTAL DE HR", "NRO DE OPE.", "DESMEDRO", "REFERENCIAS GRÁFICAS DEL LOTE RECEPCIONADO",
              "OBSERVACIONES"):
        assert t in vals, t
    assert any(isinstance(v, str) and v.startswith("=D") and "*E" in v for v in vals)  # TOTAL con fórmula del formato
    assert len(ws._images) >= 1  # logo


def test_pantallas(app, s, U):
    from app.services import process_liq as pl
    r = _parte(s)
    liq = pl.create(U("asist.produccion"), [r.id])
    s.commit()
    c = app.test_client()
    login(c, "asist.produccion")
    for p in ("/produccion/liquidacion-proceso", f"/produccion/liquidacion-proceso?sku=PT-905&lote=L-DEMO-1",
              f"/produccion/liquidacion-proceso/{liq.id}", f"/produccion/liquidacion-proceso/{liq.id}?imprimir=1"):
        html = c.get(p).get_data(as_text=True)
        assert "LIQUIDACION DE PROCESO" in html or "Liquidación de proceso" in html, p
    html = c.get(f"/produccion/liquidacion-proceso/{liq.id}").get_data(as_text=True)
    assert "1ER PROCESO : TRUER" in html and "N° 00001" in html and "REFERENCIAS GRÁFICAS" in html
    assert c.get(f"/produccion/liquidacion-proceso/{liq.id}/excel").data[:2] == b"PK"
    assert "Liquidación de proceso (CP-OPE-004)" in c.get("/").get_data(as_text=True)
    c2 = app.test_client()
    login(c2, "asist.calidad")
    assert c2.get(f"/produccion/liquidacion-proceso/{liq.id}").status_code in (302, 403)
