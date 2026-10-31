"""Indicaciones de Calidad del 30-oct-2026 (V13)."""
import io
import re
from datetime import date, timedelta

import pytest

from app.services.core import BusinessError

from .conftest import alerts_of
from .test_web import login


def _get(app, user, path):
    c = app.test_client()
    login(c, user)
    r = c.get(path)
    assert r.status_code == 200, (user, path, r.status_code)
    return r.get_data(as_text=True)


def _lot(s, estado="DISPONIBLE"):
    from app.models import StockLot
    lot = s.query(StockLot).filter(StockLot.estado == "DISPONIBLE", StockLot.cantidad > 0).first()
    return lot


def test_lote_observado_alerta_inmediata_y_semanal(s, U):
    from app.services import quality_v13 as q
    lot = _lot(s)
    with pytest.raises(BusinessError):
        q.observe_lot(U("coord.almacen"), lot, "PLAGA")
    lot = q.observe_lot(U("asist.calidad"), lot, "PIEDRA", "presencia de piedra en la muestra")
    s.flush()
    assert lot.estado == "OBSERVADO" and lot.observacion["categoria"] == "PIEDRA"
    a = alerts_of(s, "LOTE_OBSERVADO", lot.id)
    assert a and a[0].prioridad == "ALTA"
    assert q.check_observed(date.today() + timedelta(days=3)) == 0
    assert q.check_observed(date.today() + timedelta(days=7)) == 1
    s.flush()
    assert len(alerts_of(s, "LOTE_OBSERVADO", lot.id)) == 2
    q.check_observed(date.today() + timedelta(days=8))  # misma semana: no duplica
    s.flush()
    assert len(alerts_of(s, "LOTE_OBSERVADO", lot.id)) == 2


def test_bloqueado_solo_lo_libera_calidad_con_solicitud(s, U):
    from app.services import quality_v13 as q
    from app.services.stock import change_lot_state
    lot = q.observe_lot(U("asist.calidad"), _lot(s), "HONGUEADO")
    s.flush()
    with pytest.raises(BusinessError, match="solo lo libera Calidad"):
        change_lot_state(lot, "DISPONIBLE", U("coord.almacen"), "ya está bien")
    with pytest.raises(BusinessError):
        q.request_review(U("coord.almacen"), lot, "LIBERACION", "")
    r = q.request_review(U("coord.almacen"), lot, "LIBERACION", "se separó lo hongueado")
    s.flush()
    assert lot.liberacion_solicitada and alerts_of(s, "REVISION_LOTE", r.id)
    with pytest.raises(BusinessError):
        q.resolve_review(U("coord.almacen"), r, "LIBERADO")
    q.resolve_review(U("jefe.calidad"), r, "LIBERADO", "conforme")
    s.flush()
    lot2 = q._segment(lot, "DISPONIBLE")
    assert lot2.estado == "DISPONIBLE" and r.estado == "RESUELTA" and not lot2.liberacion_solicitada


def test_reservas_solo_disponibles(s, U):
    from app.services import quality_v13 as q
    from app.services.production import available_in
    lot = _lot(s)
    antes = available_in(lot.product_code)
    q.observe_lot(U("asist.calidad"), lot, "VENCIMIENTO")
    s.flush()
    assert available_in(lot.product_code) < antes


def test_fumigacion_con_inicio_y_liberacion(s, U):
    from app.services import quality_v13 as q
    lot = _lot(s)
    with pytest.raises(BusinessError):
        q.set_fumigation(U("asist.calidad"), lot, date.today(), None)
    with pytest.raises(BusinessError):
        q.set_fumigation(U("asist.calidad"), lot, date.today(), date.today() - timedelta(days=1))
    lot = q.set_fumigation(U("asist.calidad"), lot, date.today(), date.today() + timedelta(days=4), "Almacén externo 767")
    s.flush()
    assert lot.estado == "EN_FUMIGACION" and lot.fumigacion["zona"] == "Almacén externo 767"
    assert q.check_fumigation_release(date.today() + timedelta(days=3)) == 0
    assert q.check_fumigation_release(date.today() + timedelta(days=4)) == 1


def test_revision_apto_proceso_y_a_fumigacion(s, U):
    from app.services import quality_v13 as q
    lot = _lot(s)
    r = q.request_review(U("jefe.produccion"), lot, "APTO_PROCESO", "antes de la OF")
    s.flush()
    with pytest.raises(BusinessError):
        q.resolve_review(U("jefe.calidad"), r, "NO_APTO")  # sin motivo
    q.resolve_review(U("jefe.calidad"), r, "A_FUMIGACION", "gorgojo", date.today(), date.today() + timedelta(days=3))
    s.flush()
    assert q._segment(lot, "EN_FUMIGACION").estado == "EN_FUMIGACION"


def test_fechas_de_produccion_multiples_editables(s, U):
    from app.services import quality_v13 as q
    lot = _lot(s)
    d1, d2 = date.today() - timedelta(days=10), date.today() - timedelta(days=5)
    q.add_production_dates(U("coord.almacen"), lot, [{"fecha": d2.isoformat(), "cantidad": "10"},
                                                     {"fecha": d1.isoformat(), "estado": "OBSERVADO"}])
    assert [x["fecha"] for x in lot.fechas_produccion] == [d1.isoformat(), d2.isoformat()]
    with pytest.raises(BusinessError):
        q.add_production_dates(U("coord.almacen"), lot, [{"fecha": d1.isoformat()}])
    q.update_production_date(U("asist.calidad"), lot, d1.isoformat(), "LIBERADO", "revisado")
    assert lot.fechas_produccion[0]["estado"] == "LIBERADO"


def test_bandeja_de_calidad(app, s, U):
    from app.models import Incident, RepackRequest
    from app.services import quality_v13 as q
    s.add(Incident(tipo="INFESTACION_INTERNA", grupo="CALIDAD", descripcion="gorgojo en la entrega"))
    lot = _lot(s)
    s.add(RepackRequest(tipo="CAMBIO_PRESENTACION", lot_id=lot.id, cantidad=5, producto_destino=lot.product_code))
    q.request_review(U("coord.almacen"), lot, "FUMIGACION", "se vio polilla")
    s.commit()
    d = q.inbox()
    assert d["devoluciones"] and d["devoluciones"][0]["calidad"] and d["revisiones"]
    html = _get(app, "jefe.calidad", "/calidad/bandeja")
    for t in ("Bandeja de Calidad", "por motivo de calidad", "¿Va a fumigación?", "Solicitudes de reempaque", "Marca"):
        assert t in html, t


def test_observacion_en_linea_por_pallet_con_fases(s, U):
    from app.services import quality_v13 as q
    o = q.observe_line(U("asist.calidad"), "TRUER", "granos partidos sobre lo tolerado",
                       [{"pallet": "P-01", "detalle": "parte superior"}, {"pallet": "P-02"}, {"pallet": ""}])
    s.flush()
    assert len(o.pallets) == 2 and alerts_of(s, "OBS_LINEA", o.id)
    q.advance_observation(U("jefe.produccion"), o, "EN_ATENCION")
    with pytest.raises(BusinessError):
        q.advance_observation(U("jefe.produccion"), o, "CERRADA", "listo")  # cierra Calidad
    with pytest.raises(BusinessError):
        q.advance_observation(U("asist.calidad"), o, "CERRADA", "")
    q.advance_observation(U("asist.calidad"), o, "CERRADA", "se reajustó la máquina")
    assert o.fase == "CERRADA" and [h["fase"] for h in o.historial] == ["NOTIFICADA", "EN_ATENCION", "CERRADA"]


def test_merma_validacion_reproceso_y_vencimiento(s, U):
    from app.models import ProductionReport, now
    from app.services import quality_v13 as q
    r = ProductionReport(linea="TRUER", tipo="SELECCION", fecha=date.today(), product_code="PT-905", peso_entrada_kg=1000,
                         granel_kg=900, desmedro_pct=9.0, estado="PENDIENTE")
    s.add(r)
    s.flush()
    mv = q.open_merma_validation(r)
    s.flush()
    assert mv.limite_pct == 5 and q.open_merma_validation(r) is None
    assert alerts_of(s, "MERMA_VALIDAR", mv.id)[0].prioridad == "ALTA"
    with pytest.raises(BusinessError):
        q.validate_merma(U("asist.calidad"), mv, False, 30, None)
    q.validate_merma(U("asist.calidad"), mv, False, 30, 60, ["foto1.jpg", "foto2.jpg"], "recuperar grano bueno")
    s.flush()
    assert mv.estado == "REPROCESO" and len(mv.fotos) == 2 and alerts_of(s, "REPROCESO_SOLICITADO", mv.id)
    assert q.check_reprocess(now() + timedelta(hours=47)) == 0
    assert q.check_reprocess(now() + timedelta(hours=49)) == 1
    with pytest.raises(BusinessError):
        q.attend_reprocess(U("asist.calidad"), mv)
    q.attend_reprocess(U("jefe.produccion"), mv, "reprocesado en Cimbria")
    assert mv.estado == "REPROCESADO"


def test_material_para_produccion_y_saldo_por_linea(app, s, U):
    from app.models import ProductionReport
    from app.services import quality_v13 as q
    s.add(ProductionReport(linea="TRUER", tipo="SELECCION", fecha=date.today(), product_code="PT-905", peso_entrada_kg=1000,
                           granel_kg=900, desmedro_pct=5.0, sobrante_kg=40, estado="PENDIENTE"))
    s.commit()
    b = q.balance_by_line()
    f = next(x for x in b["filas"] if x["linea"] == "TRUER")
    assert f["saldo"] == 10.0
    assert isinstance(q.materials_for_production(), list)
    html = _get(app, "asist.calidad", "/calidad/material")
    assert "Saldo por proceso" in html and "TRUER" in html


def test_inspeccion_externa_fases_alertas_y_hallazgos(s, U):
    from app.services import quality_mgmt as qm
    with pytest.raises(BusinessError):
        qm.schedule_inspection(U("asist.calidad"), {"tipo": "EXTERNA", "fecha": date.today().isoformat()})
    i = qm.schedule_inspection(U("asist.calidad"), {"tipo": "EXTERNA", "fecha": date.today().isoformat(), "hora_inicio": "09:30",
                                                    "cliente": "Cliente Demo SAC", "empresa": "Certificadora Demo"})
    s.flush()
    a = alerts_of(s, "INSPECCION", i.id)
    assert a
    with pytest.raises(BusinessError):
        qm.close_inspection(U("asist.calidad"), i, "APROBADO")
    qm.start_inspection(U("asist.calidad"), i)
    qm.add_finding(U("asist.calidad"), i, "Malla rota en ventana del almacén", "Almacén", date.today() + timedelta(days=5))
    s.flush()
    assert alerts_of(s, "HALLAZGO_TAREA", i.id)
    qm.close_inspection(U("asist.calidad"), i, "APROBADO", "92.5", None, ["r.jpg"], ["h1.jpg", "h2.jpg"])
    assert i.fase == "CERRADA" and i.n_hallazgos == 1 and i.puntaje == 92.5 and len(i.fotos_hallazgos) == 2
    assert i.resultado == "BUENO"  # V14: 86–95 % bueno
    with pytest.raises(BusinessError):
        qm.close_finding(U("asist.calidad"), i, 1, "")
    qm.close_finding(U("asist.calidad"), i, 1, "Se cambió la malla")
    assert i.hallazgos[0]["estado"] == "LEVANTADO"


def test_alerta_inspeccion_llega_a_las_areas(s):
    from app.catalogs import ALERT_TYPES
    roles = next(t[5] for t in ALERT_TYPES if t[0] == "INSPECCION")
    for r in ("JEFE_PRODUCCION", "COORD_ALMACEN", "JEFE_LOGISTICA", "GERENCIA_OPERACIONES", "GERENTE_COMERCIAL", "COORD_FINANZAS"):
        assert r in roles, r


def test_visita_almacen_externo_confirmar_y_reprogramar(s, U):
    from app.services import quality_mgmt as qm
    from datetime import datetime
    w = qm.external_warehouses()[0]
    v = qm.propose_visit(U("asist.calidad"), w.code, datetime.combine(date.today() + timedelta(days=2), datetime.min.time()).replace(hour=10))
    s.flush()
    assert alerts_of(s, "VISITA_ALMACEN_EXTERNO", v.id)
    with pytest.raises(BusinessError):
        qm.answer_visit(U("asist.calidad"), v, True)
    with pytest.raises(BusinessError):
        qm.answer_visit(U("coord.almacen"), v, False, None, "")
    qm.answer_visit(U("coord.almacen"), v, False, (date.today() + timedelta(days=3)).isoformat() + "T15:00", "descarga de contenedor")
    assert v.estado == "REPROGRAMADA" and v.propuesta.hour == 15
    qm.answer_visit(U("coord.almacen"), v, True)
    qm.complete_visit(U("asist.calidad"), v, "sin novedades")
    qm.add_finding(U("asist.calidad"), v, "Pallets en mal estado", "Logística")
    assert v.estado == "REALIZADA" and v.hallazgos[0]["area"] == "Logística"
    f = qm.propose_visit(U("asist.calidad"), w.code, datetime.now() + timedelta(days=9), "FUMIGACION")
    assert f.tipo == "FUMIGACION"


def test_formatos_de_verificacion_por_puntos(s, U):
    from app.catalogs import CHECKLIST_CALIDAD
    from app.services import quality_mgmt as qm
    cod, nombre, grupos = qm.checklist_points("LIMPIEZA_EXTERNAS")
    assert cod == "PHS-001" and grupos[0][1]
    n = sum(len(p) for _, p in grupos)
    vals = {str(i): "C" for i in range(n)}
    vals["0"] = "NC"
    with pytest.raises(BusinessError):
        qm.register_checklist(U("asist.calidad"), "LIMPIEZA_EXTERNAS", None, vals, {})
    r = qm.register_checklist(U("asist.calidad"), "LIMPIEZA_EXTERNAS", None, vals, {"0": "se limpió de nuevo"})
    assert r.resultado == "NO_CONFORME"
    assert "HIGIENE_PERSONAL" not in CHECKLIST_CALIDAD and "BPM" not in CHECKLIST_CALIDAD  # V14: PHS-003 propio


def test_consolidado_exportar_e_importar_sin_duplicar(s, U):
    from openpyxl import load_workbook
    from app.models import QualityRecord
    from app.services import quality as qsvc
    from app.services import quality_mgmt as qm
    qsvc.register(U("asist.calidad"), "QUEJA", {"cliente": "Cliente Demo", "motivo": "polvo"}, None, "L-1")
    qsvc.register(U("asist.calidad"), "RECEPCION_MP", qsvc.parse_form("RECEPCION_MP", {"humedad": "12", "grupo_humedad": "TRIGO_CEBADA_CRUDA"}),
                  None, "L-2")
    s.flush()
    data = qm.export_consolidated()
    wb = load_workbook(io.BytesIO(data))
    assert "Resumen" in wb.sheetnames and "QUEJA" in wb.sheetnames and "RECEPCION_MP" in wb.sheetnames
    mp = load_workbook(io.BytesIO(qm.export_consolidated(solo_mp=True)))
    assert "QUEJA" not in mp.sheetnames and "RECEPCION_MP" in mp.sheetnames
    antes = s.query(QualityRecord).count()
    import tempfile, os
    p = os.path.join(tempfile.mkdtemp(), "base.xlsx")
    open(p, "wb").write(data)
    r1 = qm.import_consolidated(U("asist.calidad"), p)
    assert r1["cargados"] == 2
    r2 = qm.import_consolidated(U("asist.calidad"), p)
    assert r2["cargados"] == 0 and r2["omitidos"] == 2
    assert s.query(QualityRecord).count() == antes + 2


def test_pcc1_y_formatos_hora_zona_retirados(app, s, U):
    from app.catalogs import PCC1_HUMEDAD
    from app.services import quality as qsvc
    g = {c: (n, l) for c, n, l in PCC1_HUMEDAD}
    assert "trigo mote" in g["QUINUA_CANIHUA_MOTE"][0].lower() and g["QUINUA_CANIHUA_MOTE"][1] == 12.5
    assert "trigo entero pelado" in g["TRIGO_CEBADA_CRUDA"][0].lower() and g["TRIGO_CEBADA_CRUDA"][1] == 14.5
    tipos = {k: t for k, _, t in qsvc.FORMATOS["CONTROL_PESOS"][2]}
    assert tipos["hora"] == "time"
    assert qsvc.FORMATOS["FUMIGACION"][2][0][2] == "zona"
    d = qsvc.parse_form("FUMIGACION", {"zona": "OTRO", "zona_otro": "contenedor 3"})
    assert d["zona"] == "Otro: contenedor 3"
    with pytest.raises(BusinessError):
        qsvc.parse_form("FUMIGACION", {"zona": "OTRO"})
    with pytest.raises(BusinessError):
        qsvc.register(U("asist.calidad"), "PROCESO", {})
    with pytest.raises(BusinessError):
        qsvc.register(U("asist.calidad"), "LOTE_PRODUCIDO", {})
    html = _get(app, "asist.calidad", "/calidad/")
    nuevo = html.split('name="formato_sel"')[1].split("</select>")[0]  # selector «Nuevo registro»
    assert 'value="PROCESO"' not in nuevo and 'value="LOTE_PRODUCIDO"' not in nuevo and "formato=LOTE_PRODUCIDO" not in html
    assert "Control de limpieza y desinfección de áreas externas" in html
    assert 'type="time"' in _get(app, "asist.calidad", "/calidad/registro/nuevo/CONTROL_PESOS")
    assert "¿dónde?" in _get(app, "asist.calidad", "/calidad/registro/nuevo/FUMIGACION")
    assert _get(app, "jefe.produccion", "/calidad/registro/nuevo/LOTE_PRODUCIDO")


def test_revision_por_fecha_y_varios_tipos(app, s, U):
    from app.services import quality as qsvc
    a = qsvc.register(U("asist.calidad"), "QUEJA", {"cliente": "X"}, fecha=date.today() - timedelta(days=3))
    b = qsvc.register(U("asist.calidad"), "NC_ACCION", {"descripcion": "Y"})
    c = qsvc.register(U("asist.calidad"), "DISENO_BOBINA", {"proveedor": "Z"})
    s.commit()
    html = _get(app, "asist.calidad", f"/calidad/?dia={date.today().isoformat()}&formato=QUEJA&formato=NC_ACCION")
    assert f"rid={b.id}" in html.replace("/calidad/registro/", "rid=") and f"/calidad/registro/{a.id}\"" not in html
    assert f"/calidad/registro/{c.id}\"" not in html


def test_certificado_procesado_y_envasado_y_datos_adicionales(s, U):
    from app.services import quality as qsvc
    lot = _lot(s)
    c = qsvc.generate_certificate(lot, U("jefe.calidad"))
    s.flush()
    qsvc.review_certificate(c, U("jefe.calidad"), {"procesado_etiqueta": "IMPORTADO, PROCESADO Y ENVASADO POR",
                                                   "datos_adicionales": [("Puerto de embarque", "Callao"), ("", "")]})
    assert c.datos["procesado_etiqueta"].startswith("IMPORTADO") and c.datos["datos_adicionales"] == [
        {"etiqueta": "Puerto de embarque", "descripcion": "Callao"}]
    with pytest.raises(BusinessError):
        qsvc.review_certificate(c, U("jefe.calidad"), {"datos_adicionales": [("", "sin etiqueta")]})
    with pytest.raises(BusinessError):
        qsvc.review_certificate(c, U("jefe.calidad"), {"procesado_etiqueta": "OTRA"})


def test_quejas_y_nc_reporte(app, s, U):
    from app.services import quality as qsvc
    from app.services import quality_mgmt as qm
    qsvc.register(U("asist.calidad"), "QUEJA", {"cliente": "Cliente Demo"})
    qsvc.register(U("asist.calidad"), "NC_ACCION", {"estado_accion": "ABIERTA"})
    qsvc.register(U("asist.calidad"), "NC_ACCION", {"estado_accion": "CERRADA"})
    s.commit()
    r = qm.complaints_report()
    assert len(r["quejas"]) == 1 and len(r["ncs"]) == 2 and len(r["abiertas"]) == 1
    assert r["por_cliente"][0] == ("Cliente Demo", 1)
    html = _get(app, "asist.calidad", "/calidad/gestion?tab=quejas")
    assert "Reporte de quejas" in html and "Acciones abiertas" in html


def test_menu_calidad_sin_productividad_asistencia_reposicion(app):
    html = _get(app, "jefe.calidad", "/calidad/")
    for t in ("Productividad por persona", "Análisis de reposición", ">Asistencia<"):
        assert t not in html, t
    for t in ("Bandeja de Calidad", "En la línea y merma", "Llegadas de compras e importaciones"):
        assert t in html, t


def test_paginas_v13_cargan(app, s, U):
    from app.services import quality_mgmt as qm
    i = qm.schedule_inspection(U("asist.calidad"), {"tipo": "INTERNA", "fecha": date.today().isoformat()})
    v = qm.propose_visit(U("asist.calidad"), qm.external_warehouses()[0].code, date.today().isoformat() + "T10:00")
    lot = _lot(s)
    s.commit()
    for p in ("/calidad/linea", "/calidad/gestion", "/calidad/gestion?tab=visitas", f"/calidad/inspeccion/{i.id}",
              f"/calidad/visita/{v.id}", "/calidad/verificacion/LIMPIEZA_LINEA", "/calidad/verificacion/LIMPIEZA_EXTERNAS", "/calidad/higiene-personal", "/calidad/higiene-equipos",
              f"/calidad/lote/{lot.id}/fechas"):
        _get(app, "asist.calidad", p)
    _get(app, "coord.almacen", f"/calidad/visita/{v.id}")
    c = app.test_client()
    login(c, "asist.calidad")
    r = c.get("/calidad/consolidado")
    assert r.status_code == 200 and r.data[:2] == b"PK"


def test_bitacora_tiene_zonas_demo(app):
    html = _get(app, "jefe.logistica", "/mantenimiento/?tab=saneamiento")
    assert "Almacén PT – zona A (demo)" in html


def test_oportunidad_proyeccion_por_producto(s, U, cust):
    from werkzeug.datastructures import MultiDict
    from app.models import Product
    from app.services import opportunities as op
    p1, p2 = [p.code for p in s.query(Product).limit(2)]
    o = op.save_opportunity(U("coord.ventas"), cust(), MultiDict([("titulo", "Ingreso a cadena"), ("productos", p1), ("prod_kg", "100"),
                                                                 ("prod_soles", "500"), ("productos", p2), ("prod_kg", "50"),
                                                                 ("prod_soles", "300")]))
    assert o.detalle_productos == [{"producto": p1, "kg_mes": 100.0, "soles_mes": 500.0},
                                   {"producto": p2, "kg_mes": 50.0, "soles_mes": 300.0}]
    assert o.proyeccion_mensual == 800 and o.proyeccion_kg == 150
    o2 = op.save_opportunity(U("coord.ventas"), cust(), MultiDict([("titulo", "Sin detalle"), ("productos", p1), ("prod_kg", ""),
                                                                  ("prod_soles", ""), ("proyeccion_mensual", "1200")]))
    assert o2.proyeccion_mensual == 1200


def test_devolucion_por_calidad_alerta_a_calidad(s, U):
    from app.models import Incident
    from app.services import incidents_flow as fl
    i = Incident(tipo="INFESTACION_INTERNA", grupo="CALIDAD", entrega="PARCIAL", descripcion="gorgojo")
    j = Incident(tipo="CLIENTE_AUSENTE", grupo="CLIENTE", entrega="NINGUNA", descripcion="no estaba")
    s.add_all([i, j])
    s.flush()
    fl.register_part2(U("coord.distribucion"), i, [{"producto": "PT-905", "cantidad": 2}])
    fl.register_part2(U("coord.distribucion"), j, [{"producto": "PT-905", "cantidad": 1}])
    s.flush()
    assert alerts_of(s, "DEVOLUCION_CALIDAD", i.id) and not alerts_of(s, "DEVOLUCION_CALIDAD", j.id)
