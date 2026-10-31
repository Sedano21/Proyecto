"""Indicaciones del 31-oct-2026 (V14)."""
import io
from datetime import date, timedelta

import pytest
from werkzeug.datastructures import MultiDict

from app.services.core import BusinessError

from .conftest import alerts_of
from .test_web import login


def _get(app, user, path):
    c = app.test_client()
    login(c, user)
    r = c.get(path)
    assert r.status_code == 200, (user, path, r.status_code)
    return r


def test_movimiento_de_granel_y_materia_prima(s):
    from app.models import Product
    from app.services import supply
    from app.services.stock import product_kind
    gr = next(p for p in s.query(Product) if product_kind(p) == "GR")
    mp = next(p for p in s.query(Product) if product_kind(p) == "MP")
    c, det = supply.move_cargo({"tipo": "GRANEL", "kg": "500", "mp": gr.code, "sacos": "10"})
    assert c == {"tipo": "GRANEL", "kg": 500.0, "mp": gr.code, "sacos": 10} and "Granel" in det and "10 sacos" in det
    c, det = supply.move_cargo({"tipo": "MATERIA_PRIMA", "kg": "250", "mp": mp.code})
    assert c["mp"] == mp.code and det.startswith("Materia prima")
    with pytest.raises(BusinessError):  # el artículo debe ser del tipo elegido
        supply.move_cargo({"tipo": "GRANEL", "kg": "1", "mp": mp.code})


def test_capacidad_raumak_1_y_2_y_plan_capacidad_programada_ejecutada(app, s):
    from app.models import ProductionReport
    from app.services import production as psvc
    lunes = date.today() - timedelta(days=date.today().weekday())
    s.add(ProductionReport(linea="R1", tipo="SELECCION", fecha=lunes, product_code="PT-905", peso_entrada_kg=5000,
                           granel_kg=4800, estado="PENDIENTE"))
    s.commit()
    d = psvc.capacity_vs_plan(lunes, lunes)
    r12 = next(x for x in d["filas"] if x["grupo"] == "RAUMAK 1 Y 2")
    assert r12["kg_dia"] == 18700 and r12["capacidad"] == 18700 and r12["ejecutada"] == 4800
    assert r12["pct_ejecutada"] == round(100 * 4800 / 18700, 1)
    html = _get(app, "jefe.produccion", "/produccion/").get_data(as_text=True)
    assert "18,700" in html or "18 700" in html or "18.700" in html
    assert "compartida: RAUMAK 1 Y 2" in html
    html = _get(app, "jefe.produccion", "/produccion/plan").get_data(as_text=True)
    assert "Capacidad real vs programada vs ejecutada" in html


def test_compras_las_aprueba_el_gerente_general_adjunto(s, U):
    from app import catalogs as C
    from app.services.core import has_perm
    assert has_perm(U("gerente.adjunto"), "compras.aprobar_mayor_volumen")
    assert has_perm(U("gerente.adjunto"), "compras.autorizar_bajo_observacion")
    assert not has_perm(U("gerente.operaciones"), "compras.aprobar_mayor_volumen")
    roles = next(t[5] for t in C.ALERT_TYPES if t[0] == "COMPRA_APROBACION_GERENCIA")
    assert "GERENTE_GENERAL_ADJUNTO" in roles and "GERENCIA_OPERACIONES" not in roles
    gg = {p for p, *_ in C.ROLE_PERMISSIONS["GERENCIA_GENERAL"]}
    assert gg <= {p for p, *_ in C.ROLE_PERMISSIONS["GERENTE_GENERAL_ADJUNTO"]}  # mismo poder que GG


def test_lugar_de_entrega_con_varios_contactos_y_correos(s, U, cust):
    from app.services import customers as csv_
    c = cust()
    datos = MultiDict([("nombre", "Tienda Demo"), ("direccion", "Av. Demo 123"), ("distrito", "Lima"), ("provincia", "Lima"),
                       ("ubigeo", "150101"), ("c_nombre", "Contacto Uno"), ("c_telefono", "999000111"),
                       ("c_correo", "uno@demo.pe, dos@demo.pe"), ("c_nombre", "Contacto Dos"), ("c_telefono", ""),
                       ("c_correo", "tres@demo.pe"), ("c_nombre", ""), ("c_telefono", ""), ("c_correo", "")])
    site = csv_.save_site(U("coord.ventas"), c, datos)
    assert len(site.contactos) == 2 and site.contactos[0]["correo"] == "uno@demo.pe, dos@demo.pe"
    assert site.contacto == "Contacto Uno" and site.telefono == "999000111"
    datos.setlist("c_correo", ["no-es-correo", "", ""])
    with pytest.raises(BusinessError):
        csv_.save_site(U("coord.ventas"), c, datos, site)


def test_higiene_del_personal_phs003(s, U):
    from app.services import quality_mgmt as qm
    ok = {k: "C" for k in ("HP", "ES", "UT", "UN", "H")}
    with pytest.raises(BusinessError):  # X sin observación
        qm.register_personal_hygiene(U("asist.calidad"), None, [{"nombre": "Persona A", "area": "SM", **ok, "H": "X"}])
    r = qm.register_personal_hygiene(U("asist.calidad"), None, [
        {"nombre": "Persona A", "area": "SM", **ok}, {"nombre": "Persona B", "area": "P", **ok, "UN": "X", "obs": "se le pidió cortarlas"},
        {"nombre": "", "area": ""}])
    s.flush()
    assert r.codigo_formato == "PHS-003" and len(r.datos["personas"]) == 2 and r.resultado == "NO_CONFORME"
    assert alerts_of(s, "CALIDAD_NO_CONFORME", r.id)


def test_reportes_de_registros_pdf_y_excel(app, s, U):
    from openpyxl import load_workbook
    from app.services import quality_mgmt as qm
    ok = {k: "C" for k in ("HP", "ES", "UT", "UN", "H")}
    r = qm.register_personal_hygiene(U("asist.calidad"), None, [{"nombre": "Persona A", "area": "SM", **ok}])
    n = sum(len(p) for _, p in qm.checklist_points("LIMPIEZA_LINEA")[2])
    c2 = qm.register_checklist(U("asist.calidad"), "LIMPIEZA_LINEA", None, {str(i): "C" for i in range(n)}, {})
    s.commit()
    wb = load_workbook(io.BytesIO(qm.report_excel(qm.report_records(["HIGIENE_PERSONAL", "LIMPIEZA_LINEA"]))))
    assert "HIGIENE_PERSONAL" in wb.sheetnames and "LIMPIEZA_LINEA" in wb.sheetnames
    assert wb["HIGIENE_PERSONAL"]["C3"].value == "Persona A"
    cli = app.test_client()
    login(cli, "asist.calidad")
    assert cli.get(f"/calidad/reporte?id={r.id}&salida=pdf").data[:4] == b"%PDF"
    assert cli.get(f"/calidad/reporte?formato=LIMPIEZA_LINEA&salida=xlsx").data[:2] == b"PK"
    html = cli.get(f"/calidad/reporte?dia={date.today().isoformat()}").get_data(as_text=True)
    assert "Persona A" in html and "PHS-002" in html
    assert cli.get(f"/calidad/registro/{r.id}").status_code == 302  # el PHS-003 se ve en el reporte
    assert "Control de limpieza y desinfección del área de óptica" in html


def test_ficha_de_muestreo_aprueba_la_muestra_de_la_compra(s, U):
    from app.models import PurchaseRequest
    from app.services import quality as qsvc
    r = PurchaseRequest(product_code="MP-901", cantidad=10, estado="MUESTRA_EN_CALIDAD", requiere_muestra=True)
    s.add(r)
    s.flush()
    qsvc.register(U("asist.calidad"), "MUESTREO", {"humedad": 12}, "MP-901", "L-M", None, "APROBADO", "ok", "COMPRA", r.id)
    assert r.estado == "MUESTRA_EN_CALIDAD"  # la ficha queda; la decisión la toma quien evalúa muestras
    qsvc.register(U("jefe.calidad"), "MUESTREO", {"humedad": 12}, "MP-901", "L-M", None, "APROBADO", "ok", "COMPRA", r.id)
    assert r.estado == "EN_COTIZACION" and r.resultado_muestra == "APROBADA"


def test_fumigacion_prellenada_desde_el_lote(app, s, U):
    from app.models import StockLot
    from app.services import quality_v13 as q
    lot = s.query(StockLot).filter(StockLot.estado == "DISPONIBLE", StockLot.cantidad > 0).first()
    ini, lib = date.today(), date.today() + timedelta(days=5)
    lot = q.set_fumigation(U("asist.calidad"), lot, ini, lib, "Almacén PT – zona A (demo)")
    s.commit()
    html = _get(app, "asist.calidad", f"/calidad/registro/nuevo/FUMIGACION?lot_id={lot.id}").get_data(as_text=True)
    assert f'name="inicio_fecha" value="{ini.isoformat()}"' in html and f'name="fin_fecha" value="{lib.isoformat()}"' in html


def test_calificacion_de_inspecciones():
    from app.services.quality_mgmt import calificacion
    assert [calificacion(x) for x in (50, 69, 70, 85, 86, 95, 96, 100)] == [
        "DESAPROBADO", "DESAPROBADO", "REGULAR", "REGULAR", "BUENO", "BUENO", "MUY_BUENO", "MUY_BUENO"]


def test_precertificado_toma_registros_del_proceso(s, U):
    from app.models import ProductionReport, QualityRecord, StockLot
    from app.services import quality as qsvc
    lot = s.query(StockLot).filter(StockLot.cantidad > 0).first()
    s.add(ProductionReport(linea="TRUER", tipo="SELECCION", fecha=date.today(), product_code=lot.product_code, peso_entrada_kg=100,
                           lote="MP-LOTE-DEMO", lote_pt=lot.lote, estado="PENDIENTE"))
    s.add(QualityRecord(formato="CONTROL_GRANOS", codigo_formato="REG-CAL-014", lote="MP-LOTE-DEMO",
                        datos={"humedad": 11.8, "grano_partido": 0.5}, resultado="CONFORME", user_id=U("asist.calidad").id))
    s.flush()
    d = qsvc.certificate_data(lot)
    assert d["humedad"]["resultado"] == "11.8%"


def test_inicio_de_calidad_lineas_en_proceso_y_firma(app, s, U):
    from app.models import ProductionReport, now
    s.add(ProductionReport(linea="TRUER", tipo="SELECCION", fecha=date.today(), product_code="PT-905", peso_entrada_kg=1000,
                           estado="EN PROCESO", inicio_ts=now() - timedelta(minutes=60)))
    s.commit()
    html = _get(app, "jefe.calidad", "/calidad/").get_data(as_text=True)
    assert "Líneas en proceso de producción (1)" in html and "merma" in html
    assert "Firma de la Jefatura de Calidad" in _get(app, "jefe.calidad", "/calidad/certificados").get_data(as_text=True)


def test_interfaz_hora_fotos_y_notificaciones(app):
    html = _get(app, "asist.calidad", "/calidad/registro/nuevo/CONTROL_PESOS").get_data(as_text=True)
    assert "v14.js" in html and 'type="time"' in html and 'name="adjuntos"' in html
    js = app.test_client().get("/static/v14.js").get_data(as_text=True)
    for t in ("Ahora", "Cámara", "+ Añadir otro", "ev-thumb"):
        assert t in js, t
    css = app.test_client().get("/static/style.css").get_data(as_text=True)
    assert ".toast{background:#c62828" in css


# ------------------------------------------------------------------ V15: conexión con SAP tolerante a cortes
class _SL:
    """Service Layer simulado: existente (OV ya creada en un envío cortado), estampa de SAP y caída de conexión."""
    def __init__(self, existente=None, stamp=None, caido=False):
        self.cfg = {"company": "ZZZ_CORPRODUC_TEST260826"}
        self.llamadas, self.existente, self.stamp, self.caido = [], existente, stamp or {}, caido

    def _req(self, method, path, body=None, headers=None):
        from app.services.sap_sl import SapConnectionError
        self.llamadas.append((method, path))
        if self.caido:
            raise SapConnectionError("No se pudo conectar con SAP: Connection refused")
        if method == "GET" and path.startswith("Orders?"):
            return {"value": [self.existente] if self.existente else []}
        if method == "GET" and "$select=UpdateDate" in path:
            return dict({"UpdateDate": "2026-10-31T00:00:00Z", "UpdateTime": "10:00:00", "DocumentStatus": "bost_Open",
                         "Cancelled": "tNO"}, **self.stamp)
        if method == "POST" and path == "Orders":
            return {"DocEntry": 90, "DocNum": 5090}
        return {}


def _ov_sap(s, U, cust):
    from app.services import orders as osvc
    c = cust()
    c.codigo_sap = "C20000000001"
    o = osvc.create_order(U("coord.ventas"), "SAP-V15", c.id, date.today() + timedelta(days=4),
                          [{"product_code": "PT-905", "cantidad": 5, "precio_unitario": 999}], oc_archivo="oc.pdf",
                          oc_numero="OC-9", evaluar=False)
    s.flush()
    if o.retenida_precio or o.estado == "BORRADOR":
        pytest.skip("la demo retuvo la OV")
    return o


def test_sap_no_duplica_una_ov_ya_creada(s, U, cust):
    from app.services import sap_sl
    o = _ov_sap(s, U, cust)
    sl = _SL(existente={"DocEntry": 55, "DocNum": 5055})
    assert sap_sl.push_order(o, None, sl) == "CREAR"
    assert ("POST", "Orders") not in sl.llamadas and o.sap_docnum == "5055"


def test_sap_sin_conexion_deja_la_ov_en_cola(s, U, cust):
    from app.services import sap_sl
    o = _ov_sap(s, U, cust)
    with pytest.raises(sap_sl.SapConnectionError):
        sap_sl.push_order(o, None, _SL(caido=True))
    assert o.sap_envio["estado"] == "EN_COLA" and not alerts_of(s, "SAP_OV_ERROR", o.id)
    assert sap_sl.pending_action(o) == "CREAR"
    assert sap_sl.push_order(o, None, _SL()) == "CREAR" and o.sap_docnum == "5090"


def test_sap_cambiado_en_sap_no_se_sobrescribe(s, U, cust):
    from app.services import orders as osvc
    from app.services import sap_sl
    o = _ov_sap(s, U, cust)
    sap_sl.push_order(o, None, _SL())
    osvc.modify_order(o, U("coord.ventas"), {"lines": [{"product_code": "PT-905", "cantidad": 8, "precio_unitario": 999}]})
    sl = _SL(stamp={"UpdateTime": "11:30:00"})  # alguien la editó en SAP
    assert sap_sl.push_order(o, None, sl) is None
    assert o.sap_envio["conflicto"] and not any(m == "PATCH" for m, _ in sl.llamadas)
    assert sap_sl.pending_action(o) is None and alerts_of(s, "SAP_OV_ERROR", o.id)
    sl2 = _SL(stamp={"UpdateTime": "11:30:00"})
    assert sap_sl.push_order(o, None, sl2, forzar=True) == "ACTUALIZAR" and any(m == "PATCH" for m, _ in sl2.llamadas)
    assert not o.sap_envio.get("conflicto")
    sl3 = _SL(stamp={"DocumentStatus": "bost_Close"})
    osvc.modify_order(o, U("coord.ventas"), {"lines": [{"product_code": "PT-905", "cantidad": 9, "precio_unitario": 999}]})
    assert sap_sl.push_order(o, None, sl3) is None and "cerrada" in o.sap_envio["conflicto"]


def test_estado_de_conexion_sap_alerta_y_aviso(app, s, U):
    from app.services import sap_sl
    sap_sl._set_state(False, "No se pudo conectar con SAP: Connection refused")
    s.commit()
    assert sap_sl.connection_state()["ok"] is False and alerts_of(s, "SAP_SIN_CONEXION")
    html = _get(app, "coord.ventas", "/").get_data(as_text=True)
    assert "SAP sin conexión desde" in html
    sap_sl._set_state(True)
    s.commit()
    assert "SAP sin conexión desde" not in _get(app, "coord.ventas", "/").get_data(as_text=True)


def test_ruc_de_clientes_sap_y_diagnostico(app, s, monkeypatch):
    from app.services import sap_sl
    assert sap_sl._ruc_de({"FederalTaxID": "", "CardCode": "C20123456789"}) == "20123456789"
    assert sap_sl._ruc_de({"FederalTaxID": "RUC 20-1234567-89", "CardCode": "X"}) == "20123456789"

    class SL(_SL):
        def login(self):
            return {"Version": "1000160", "SessionTimeout": 30}

        def logout(self):
            pass

        def get_all(self, path, params=None):
            return [{"WarehouseCode": "ALM-001"}] if path == "Warehouses" else [{"CardCode": "C1", "FederalTaxID": ""}]

        def _req(self, method, path, body=None, headers=None):
            return 7 if "$count" in path else {}
    monkeypatch.setattr(sap_sl, "ServiceLayer", SL)
    out = "\n".join(sap_sl.diagnostics())
    assert "Ingreso a SAP: OK (versión 1000160" in out and "Clientes: 7" in out
    assert "NO existen en SAP" in out and "Clientes sin RUC en SAP (FederalTaxID vacío): 1" in out


def test_excel_de_conteo(app, s, U):
    from app.services import inventory as isvc
    c = isvc.open_count(U("coord.almacen"), "ALM-001")
    s.commit()
    cli = app.test_client()
    login(cli, "coord.almacen")
    assert cli.get(f"/almacen/conteos/{c.id}/excel").data[:2] == b"PK"


def test_aviso_del_sistema_y_login_responsivo(app):
    js = app.test_client().get("/static/v14.js").get_data(as_text=True)
    assert "cpNotificar" in js and "Notification.requestPermission" in js
    css = app.test_client().get("/static/style.css").get_data(as_text=True)
    assert ".public{flex-direction:column" in css
