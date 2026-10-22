"""Marcha blanca (respuestas del 22-oct-2026): aviso de datos y contraseñas, observaciones, turnos y cumplimiento,
conciliación con SAP, contingencia y guía de SAP por defecto."""
import io
import os
import re
from datetime import date, datetime, timedelta

import openpyxl
import pytest

from app.services import uploads as up
from app.services.core import BusinessError, param, set_param

from .conftest import alerts_of
from .test_web import csrf, login


def _nuevo_usuario(s, username="persona.real", role="AUX_ALMACEN", clave="Inicial01"):
    from app.seed import create_user
    u = create_user(username, "Persona Real", role, clave)
    u.debe_cambiar_clave = True
    u.almacen = "ALM-001"
    s.commit()
    return u


# ------------------------------------------------------------------ acceso
def test_primer_ingreso_aviso_y_cambio_de_clave(app, s):
    _nuevo_usuario(s)
    c = app.test_client()
    tok, r = login(c, "persona.real", "Inicial01")
    assert r.status_code == 302
    r = c.get("/")
    assert r.status_code == 302 and "/aviso-datos" in r.headers["Location"]  # primero el aviso de datos
    html = c.get("/aviso-datos").get_data(as_text=True)
    assert "Corporación Productiva" in html and "29733" in html and "contraseña" in html
    tok = csrf(c, "/aviso-datos")
    r = c.post("/aviso-datos", data={"_csrf": tok, "accion": "aceptar", "acepto_datos": "1"})  # falta una casilla
    assert "/aviso-datos" in r.headers["Location"]
    c.post("/aviso-datos", data={"_csrf": tok, "accion": "aceptar", "acepto_datos": "1", "acepto_responsabilidad": "1"})
    r = c.get("/")
    assert r.status_code == 302 and "/cambiar-clave" in r.headers["Location"]  # luego la contraseña propia
    tok = csrf(c, "/cambiar-clave")
    c.post("/cambiar-clave", data={"_csrf": tok, "actual": "Inicial01", "nueva": "corta", "repetir": "corta"})
    assert c.get("/").status_code == 302  # menos de 8 caracteres: no se acepta
    c.post("/cambiar-clave", data={"_csrf": tok, "actual": "Inicial01", "nueva": "Nueva2026", "repetir": "Nueva2026"})
    assert c.get("/").status_code == 200
    from app.models import AuditLog
    assert s.query(AuditLog).filter_by(accion="AVISO_DATOS_ACEPTADO").count() == 1


def test_rechazar_aviso_cierra_sesion(app, s):
    _nuevo_usuario(s)
    c = app.test_client()
    login(c, "persona.real", "Inicial01")
    tok = csrf(c, "/aviso-datos")
    r = c.post("/aviso-datos", data={"_csrf": tok, "accion": "rechazar"})
    assert "/login" in r.headers["Location"]
    assert c.get("/").status_code == 302


def test_nueva_version_del_aviso_se_vuelve_a_pedir(app, s, U):
    c = app.test_client()
    login(c, "coord.ventas")
    assert c.get("/").status_code == 200
    set_param("aviso_datos_version", "1.0")
    s.commit()
    r = c.get("/")
    assert r.status_code == 302 and "/aviso-datos" in r.headers["Location"]


def test_bloqueo_tras_5_intentos_y_desbloqueo(app, s, U):
    c = app.test_client()
    for _ in range(5):
        login(c, "coord.ventas", "incorrecta")
    tok, r = login(c, "coord.ventas")  # contraseña correcta, pero bloqueado
    assert r.status_code == 200 and "bloqueado" in r.get_data(as_text=True)
    s.expire_all()
    assert U("coord.ventas").bloqueado_hasta is not None
    adm = app.test_client()
    login(adm, "sistemas")
    tok = csrf(adm, "/usuarios")
    adm.post("/usuarios", data={"_csrf": tok, "id": U("coord.ventas").id, "accion": "desbloquear"})
    s.expire_all()
    assert U("coord.ventas").bloqueado_hasta is None
    tok, r = login(c, "coord.ventas")
    assert r.status_code == 302


def test_clave_vencida_pide_cambio(app, s, U):
    u = U("coord.ventas")
    u.clave_ts = datetime.now() - timedelta(days=param("clave_dias_vigencia") + 1)
    s.commit()
    c = app.test_client()
    login(c, "coord.ventas")
    r = c.get("/ov/")
    assert r.status_code == 302 and "/cambiar-clave" in r.headers["Location"]
    assert "venció" in c.get("/cambiar-clave").get_data(as_text=True)


# ------------------------------------------------------------------ observaciones
def test_reportar_observacion_y_atenderla(app, s, U):
    c = app.test_client()
    login(c, "aux.almacen")
    html = c.get("/").get_data(as_text=True)
    assert "Reportar observación" in html and 'id="obsdlg"' in html
    tok = csrf(c, "/")
    r = c.post("/observaciones/nueva", data={"_csrf": tok, "comentario": "El botón no responde", "url": "/almacen/picking",
                                             "pantalla": "Programación de picking", "tipo": "ERROR", "impide": "1",
                                             "captura": (io.BytesIO(b"\x89PNG demo"), "captura.png")},
               content_type="multipart/form-data")
    assert r.status_code == 302 and r.headers["Location"].endswith("/almacen/picking")
    from app.models import Observation
    o = s.query(Observation).one()
    assert o.clasificacion == "BLOQUEANTE" and o.captura and o.version
    assert alerts_of(s, "OBSERVACION_REPORTADA")
    assert "El botón no responde" in c.get("/observaciones").get_data(as_text=True)  # «Mis observaciones»
    r = c.post("/observaciones", data={"_csrf": tok, "id": o.id, "estado": "CORREGIDA", "respuesta": "x"})
    s.expire_all()
    assert o.estado == "ABIERTA"  # el auxiliar no la atiende
    gg = app.test_client()
    login(gg, "gerente.general")
    tok = csrf(gg, "/observaciones")
    gg.post("/observaciones", data={"_csrf": tok, "id": o.id, "clasificacion": "MAYOR", "estado": "CORREGIDA",
                                    "respuesta": "Corregido en la versión de la noche"})
    s.expire_all()
    assert o.estado == "CORREGIDA" and o.clasificacion == "MAYOR"
    assert gg.get("/observaciones?formato=xlsx").status_code == 200


# ------------------------------------------------------------------ turnos
def test_turnos_rotativos_y_nocturno(s, U, today):
    from app.services import shifts as ss
    adm = U("sistemas")
    lunes = today - timedelta(days=today.weekday())
    dia = ss.save_shift(adm, "Día L-S", [0, 1, 2, 3, 4, 5], "07:00", "16:00", 60)
    noche = ss.save_shift(adm, "Noche L-S", [0, 1, 2, 3, 4, 5], "22:00", "06:00", 30)
    p = U("aux.almacen")
    sem = lunes.isocalendar()[1] % 2
    ss.assign_shift(adm, p.id, dia.id, lunes - timedelta(days=14), None, 2, sem)       # semanas pares: día
    ss.assign_shift(adm, p.id, noche.id, lunes - timedelta(days=14), None, 2, 1 - sem)  # semanas impares: noche
    s.commit()
    assert ss.works_on(p.id, lunes) and ss.windows(p.id, lunes)[0][3] == "Día L-S"
    otra = lunes + timedelta(days=7)
    assert ss.windows(p.id, otra)[0][3] == "Noche L-S"
    assert ss.in_shift_at(p.id, datetime.combine(otra + timedelta(days=1), datetime.min.time()).replace(hour=3))
    assert not ss.works_on(p.id, lunes + timedelta(days=6))  # domingo
    with pytest.raises(BusinessError):
        ss.save_shift(adm, "x", [0], "08:00", "09:00", 90)  # refrigerio mayor que el turno
    with pytest.raises(BusinessError):
        ss.save_shift(U("aux.almacen"), "x", [0], "08:00", "16:00", 0)  # sin permiso


def test_picking_solo_personal_en_turno_y_refuerzo(s, U, new_order, today):
    from app.models import PickingTask
    from app.services import picking as pk
    from app.services import shifts as ss
    adm = U("sistemas")
    t = ss.save_shift(adm, "Todos los días", list(range(7)), "07:00", "16:00", 60)
    turno = U("aux.almacen2")
    ss.assign_shift(adm, turno.id, t.id, today - timedelta(days=1))
    s.commit()
    o = new_order([("PT-905", 20)], dias=2)
    from app.services import orders as osvc
    osvc.validate_picking_plan(o, U("lider.almacen"))
    tarea = s.query(PickingTask).filter_by(order_id=o.id).one()
    assert tarea.asignados == [turno.id]  # el único del almacén con turno
    otro = U("aux.almacen")
    with pytest.raises(BusinessError, match="no está"):
        pk.assign(tarea, U("lider.almacen"), [otro.id])
    ss.add_reinforcement(adm, otro.id, tarea.fecha, tarea.fecha, "08:00", "14:00", "Pedido urgente")
    pk.assign(tarea, U("lider.almacen"), [otro.id])
    assert tarea.asignados == [otro.id]


def test_tarea_vencida_alerta_al_jefe_una_sola_vez(s, U):
    from app.models import Alert
    from app.services import alerts
    from app.services.shifts import check_overdue
    a = alerts.emit("LIQUIDACION_MONTACARGUISTA", "Mover PT de prueba", "", "Liquidation", 1)
    a.asignado_a = U("montacarguista").id
    a.creado_ts = datetime.now() - timedelta(hours=30)
    s.commit()
    assert check_overdue() == 1
    av = s.query(Alert).filter_by(type_code="TAREA_INCUMPLIDA").one()
    assert "LIDER_ALMACEN" in av.roles and "Montacarguista" in av.mensaje
    av.estado = "COMPLETADA"
    s.commit()
    assert check_overdue() == 0  # no se repite aunque el jefe ya la cerró


def test_informe_cumplimiento(app, s, U):
    c = app.test_client()
    login(c, "coord.ventas")
    c.get("/")
    gg = app.test_client()
    login(gg, "gerente.general")
    html = gg.get("/informes/cumplimiento").get_data(as_text=True)
    assert "Interacciones y cumplimiento" in html and "Coordinador de Ventas" in html
    assert gg.get("/informes/cumplimiento?formato=xlsx").status_code == 200
    jefe = app.test_client()
    login(jefe, "aux.almacen")
    assert jefe.get("/informes/cumplimiento").status_code == 403


# ------------------------------------------------------------------ conciliación con SAP
def test_conciliacion_stock_alerta_y_resolucion(app, s, U, tmp_path, today):
    from app.demo_data import write_stock
    from app.models import SapDiscrepancy
    from app.services import conciliation as cs
    p = os.path.join(tmp_path, "stock.xlsx")
    write_stock(p, today, extra_rows=[("PT-901", "LZZ9", "ALM-001", 7, today + timedelta(days=200), today, today)])
    b = up.create_batch(U("coord.almacen"), "STOCK", p, "stock.xlsx")
    res = up.confirm(b, U("coord.almacen"))
    s.commit()
    assert res["diferencias_sap"] >= 1
    d = s.query(SapDiscrepancy).filter_by(lote="LZZ9").one()
    assert d.valor_sap == "7" and d.valor_modulo == "0" and d.area == "Almacén"
    assert alerts_of(s, "CONCILIACION_STOCK")
    with pytest.raises(BusinessError):
        cs.resolve(U("coord.ventas"), [d.id], "LEVANTADA", "x")  # otra área
    pend = [x.id for x in s.query(SapDiscrepancy).filter_by(upload_id=b.id)]
    cs.resolve(U("coord.almacen"), pend, "REGULARIZADA", "Ingreso registrado en SAP el sábado")
    s.commit()
    a = alerts_of(s, "CONCILIACION_STOCK")[0]
    assert a.estado == "COMPLETADA" or a.resultado == "SISTEMA"
    c = app.test_client()
    login(c, "coord.almacen")
    assert "Conciliación SAP" in c.get("/conciliacion?estado=").get_data(as_text=True)


def test_conciliacion_ov_alerta_sin_modificar(s, U, tmp_path, today):
    from app.demo_data import write_ovs
    from app.models import SalesOrder, SapDiscrepancy
    p = os.path.join(tmp_path, "ovs.xlsx")
    write_ovs(p, today)
    up.confirm(up.create_batch(U("coord.ventas"), "OV", p, "ovs.xlsx"), U("coord.ventas"))
    o = s.query(SalesOrder).filter_by(numero="SAP-9001").one()
    assert o.estado_stock == "COMPLETA"  # la OV abierta cargada reserva stock al ingresar
    wb = openpyxl.load_workbook(p)
    ws = wb["OVS"]
    for row in ws.iter_rows(min_row=2):
        if row[1].value == "SAP-9001" and row[12].value == "PT-905":
            row[14].value = 500
    wb.save(p)
    res = up.confirm(up.create_batch(U("coord.ventas"), "OV", p, "ovs.xlsx"), U("coord.ventas"))
    s.commit()
    d = s.query(SapDiscrepancy).filter_by(referencia="SAP-9001", product_code="PT-905").one()
    assert d.valor_sap == "500" and d.valor_modulo == "400" and res["ov_modificadas"] == 0
    assert [l.cantidad for l in o.lines if l.product_code == "PT-905"] == [400]  # no se cambia: se alerta
    assert alerts_of(s, "CONCILIACION_OV")


def test_conciliacion_despachos_guia_y_factura(s, U, tmp_path, new_order):
    from app.models import SapDiscrepancy
    o = new_order([("PT-905", 10)])
    p = os.path.join(tmp_path, "desp.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DESPACHOS"
    ws.append(["Número de OV", "Número de guía", "Número de factura", "Fecha", "Código de producto", "Cantidad"])
    ws.append([o.numero, "T001-000777", "F001-000888", date.today(), "PT-905", 10])
    ws.append(["OV-NO-EXISTE", "T001-1", "", date.today(), "PT-905", 1])
    wb.save(p)
    b = up.create_batch(U("coord.distribucion"), "DESPACHOS", p, "desp.xlsx")
    assert b.estado != "VALIDADO_CON_ERRORES"
    res = up.confirm(b, U("coord.distribucion"))
    s.commit()
    campos = {(d.origen, d.campo) for d in s.query(SapDiscrepancy).filter_by(upload_id=b.id)}
    assert ("DESPACHO", "Guía") in campos and ("FACTURA", "Factura") in campos
    assert ("DESPACHO", "Cantidad despachada") in campos and ("DESPACHO", "Existencia") in campos
    assert res["diferencias_sap"] == 4 and alerts_of(s, "CONCILIACION_FACTURA")


# ------------------------------------------------------------------ contingencia y guía
def test_contingencia_marca_lo_registrado(app, s, U, new_order):
    from app.models import AuditLog
    c = app.test_client()
    login(c, "sistemas")
    tok = csrf(c, "/contingencia")
    c.post("/contingencia", data={"_csrf": tok, "accion": "activar", "motivo": "Caída de 08:00 a 09:30"})
    assert "Registro en contingencia activo" in c.get("/").get_data(as_text=True)
    o = new_order([("PT-905", 5)])
    ult = s.query(AuditLog).filter_by(entidad="SalesOrder").order_by(AuditLog.id.desc()).first()
    assert ult.detalle.get("registrado_en_contingencia") is True
    c.post("/contingencia", data={"_csrf": tok, "accion": "desactivar"})
    assert "Registro en contingencia activo" not in c.get("/").get_data(as_text=True)
    assert alerts_of(s, "CONTINGENCIA")
    assert o.numero in c.get("/contingencia").get_data(as_text=True)


def test_guia_del_sistema_solo_como_contingencia(s, U, new_order):
    from app.models import RouteStop
    from app.services import despacho as dsp
    from app.services import distribution as dist
    o = new_order([("PT-905", 10)])
    from app.services import orders as osvc
    osvc.validate_picking_plan(o, U("lider.almacen"))
    for l in o.lines:
        for r in [r for r in l.reservations if r.estado == "ACTIVA"]:
            osvc.register_picking(o, U("lider.almacen"), l.id, r.cantidad, r.lot.lote)
    st = dist.active_stop(o)
    dist.confirm_route(st.route, U("coord.distribucion"))
    with pytest.raises(BusinessError, match="contingencia"):
        dsp.set_guide(st, U("coord.distribucion"), "", "SISTEMA", "", motivo="")
    dsp.set_guide(st, U("coord.distribucion"), "", "SISTEMA", "", motivo="SAP no disponible")
    assert st.guia_origen == "SISTEMA" and "contingencia" in st.observacion
