"""Rediseño de informes y tableros (28-oct-2026): pertinencia de indicadores por perfil, informe predeterminado y
personalizado con plantillas; las exclusiones rigen en todas las vistas y exportaciones."""
import io

import pytest

from .test_web import login


def _get(app, user, path):
    c = app.test_client()
    login(c, user)
    r = c.get(path)
    assert r.status_code == 200, (user, path, r.status_code)
    return r


def test_matriz_produccion_no_ve_indicadores_comerciales(s, U):
    from app.services import indicators as ind
    jp = U("jefe.produccion")
    for c in ("OTIF_SALIDA", "FILL_RATE", "MONTO_OV", "PRECIO_BAJO_MINIMO", "VOLUMEN_OV_KG", "MONTOS_VENTA",
              "PED_ESTADOS", "T_PREPARACION"):
        assert ind.level(jp, c) == ind.NO_VISIBLE, c
    for c in ("PROD_CUMPLIMIENTO", "PROD_PRODUCIDO", "PROD_RENDIMIENTO", "PRODUCCION_PENDIENTE", "FALTANTES"):
        assert ind.level(jp, c) == ind.OBLIGATORIO, c
    jv = U("jefe.ventas")
    assert ind.level(jv, "OTIF_SALIDA") == ind.OBLIGATORIO and ind.level(jv, "MONTO_OV") == ind.OBLIGATORIO
    assert ind.level(jv, "PROD_RENDIMIENTO") == ind.NO_VISIBLE
    gg = U("gerente.general")
    assert ind.default_set(gg) >= {"OTIF_SALIDA", "PROD_CUMPLIMIENTO", "MONTO_OV"}


def test_inicio_de_produccion_muestra_su_operacion(app):
    html = _get(app, "jefe.produccion", "/").get_data(as_text=True)
    assert "Cumplimiento del programa" in html and "Rendimiento de las líneas" in html
    for ajeno in ("OTIF a la salida", "Fill rate", "Monto de OV", "Volumen de OV", "Pedidos por estado"):
        assert ajeno not in html, ajeno


def test_inicio_comercial_sin_indicadores_de_produccion(app):
    html = _get(app, "jefe.ventas", "/").get_data(as_text=True)
    assert "OTIF" in html and "OV ingresadas" in html and "Volumen de OV" in html
    # sin duplicar: OTIF, fill rate y montos ya están en «Gestión comercial del mes»
    assert "OTIF a la salida" not in html and "Monto de OV" not in html
    assert "Cumplimiento del programa" not in html and "Rendimiento de las líneas" not in html
    html = _get(app, "coord.distribucion", "/").get_data(as_text=True)
    assert "OTIF a la salida" in html and "Monto de OV" not in html and "Cumplimiento del programa" not in html


def test_exclusion_configurable_desaparece_de_todas_las_vistas(app, s, U):
    from app.services import indicators as ind
    from app.services.core import set_param
    set_param("indicadores_ocultos", {"Comercial": ["VOLUMEN_OV_KG"], "JEFE_PRODUCCION": ["PROD_OEE"]})
    s.commit()
    try:
        assert ind.level(U("jefe.ventas"), "VOLUMEN_OV_KG") == ind.NO_VISIBLE
        assert ind.level(U("jefe.logistica"), "VOLUMEN_OV_KG") == ind.OBLIGATORIO  # no afecta a otras áreas
        assert "Volumen de OV" not in _get(app, "jefe.ventas", "/").get_data(as_text=True)
        # por perfil: desaparece del inicio, del personalizado y de su exportación
        assert "OEE promedio" not in _get(app, "jefe.produccion", "/").get_data(as_text=True)
        assert "OEE (disponibilidad" not in _get(app, "jefe.produccion", "/informes/personalizado").get_data(as_text=True)
        x = _get(app, "jefe.produccion", "/informes/personalizado?fuente=partes&cfg=1&ind=PROD_OEE&ind=PROD_PRODUCIDO"
                                         "&elem=indicadores&formato=pdf")
        assert x.data[:4] == b"%PDF"
        from app.services.indicators import values
        assert [v["codigo"] for v in values(U("jefe.produccion"), ["PROD_OEE", "PROD_PRODUCIDO"])] == ["PROD_PRODUCIDO"]
    finally:
        set_param("indicadores_ocultos", {})
        s.commit()


def test_portada_de_informes_con_predeterminado_y_personalizar(app):
    html = _get(app, "jefe.produccion", "/informes/").get_data(as_text=True)
    assert "Informe predeterminado" in html and "Personalizar informe" in html
    assert "/informes/produccion" in html and "Gestión comercial" not in html
    html = _get(app, "jefe.ventas", "/informes/").get_data(as_text=True)
    assert "/informes/gestion" in html


def test_cada_area_tiene_su_predeterminado(s, U):
    from app.services import report_center as rc
    from app.services.custom_report import sources_for
    from app.views.informes import INFORMES

    def pred(u):
        disp = {k: t for k, (t, perms) in INFORMES.items() if any(p in __import__("app.services.core", fromlist=["x"]).user_perms(u) for p in perms)}
        return rc.default_for(u, disp, sources_for(u))
    assert pred(U("jefe.produccion"))["nombre"] == "produccion"
    assert pred(U("jefe.ventas"))["nombre"] == "gestion"
    assert pred(U("coord.distribucion"))["nombre"] == "distribucion"
    p = pred(U("compras"))
    assert p and p["tipo"] == "personalizado" and p["config"]["fuente"] == "compras"


def test_personalizado_sin_montos_para_produccion_ni_en_exportaciones(app, s, U):
    from app.services import custom_report as cr
    jp = U("jefe.produccion")
    cols = [c for c, _, _ in cr.columns("ov_lineas", jp)]
    assert "precio" not in cols and "importe" not in cols and "cantidad" in cols
    r = cr.run(jp, "ov_lineas", ["producto", "precio", "importe"])  # aunque los pida, no salen
    assert "Precio S/" not in r["columnas"] and "Importe S/" not in r["columnas"]
    assert "importe" in [c for c, _, _ in cr.columns("ov_lineas", U("jefe.ventas"))]
    c = app.test_client()
    from .test_web import login
    login(c, "jefe.produccion")
    x = c.get("/informes/personalizado?fuente=ov_lineas&col=producto&col=precio&col=importe&ind=MONTO_OV&ind=OTIF_SALIDA"
              "&ind=PROD_PRODUCIDO&elem=indicadores&elem=tabla&formato=xlsx")
    assert x.status_code == 200
    from openpyxl import load_workbook
    texto = " ".join(str(v) for row in load_workbook(io.BytesIO(x.data)).active.iter_rows(values_only=True) for v in row if v)
    assert "Precio S/" not in texto and "Importe S/" not in texto and "Monto de OV" not in texto and "OTIF" not in texto
    assert "Producido frente al plan" in texto
    html = c.get("/informes/personalizado?fuente=ov_lineas").get_data(as_text=True)
    assert "Monto de OV" not in html and "OTIF" not in html and 'value="precio"' not in html


def test_plantillas_personales_y_compartidas(app, s, U):
    from werkzeug.datastructures import MultiDict
    from app.services import report_center as rc
    from app.services.core import BusinessError
    jp, ap = U("jefe.produccion"), U("asist.produccion")
    cfg = rc.config_from_args(MultiDict([("fuente", "partes"), ("col", "linea"), ("col", "rendimiento"),
                                         ("agrupar", "linea"), ("ind", "PROD_RENDIMIENTO"), ("desde", "2026-10-01"),
                                         ("hasta", "2026-10-08")]))
    assert cfg["dias"] == 7 and cfg["col"] == ["linea", "rendimiento"]
    t = rc.save_template(jp, "Rendimiento por línea", cfg, compartir=True)
    assert t.alcance == "COMPARTIDA" and t.area == "Producción"
    p = rc.save_template(ap, "Mi vista", cfg)
    s.flush()
    assert t in rc.templates_for(ap) and p in rc.templates_for(ap)
    assert t not in rc.templates_for(U("jefe.ventas")) and p not in rc.templates_for(jp)
    with pytest.raises(BusinessError, match="jefaturas"):
        rc.save_template(U("asist.calidad"), "x", cfg, compartir=True)  # sin permiso no publica para el área
    with pytest.raises(BusinessError):
        rc.delete_template(ap, t.id)
    assert rc.get_template(ap, t.id).config["agrupar"] == "linea"


def test_plantilla_compartida_no_abre_columnas_vedadas(app, s, U):
    """Una plantilla con montos guardada por Comercial no los muestra a un perfil que no puede verlos."""
    from app.services import custom_report as cr
    from app.services import report_center as rc
    from werkzeug.datastructures import MultiDict
    cfg = rc.config_from_args(MultiDict([("fuente", "ov_lineas"), ("col", "producto"), ("col", "importe")]))
    t = rc.save_template(U("jefe.ventas"), "Importe por producto", cfg)
    s.flush()
    r = cr.run(U("jefe.produccion"), "ov_lineas", t.config["col"])
    assert r["columnas"] == ["Producto"]


def test_informe_produccion_con_cabecera_institucional(app):
    html = _get(app, "jefe.produccion", "/informes/produccion").get_data(as_text=True)
    assert 'class="rep-head"' in html and "logo" in html and "Actualizado" in html and "Fuentes:" in html
    assert html.index("Cumplimiento del plan diario") < html.index("1. Capacidad de producción")
